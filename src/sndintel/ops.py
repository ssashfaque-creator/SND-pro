"""Monday / DSR / Friday operating packs on top of the this-week engine.

Monday — NSM one-pager: hole, driver, whales, DSR labels, Ask KG.
DSR beat — capacity-capped call list with owner and ask.
Friday — closed loop: listed → visited → billed vs Ask.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, timezone
from io import BytesIO
from typing import Any

import pandas as pd

from sndintel.action import (
    ACTION_CALL,
    ACTION_CONVERT,
    ACTION_HOLD,
    ACTION_LIFT,
    ACTION_RECOVER,
    ActionPack,
    _kg_text,
    _round_kg,
)
from sndintel.capacity import (
    LABEL_FINE,
    cap_shops_per_dsr,
    score_dsr_capacity,
    visit_quality_warnings,
    whale_shops,
)
from sndintel.config import DSR_DAY_CAP, EXPECTED_FORMULA
from sndintel.identity import dsr_display_name

MONDAY_WHALES = 30
MONDAY_DSRS = 15


@dataclass
class OpsPack:
    period: str
    label: str = ""
    kind: str = "monday"
    headline: str = ""
    sheets: list[tuple[str, str, str, pd.DataFrame]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def build_monday_pack(
    action: ActionPack,
    units: pd.DataFrame | None = None,
    visits: pd.DataFrame | None = None,
) -> OpsPack:
    """NSM Monday: where the tons are, who is overloaded vs not converting, which whales."""
    shops = action.raw_shops if action.raw_shops is not None and not action.raw_shops.empty else pd.DataFrame()
    warnings = visit_quality_warnings(units, visits, action.period)
    cap = score_dsr_capacity(shops, action.as_of_day, action.days_in_month, action.days_left) if not shops.empty else pd.DataFrame()
    whales = whale_shops(shops, n=MONDAY_WHALES) if not shops.empty else pd.DataFrame()
    cities = _city_driver_table(units) if units is not None else pd.DataFrame()
    headline = action.headline or f"{action.label}: Monday dispatch"
    work_dsrs = cap[cap["label"] != LABEL_FINE].head(MONDAY_DSRS) if not cap.empty else cap
    sheets = [
        ("01 Country", "Country this week", "Ask rest of month is closable drops, not the whole hole.", action.country),
        (
            "02 Cities",
            "City drivers",
            "Unbilled = conversion. Unvisited = coverage. Drop = order size. Do not send coverage actions into a 100% visit city.",
            cities,
        ),
        (
            "03 DSR labels",
            "Who to push — capacity vs skill",
            "Overloaded = not enough DSRs. Not working = effort. Not converting / not lifting = skill or commercial. Fine = leave them.",
            _present_capacity(work_dsrs if not work_dsrs.empty else cap.head(MONDAY_DSRS)),
        ),
        (
            "04 Whales",
            "Volume doors (AMS / drop ≥ 1 MT)",
            "These close the month. Kiryana seriousness lists do not. Owner is the DSR.",
            _present_whale_ops(whales),
        ),
        (
            "05 Distributors",
            "Distributor ask",
            "Ranked by rest-of-month ask KG.",
            action.distributors,
        ),
    ]
    return OpsPack(period=action.period, label=action.label, kind="monday", headline=headline, sheets=sheets, warnings=warnings)


def build_dsr_beat_pack(action: ActionPack, per_dsr: int | None = None) -> OpsPack:
    """Every DSR’s today-list, capped at what they can actually call."""
    shops = action.raw_shops if action.raw_shops is not None and not action.raw_shops.empty else pd.DataFrame()
    if shops.empty:
        return OpsPack(period=action.period, label=action.label, kind="dsr", headline="No DSR beat list.")
    work = shops[shops["action"].isin({ACTION_CALL, ACTION_CONVERT, ACTION_LIFT, ACTION_RECOVER})].copy()
    coming = shops[shops.get("coming_due").fillna(False)] if "coming_due" in shops.columns else shops.iloc[0:0]
    work = pd.concat([work, coming], ignore_index=True)
    if "store_id" in work.columns:
        work = work.drop_duplicates("store_id")
    work = cap_shops_per_dsr(work, per_dsr=per_dsr or DSR_DAY_CAP)
    if "value_score" in work.columns:
        work = work.sort_values(["dsr_name", "city", "value_score"], ascending=[True, True, False])
    cap = int(per_dsr or DSR_DAY_CAP)
    headline = (
        f"DSR beat lists — {cap} doors each (capacity cap). "
        f"{len(work)} calls on the sheet. Waiting-list doors stay in the detailed action pack."
    )
    sheets = [
        (
            "01 Beat lists",
            "Capacity-capped calls",
            f"Max {cap} doors per DSR, ranked by Ask KG. Owner = DSR. Do this is the instruction.",
            _present_beat(work),
        )
    ]
    return OpsPack(period=action.period, label=action.label, kind="dsr", headline=headline, sheets=sheets)


def build_friday_pack(outcomes: pd.DataFrame, action: ActionPack | None = None) -> OpsPack:
    """Close the loop on the list we printed, not a new ranking."""
    if outcomes is None or outcomes.empty:
        headline = (
            "No closed-loop yet. Score the warehouse twice in the same month "
            "(a later MTD cut) and Friday will compare the previous list to new bills and visits."
        )
        empty = pd.DataFrame(
            [
                {
                    "Listed": 0,
                    "Visited after list": 0,
                    "Billed after list": 0,
                    "Not visited": 0,
                    "Visited · still unbilled": 0,
                    "Ask listed (KG)": 0,
                    "Billed vs ask (KG)": 0,
                }
            ]
        )
        return OpsPack(
            period=action.period if action else "",
            label=action.label if action else "",
            kind="friday",
            headline=headline,
            sheets=[("01 Close", "Friday close", headline, empty)],
        )
    summary = _outcome_summary(outcomes)
    by_action = _outcome_by_action(outcomes)
    misses = outcomes[outcomes["outcome"].isin({"not_visited", "visited_unbilled"})].copy()
    if "ask_mt" in misses.columns:
        misses = misses.sort_values("ask_mt", ascending=False).head(40)
    headline = str(summary.get("headline") or "Friday close")
    sheets = [
        ("01 Close", "Listed → visited → billed", "This is whether the list moved volume, not a new ranking.", summary["table"]),
        ("02 By action", "By Due / Convert / Lift / Lapsing", "Visited-unbilled is commercial. Not-visited is management.", by_action),
        ("03 Still open", "Doors still open from the list", "Highest Ask first.", _present_open(misses)),
    ]
    return OpsPack(
        period=str(outcomes["period"].iloc[0]) if "period" in outcomes.columns else "",
        label=action.label if action else "",
        kind="friday",
        headline=headline,
        sheets=sheets,
    )


def score_closed_loop(
    previous: pd.DataFrame,
    shop_month: pd.DataFrame,
    visits: pd.DataFrame | None,
    period: str,
    listed_at: str | None = None,
) -> pd.DataFrame:
    """Compare a previously listed action_shops snapshot to current billed / visits."""
    if previous is None or previous.empty or not period:
        return pd.DataFrame()
    prev = previous.copy()
    work = prev[prev.get("action", pd.Series(dtype=str)).isin({ACTION_CALL, ACTION_CONVERT, ACTION_LIFT, ACTION_RECOVER})]
    if "Action" in prev.columns and work.empty:
        work = prev[prev["Action"].isin({ACTION_CALL, ACTION_CONVERT, ACTION_LIFT, ACTION_RECOVER})]
    if work.empty:
        work = prev
    billed_now = _period_billed(shop_month, period)
    visits_now = _period_visits(visits, period)
    rows = []
    stamp = listed_at or datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    for _, row in work.iterrows():
        sid = str(row.get("store_id") or "")
        if not sid:
            continue
        billed_listed = float(pd.to_numeric(pd.Series([row.get("billed_mt")]), errors="coerce").fillna(0).iloc[0] or 0)
        visits_listed = float(pd.to_numeric(pd.Series([row.get("visits")]), errors="coerce").fillna(0).iloc[0] or 0)
        billed = float(billed_now.get(sid, 0.0) or 0)
        vis = float(visits_now.get(sid, 0.0) or 0)
        ask = float(pd.to_numeric(pd.Series([row.get("week_target_mt")]), errors="coerce").fillna(0).iloc[0] or 0)
        action = str(row.get("action") or row.get("Action") or "")
        gained = max(0.0, billed - billed_listed)
        if billed > billed_listed + 0.005:
            outcome = "billed"
        elif vis > visits_listed + 0.5:
            outcome = "visited_unbilled"
        else:
            outcome = "not_visited"
        rows.append(
            {
                "period": period,
                "listed_at": stamp,
                "store_id": sid,
                "store_name": row.get("store_name") or row.get("Shop") or sid,
                "city": row.get("city") or row.get("City") or "",
                "distributor": row.get("distributor") or row.get("Distributor") or "",
                "dsr_name": row.get("dsr_name") or row.get("DSR") or "",
                "action": action,
                "ask_mt": ask,
                "billed_mt_listed": billed_listed,
                "visits_listed": visits_listed,
                "billed_mt_now": billed,
                "visits_now": vis,
                "gained_mt": gained,
                "outcome": outcome,
            }
        )
    return pd.DataFrame(rows)


def persist_outcomes(conn, outcomes: pd.DataFrame) -> None:
    from sndintel.storage import replace_table

    conn.execute(
        """CREATE TABLE IF NOT EXISTS action_outcomes (
            period TEXT NOT NULL,
            listed_at TEXT NOT NULL,
            store_id TEXT NOT NULL,
            store_name TEXT,
            city TEXT,
            distributor TEXT,
            dsr_name TEXT,
            action TEXT,
            ask_mt REAL,
            billed_mt_listed REAL,
            visits_listed REAL,
            billed_mt_now REAL,
            visits_now REAL,
            gained_mt REAL,
            outcome TEXT,
            PRIMARY KEY (period, store_id, listed_at)
        )"""
    )
    if outcomes is None or outcomes.empty:
        return
    existing = None
    try:
        from sndintel.storage import read_sql

        existing = read_sql(conn, "SELECT * FROM action_outcomes")
    except Exception:
        existing = None
    if existing is not None and not existing.empty:
        frame = pd.concat([existing, outcomes], ignore_index=True)
        if "listed_at" in frame.columns and "store_id" in frame.columns:
            frame = frame.drop_duplicates(["period", "store_id", "listed_at"], keep="last")
    else:
        frame = outcomes
    replace_table(conn, "action_outcomes", frame)


def load_outcomes(conn, period: str | None = None) -> pd.DataFrame:
    from sndintel.storage import read_sql

    try:
        df = read_sql(conn, "SELECT * FROM action_outcomes")
    except Exception:
        return pd.DataFrame()
    if df is None or df.empty:
        return pd.DataFrame()
    if period:
        part = df[df["period"].astype(str) == str(period)]
        if not part.empty:
            latest = str(part["listed_at"].max()) if "listed_at" in part.columns else None
            if latest:
                return part[part["listed_at"].astype(str) == latest].copy()
            return part.copy()
    return df


def excel_bytes(pack: OpsPack) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    from sndintel.briefing import NAVY, SLATE, _sheet_table

    buf = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "00 Cover"
    ws["A1"] = "SND Intelligence"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=NAVY)
    title = {"monday": "Monday NSM pack", "dsr": "DSR beat lists", "friday": "Friday close"}.get(pack.kind, pack.kind)
    ws["A2"] = f"{title} · {pack.label}"
    ws["A2"].font = Font(name="Calibri", size=18, bold=True, color=NAVY)
    ws["A3"] = pack.headline or ""
    ws["A3"].font = Font(name="Calibri", size=11, italic=True, color=SLATE)
    ws.merge_cells("A3:H3")
    ws["A4"] = EXPECTED_FORMULA
    ws["A4"].font = Font(name="Calibri", size=9, color=SLATE)
    row = 6
    for warning in pack.warnings:
        ws.cell(row, 1, warning)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        ws.cell(row, 1).alignment = Alignment(wrap_text=True)
        row += 1
    for sheet, heading, note, df in pack.sheets:
        _sheet_table(wb, sheet, heading, note, df)
    wb.save(buf)
    return buf.getvalue()


def pdf_bytes(pack: OpsPack) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from sndintel.action_report import _xml

    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=f"SND Intelligence · {pack.kind} · {pack.label}",
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle("t", parent=styles["Heading1"], fontSize=14, textColor=colors.HexColor("#0F172A"), spaceAfter=6)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=11, textColor=colors.HexColor("#0F172A"), spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle("b", parent=styles["Normal"], fontSize=8, textColor=colors.HexColor("#334155"), leading=11)
    kind = {"monday": "Monday NSM pack", "dsr": "DSR beat lists", "friday": "Friday close"}.get(pack.kind, pack.kind)
    story = [
        Paragraph(f"SND Intelligence · {kind}", body),
        Paragraph(_xml(pack.headline or pack.label), title),
        Paragraph(_xml(EXPECTED_FORMULA), body),
    ]
    for warning in pack.warnings:
        story.append(Paragraph(_xml(warning), body))
    for i, (_sheet, heading, note, df) in enumerate(pack.sheets, start=1):
        story.append(Paragraph(f"{i}. {_xml(heading)}", h2))
        story.append(Paragraph(_xml(note), body))
        story.append(Spacer(1, 4))
        story.append(_pdf_table(df, body))
    doc.build(story)
    return buf.getvalue()


def _pdf_table(df: pd.DataFrame, style):
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Table, TableStyle

    from sndintel.action_report import _xml

    if df is None or df.empty:
        return Paragraph("No rows at this layer.", style)
    show = df.head(40)
    header = [Paragraph(f"<b>{_xml(str(c))}</b>", style) for c in show.columns]
    data = [header]
    for _, rec in show.iterrows():
        data.append([Paragraph(_xml("" if rec[c] is None or pd.isna(rec[c]) else str(rec[c])), style) for c in show.columns])
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CBD5E1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    return table


def _city_driver_table(units: pd.DataFrame) -> pd.DataFrame:
    if units is None or units.empty:
        return pd.DataFrame()
    cities = units[units["grain"] == "city"].copy()
    if cities.empty:
        return pd.DataFrame()
    rec = pd.to_numeric(cities.get("isolated_mt"), errors="coerce").fillna(0).clip(upper=0).abs()
    unb = pd.to_numeric(cities.get("from_unbilled_mt"), errors="coerce").fillna(0)
    unv = pd.to_numeric(cities.get("from_unvisited_mt"), errors="coerce").fillna(0)
    drop = pd.to_numeric(cities.get("from_drop_size_mt"), errors="coerce").fillna(0)
    driver = []
    for u, v, d in zip(unb, unv, drop):
        parts = [("unbilled", float(u)), ("unvisited", float(v)), ("drop size", float(d))]
        parts.sort(key=lambda x: abs(x[1]), reverse=True)
        driver.append(parts[0][0] if parts[0][1] else "on expected")
    out = pd.DataFrame(
        {
            "City": cities["grain_id"].astype(str),
            "Billed (MT)": pd.to_numeric(cities.get("volume_mt"), errors="coerce").round(0),
            "Expected (MT)": pd.to_numeric(cities.get("expected_mt"), errors="coerce").round(0),
            "Gap (MT)": rec.round(0),
            "Visit %": (pd.to_numeric(cities.get("visit_rate"), errors="coerce") * 100).round(0),
            "Strike %": (pd.to_numeric(cities.get("strike_rate"), errors="coerce") * 100).round(0),
            "Driver": driver,
        }
    )
    return out.sort_values("Gap (MT)", ascending=False)


def _present_capacity(df: pd.DataFrame) -> pd.DataFrame:
    if df is None or df.empty:
        return pd.DataFrame(
            columns=["DSR", "City", "Distributor", "Label", "Universe", "Visit %", "Strike of visits %", "Span ×", "Ask rest of month (KG)", "Why"]
        )
    return pd.DataFrame(
        {
            "DSR": [dsr_display_name(v) for v in df.get("dsr_name", df.get("grain_id"))],
            "City": list(df.get("city", [])),
            "Distributor": list(df.get("distributor", [])),
            "Label": list(df.get("label", [])),
            "Universe": list(df.get("universe", [])),
            "Visit %": [None if pd.isna(v) else int(round(float(v) * 100)) for v in df.get("visit_rate", [])],
            "Strike of visits %": [None if pd.isna(v) else int(round(float(v) * 100)) for v in df.get("strike_of_visits", [])],
            "Span ×": [None if pd.isna(v) else round(float(v), 1) for v in df.get("span_unique", [])],
            "Ask rest of month (KG)": [_round_kg(v) for v in df.get("week_target_mt", [])],
            "Day cap": list(df.get("day_cap", [])),
            "Why": list(df.get("why", [])),
        }
    )


def _present_whale_ops(df: pd.DataFrame) -> pd.DataFrame:
    if df is None or df.empty:
        return pd.DataFrame(columns=["Shop", "City", "DSR", "Distributor", "Ask rest of month (KG)", "AMS (KG)", "Call", "Do this"])
    name = df["store_name"] if "store_name" in df.columns else df.get("Shop")
    return pd.DataFrame(
        {
            "Shop": list(name),
            "City": list(df.get("city", [])),
            "DSR": list(df.get("dsr_name", [])),
            "Distributor": list(df.get("distributor", [])),
            "Ask rest of month (KG)": [_round_kg(v) for v in df.get("week_target_mt", [])],
            "AMS (KG)": [_round_kg(v) for v in df.get("ams_3m", [])],
            "Billed (KG)": [_round_kg(v) for v in df.get("billed_mt", df.get("volume_mt", []))],
            "Call": list(df.get("call_status", df.get("Call", []))),
            "Do this": list(df.get("instruction", [])),
        }
    )


def _present_beat(df: pd.DataFrame) -> pd.DataFrame:
    if df is None or df.empty:
        return pd.DataFrame()
    out = pd.DataFrame(
        {
            "DSR": list(df.get("dsr_name", [])),
            "City": list(df.get("city", [])),
            "Distributor": list(df.get("distributor", [])),
            "Shop": list(df.get("store_name", [])),
            "Action": list(df.get("action", [])),
            "Ask rest of month (KG)": [_round_kg(v) for v in df.get("week_target_mt", [])],
            "Next order (KG)": [_round_kg(v) for v in df.get("next_drop_mt", [])],
            "Days since bill": [None if pd.isna(v) else int(round(float(v))) for v in df.get("days_since_bill", [])],
            "Cover left (days)": [None if pd.isna(v) else int(round(float(v))) for v in df.get("cover_left_days", [])],
            "Owner": list(df.get("dsr_name", [])),
            "Do this": list(df.get("instruction", [])),
        }
    )
    return out


def _period_billed(shop_month: pd.DataFrame, period: str) -> dict[str, float]:
    if shop_month is None or shop_month.empty:
        return {}
    cur = shop_month[shop_month["period"].astype(str) == str(period)]
    if cur.empty:
        return {}
    g = cur.groupby(cur["store_id"].astype(str))["volume_mt"].sum()
    return {str(k): float(v) for k, v in g.items()}


def _period_visits(visits: pd.DataFrame | None, period: str) -> dict[str, float]:
    if visits is None or visits.empty or "store_id" not in visits.columns:
        return {}
    v = visits.copy()
    if "period" in v.columns:
        v = v[v["period"].astype(str) == str(period)]
    if v.empty:
        return {}
    col = "visits" if "visits" in v.columns else None
    if col is None:
        return {str(s): 1.0 for s in v["store_id"].astype(str)}
    g = v.groupby(v["store_id"].astype(str))[col].sum()
    return {str(k): float(v) for k, v in g.items()}


def _outcome_summary(outcomes: pd.DataFrame) -> dict[str, Any]:
    n = int(len(outcomes))
    billed = int((outcomes["outcome"] == "billed").sum())
    vis_u = int((outcomes["outcome"] == "visited_unbilled").sum())
    skip = int((outcomes["outcome"] == "not_visited").sum())
    ask = float(pd.to_numeric(outcomes.get("ask_mt"), errors="coerce").fillna(0).sum())
    gained = float(pd.to_numeric(outcomes.get("gained_mt"), errors="coerce").fillna(0).sum())
    headline = (
        f"{billed} of {n} listed doors billed after the list "
        f"({_kg_text(gained)} against {_kg_text(ask)} Ask). "
        f"{skip} were not visited — management. {vis_u} were visited and still unbilled — commercial."
    )
    table = pd.DataFrame(
        [
            {
                "Listed": n,
                "Billed after list": billed,
                "Visited · still unbilled": vis_u,
                "Not visited": skip,
                "Ask listed (KG)": _round_kg(ask),
                "Billed vs ask (KG)": _round_kg(gained),
            }
        ]
    )
    return {"headline": headline, "table": table}


def _outcome_by_action(outcomes: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for action, g in outcomes.groupby(outcomes["action"].astype(str)):
        rows.append(
            {
                "Action": action,
                "Listed": int(len(g)),
                "Billed": int((g["outcome"] == "billed").sum()),
                "Visited · unbilled": int((g["outcome"] == "visited_unbilled").sum()),
                "Not visited": int((g["outcome"] == "not_visited").sum()),
                "Ask (KG)": _round_kg(g["ask_mt"].sum()) if "ask_mt" in g.columns else 0,
                "Gained (KG)": _round_kg(g["gained_mt"].sum()) if "gained_mt" in g.columns else 0,
            }
        )
    return pd.DataFrame(rows)


def _present_open(df: pd.DataFrame) -> pd.DataFrame:
    if df is None or df.empty:
        return pd.DataFrame(columns=["Shop", "DSR", "City", "Action", "Outcome", "Ask (KG)"])
    return pd.DataFrame(
        {
            "Shop": list(df.get("store_name", [])),
            "DSR": list(df.get("dsr_name", [])),
            "City": list(df.get("city", [])),
            "Action": list(df.get("action", [])),
            "Outcome": list(df.get("outcome", [])),
            "Ask (KG)": [_round_kg(v) for v in df.get("ask_mt", [])],
        }
    )
