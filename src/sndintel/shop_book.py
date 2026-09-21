"""Shop-wise issues pack: every door in a scope, MTD or a closed month.

Does not invent Ask, last bill, or Target. Shop Expected is the shop's own
robust last-3 / last-6 run-rate (winsorised, shrunk toward its city) and is
printed **unscaled**. The official cascade Expected for the scope is shown
beside it with the reconciliation factor — the two are never blended, so a
shop's miss is judged against its own history, not a city-wide multiplier.

Closed month and open MTD share one skeleton (KPIs → roll-up → issue mix →
shop list) and differ in what counts as an issue:

- Closed month is a result. Cover Billed, Expected, and Gap are this pack’s
  shops: Gap = max(0, Expected − billed). The mix totals to those figures
  (beat shops net against misses). Issues are the doors that missed.
- Open MTD is the beat. Full-month Expected − billed so far is not a miss
  mid-month, so it is not the issue list. Issues are due this week (Ask),
  visited with no bill, and lost doors.

Materiality is Pareto, not a fixed cut: within each DSR the doors that make
up 80% of Expected (or billed) are core; the rest is the tail, reported as a
coverage panel (doors billed vs usual) rather than as individual holes. Any
door at or above 50 kg is always core; under 10 kg is always tail.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape as xml_escape

import pandas as pd
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from sndintel.action import ActionPack, build_action_pack
from sndintel.briefing import _sheet_table
from sndintel.mtd import period_state
from sndintel.plan import attach_plan

HOLE_MT = 0.05  # a door with Expected or billed at/above this is always core
MATERIAL_FLOOR_MT = 0.01  # under this a door is always tail
PARETO_SHARE = 0.80  # within a DSR, the doors making up this share of Expected are core
MISS_TOL = 0.20  # miss / beat when off Expected by more than this share ...
MISS_FLOOR_MT = 0.01  # ... and by at least this much
ASK_MT = 0.0005
BILL_MT = 0.005
PDF_ISSUE_N = 200
MIX_OTHER = "Tail shops"
MIX_TOTAL = "Total"
ISSUE_LAPSED = "Lapsed (lost door)"
ISSUE_MISSED = "Missed Expected"
ISSUE_NOT_DUE = "Not due (cycle timing)"
ISSUE_MIGRATED = "Code migrated (see new POP)"
ISSUE_UNBILLED = "Unbilled"
ISSUE_UNVISITED = "Unvisited"
# Mix billed must be 0 for these. Any rounding dust moves into MIX_OTHER.
NO_BILL_ISSUES = {
    ISSUE_UNBILLED,
    ISSUE_UNVISITED,
    ISSUE_LAPSED,
    ISSUE_MIGRATED,
    "Visited, no bill",
    "Due · no bill",
    "Due · unvisited",
}

CLOSED_ISSUE_RANK = {
    ISSUE_MISSED: 0,
    ISSUE_LAPSED: 1,
    ISSUE_NOT_DUE: 7,
    ISSUE_MIGRATED: 7,
    "Beat Expected": 8,
    "On Expected": 9,
    MIX_OTHER: 10,
    "No run-rate": 11,
}
MTD_ISSUE_RANK = {
    "Due · unvisited": 0,
    "Due · no bill": 1,
    "Due · light drop": 2,
    "Visited, no bill": 3,
    ISSUE_LAPSED: 4,
    ISSUE_UNVISITED: 6,
    ISSUE_MIGRATED: 7,
    "On cycle": 9,
    MIX_OTHER: 10,
    "No run-rate": 11,
}
MIX_SKIP = {"No run-rate", MIX_OTHER}
ISSUE_LAG = {
    "Due · unvisited",
    "Due · no bill",
    "Due · light drop",
    "Visited, no bill",
    ISSUE_MISSED,
    ISSUE_UNBILLED,
    ISSUE_UNVISITED,
    ISSUE_LAPSED,
}


@dataclass
class ShopBook:
    period: str
    label: str
    scope: str = "national"
    scope_label: str = "Country"
    open_mtd: bool = False
    headline: str = ""
    weather: str = ""
    how_to_read: list[str] = field(default_factory=list)
    kpis: dict[str, Any] = field(default_factory=dict)
    mix: pd.DataFrame = field(default_factory=pd.DataFrame)
    issues: pd.DataFrame = field(default_factory=pd.DataFrame)
    shops: pd.DataFrame = field(default_factory=pd.DataFrame)
    raw: pd.DataFrame = field(default_factory=pd.DataFrame)
    rollup: pd.DataFrame = field(default_factory=pd.DataFrame)
    bridge: pd.DataFrame = field(default_factory=pd.DataFrame)
    tail: pd.DataFrame = field(default_factory=pd.DataFrame)
    flags: pd.DataFrame = field(default_factory=pd.DataFrame)


def empty_shop_book(period: str = "") -> ShopBook:
    return ShopBook(period=period or "", label=period or "")


def build_shop_book(
    *,
    action: ActionPack | None = None,
    shop_month: pd.DataFrame | None = None,
    stores: pd.DataFrame | None = None,
    shop_day: pd.DataFrame | None = None,
    visits: pd.DataFrame | None = None,
    ledger: pd.DataFrame | None = None,
    shop_targets: pd.DataFrame | None = None,
    units: pd.DataFrame | None = None,
    period: str | None = None,
    scope: str = "national",
    city: str | None = None,
    distributor: str | None = None,
    dsr: str | None = None,
) -> ShopBook:
    """Classify existing shop rows. Does not recompute Expected or Ask."""
    period = str(period or (getattr(action, "period", None) or "") or "")
    if action is None or getattr(action, "raw_shops", None) is None or action.raw_shops.empty:
        if shop_month is None or shop_month.empty:
            return empty_shop_book(period)
        action = build_action_pack(
            shop_month,
            stores,
            shop_day=shop_day,
            visits=visits,
            ledger=ledger,
            period=period or None,
        )
    period = str(action.period or period or "")
    shops = action.raw_shops.copy()
    if shops.empty:
        book = empty_shop_book(period)
        book.label = action.label or period
        book.open_mtd = bool(action.open_mtd)
        return book

    shops = _attach_shop_targets(shops, shop_targets)
    scope = (scope or "national").strip().lower()
    if scope in {"country", "national hq", ""}:
        scope = "national"
    shops = _scope_shops(shops, scope, city, distributor, dsr)
    open_mtd = bool(action.open_mtd)
    if ledger is not None:
        open_mtd = bool(period_state(ledger, period).get("open"))
    shops = _absorb_missing_billed(shops, shop_month, period, scope, city, distributor, dsr)
    units = _ensure_plan(units, shop_targets)
    row = _scope_unit(units, scope, city, distributor, dsr)
    shops = _attach_flags(shops, shop_month, shop_day, period)
    shops = _classify_rows(shops, open_mtd, period=period)
    shops = _sort_rows(shops, open_mtd)

    scope_label = _scope_label(scope, city, distributor, dsr)
    kpis = _kpis(shops, open_mtd)
    kpis = _apply_cover_kpis(
        kpis,
        units=units,
        shop_targets=shop_targets,
        scope=scope,
        city=city,
        distributor=distributor,
        dsr=dsr,
        open_mtd=open_mtd,
        row=row,
    )
    mix = _issue_mix(shops, open_mtd)
    rollup = _rollup(shops, scope, open_mtd)
    bridge = _door_bridge(shops, shop_month, period, scope)
    tail = _tail_panel(shops, shop_month, period, scope)
    headline, weather = _headline(kpis, scope_label, action.label or period, open_mtd)
    presented = _present(shops, scope, open_mtd, bool(kpis.get("has_plan")), period=period)
    issues = presented[presented["Issue"].isin(ISSUE_LAG)].copy() if not presented.empty else presented
    flags = _present_flags(shops, scope)
    return ShopBook(
        period=period,
        label=action.label or period,
        scope=scope,
        scope_label=scope_label,
        open_mtd=open_mtd,
        headline=headline,
        weather=weather,
        how_to_read=_how_to_read(open_mtd),
        kpis=kpis,
        mix=mix,
        issues=issues,
        shops=presented,
        raw=shops,
        rollup=rollup,
        bridge=bridge,
        tail=tail,
        flags=flags,
    )


def list_shop_book_entities(shops: pd.DataFrame | None, kind: str) -> list[str]:
    kind = (kind or "").strip().lower()
    if shops is None or shops.empty:
        return []
    work = shops.copy()
    if kind == "city":
        return sorted({str(x).strip() for x in work.get("city", pd.Series(dtype=str)).dropna().astype(str) if str(x).strip() not in {"", "nan", "(unmapped)"}})
    if kind == "distributor":
        out = []
        for _, r in work.iterrows():
            name = str(r.get("distributor") or "").strip()
            city = str(r.get("city") or "").strip()
            if not name or name in {"nan", "(unmapped)"}:
                continue
            out.append(f"{city} · {name}" if city and city not in {"nan", "(unmapped)"} else name)
        return sorted(set(out))
    if kind == "dsr":
        out = []
        for _, r in work.iterrows():
            name = str(r.get("dsr_name") or "").strip()
            city = str(r.get("city") or "").strip()
            dist = str(r.get("distributor") or "").strip()
            if not name or name in {"nan", "(unnamed)"}:
                continue
            if city and city not in {"nan", "(unmapped)"} and dist and dist not in {"nan", "(unmapped)"}:
                out.append(f"{city} · {dist} · {name}")
            elif city and city not in {"nan", "(unmapped)"}:
                out.append(f"{city} · {name}")
            else:
                out.append(name)
        return sorted(set(out))
    return []


def parse_shop_scope(kind: str, entity: str | None) -> dict[str, str | None]:
    parts = [p.strip() for p in str(entity or "").split(" · ") if str(p).strip()]
    kind = (kind or "national").strip().lower()
    if kind == "city":
        return {"city": entity, "distributor": None, "dsr": None}
    if kind == "distributor":
        if len(parts) >= 2:
            return {"city": parts[0], "distributor": " · ".join(parts[1:]), "dsr": None}
        return {"city": None, "distributor": entity, "dsr": None}
    if kind == "dsr":
        if len(parts) >= 3:
            return {"city": parts[0], "distributor": parts[1], "dsr": " · ".join(parts[2:])}
        if len(parts) == 2:
            return {"city": parts[0], "distributor": None, "dsr": parts[1]}
        return {"city": None, "distributor": None, "dsr": entity}
    return {"city": None, "distributor": None, "dsr": None}


def _attach_shop_targets(shops: pd.DataFrame, shop_targets: pd.DataFrame | None) -> pd.DataFrame:
    out = shops.copy()
    out["shop_target_mt"] = 0.0
    if shop_targets is None or shop_targets.empty or "store_id" not in shop_targets.columns:
        return out
    work = shop_targets.copy()
    work["store_id"] = work["store_id"].astype(str).str.strip()
    method = work["match_method"].fillna("unmatched").astype(str) if "match_method" in work.columns else pd.Series("id", index=work.index)
    matched = work.loc[method.ne("unmatched") & work["store_id"].ne("") & work["store_id"].ne("nan")]
    if matched.empty or "target_mt" not in matched.columns:
        return out
    tgt = matched.groupby(matched["store_id"].astype(str))["target_mt"].sum()
    out["shop_target_mt"] = out["store_id"].astype(str).map(tgt).fillna(0.0)
    return out


def _scope_shops(
    shops: pd.DataFrame,
    scope: str,
    city: str | None,
    distributor: str | None,
    dsr: str | None,
) -> pd.DataFrame:
    out = shops.copy()
    if scope == "national" or not scope:
        return out
    if city:
        out = out[out["city"].astype(str) == str(city)]
    if scope in {"distributor", "dsr"} and distributor:
        out = out[out["distributor"].astype(str) == str(distributor)]
    if scope == "dsr" and dsr:
        out = out[out["dsr_name"].astype(str) == str(dsr)]
    return out.reset_index(drop=True)


def _geo_filter(
    df: pd.DataFrame,
    scope: str,
    city: str | None,
    distributor: str | None,
    dsr: str | None,
) -> pd.DataFrame:
    out = df
    if city and "city" in out.columns:
        out = out[_series_eq(out["city"], city)]
    if scope in {"distributor", "dsr"} and distributor and "distributor" in out.columns:
        out = out[_series_eq(out["distributor"], distributor)]
    if scope == "dsr" and dsr and "dsr_name" in out.columns:
        out = out[_series_eq(out["dsr_name"], dsr)]
    return out


def _absorb_missing_billed(
    shops: pd.DataFrame,
    shop_month: pd.DataFrame | None,
    period: str,
    scope: str,
    city: str | None,
    distributor: str | None,
    dsr: str | None,
) -> pd.DataFrame:
    """Bring in doors that billed this month but are missing from the action universe."""
    if shops is None:
        shops = pd.DataFrame()
    if shop_month is None or shop_month.empty or not period:
        return shops
    cur = shop_month.copy()
    cur["store_id"] = cur["store_id"].astype(str)
    cur = cur[cur["period"].astype(str) == str(period)]
    cur = _geo_filter(cur, scope, city, distributor, dsr)
    if cur.empty:
        return shops
    billed = cur.groupby("store_id")["volume_mt"].sum()
    billed = billed[pd.to_numeric(billed, errors="coerce").fillna(0) > ASK_MT]
    have = set(shops["store_id"].astype(str)) if not shops.empty and "store_id" in shops.columns else set()
    missing = [sid for sid in billed.index.astype(str) if sid not in have]
    if not missing:
        return shops
    attrs = cur.sort_values("period").drop_duplicates("store_id", keep="last")
    attr_map = attrs.set_index("store_id") if not attrs.empty else pd.DataFrame()
    rows = []
    for sid in missing:
        rec = {c: (0.0 if c.endswith("_mt") else "") for c in shops.columns} if not shops.empty else {}
        rec["store_id"] = sid
        rec["billed_mt"] = float(billed.loc[sid])
        rec["expected_mt"] = 0.0
        rec["remaining_mt"] = 0.0
        rec["week_target_mt"] = 0.0
        rec["shop_target_mt"] = 0.0
        rec["is_lapsed"] = False
        rec["call_status"] = "Billed"
        if sid in attr_map.index:
            hit = attr_map.loc[sid]
            if isinstance(hit, pd.DataFrame):
                hit = hit.iloc[0]
            for col in ("store_name", "city", "distributor", "dsr_name", "section"):
                if col in hit.index:
                    rec[col] = str(hit.get(col) or "")
        rec.setdefault("store_name", sid)
        rows.append(rec)
    extra = pd.DataFrame(rows)
    if shops.empty:
        return extra
    return pd.concat([shops, extra], ignore_index=True, sort=False)


def _official_expected(row: pd.Series | None, open_mtd: bool) -> float | None:
    if row is None:
        return None
    today = _f(row.get("expected_mt"))
    if today <= 1e-9:
        return None
    if not open_mtd:
        return today
    pace = _f(row.get("intra_month_frac"), 1.0) or 1.0
    full = today / pace if pace > 1e-6 else today
    return full if full > 1e-9 else None


def _attach_flags(
    shops: pd.DataFrame,
    shop_month: pd.DataFrame | None,
    shop_day: pd.DataFrame | None,
    period: str,
) -> pd.DataFrame:
    """Duplicate-code / migrated-POP flags from the dupes module (never changes volume)."""
    out = shops.copy()
    for col in ("flag", "flag_detail", "flag_pair"):
        if col not in out.columns:
            out[col] = ""
    if out.empty:
        return out
    try:
        from sndintel.dupes import flag_shops

        flags = flag_shops(shop_month, shop_day, period, scope_ids=set(out["store_id"].astype(str)))
    except Exception:
        flags = pd.DataFrame()
    if flags is None or flags.empty:
        return out
    flags = flags.drop_duplicates("store_id").set_index("store_id")
    ids = out["store_id"].astype(str)
    for col in ("flag", "flag_detail", "flag_pair"):
        if col in flags.columns:
            out[col] = ids.map(flags[col]).fillna("").astype(str)
    return out


def _dsr_key(frame: pd.DataFrame) -> pd.Series:
    parts = []
    for col in ("city", "distributor", "dsr_name"):
        if col in frame.columns:
            parts.append(frame[col].fillna("").astype(str).str.strip().str.casefold())
        else:
            parts.append(pd.Series("", index=frame.index))
    return parts[0] + " | " + parts[1] + " | " + parts[2]


def flag_material(shops: pd.DataFrame) -> pd.Series:
    """Core vs tail doors.

    Size is ``max(Expected, billed)``: a door that billed 2 MT on a 30 kg
    run-rate is material even though its history is not. A door is core when
    its size is at least ``HOLE_MT``, or when it sits inside the top
    ``PARETO_SHARE`` of its DSR's size (so a DSR made of small doors still has
    a core). Under ``MATERIAL_FLOOR_MT`` is always tail.
    """
    if shops is None or shops.empty:
        return pd.Series(dtype=bool)
    expected = _num_col(shops, "expected_mt")
    billed = _num_col(shops, "billed_mt")
    size = pd.concat([expected, billed], axis=1).max(axis=1)
    material = size >= HOLE_MT
    small = size < MATERIAL_FLOOR_MT
    key = _dsr_key(shops)
    order = size.sort_values(ascending=False, kind="mergesort")
    cum = order.groupby(key.reindex(order.index)).cumsum()
    total = size.groupby(key).transform("sum").reindex(order.index)
    share_before = (cum - order) / total.replace(0, float("nan"))
    in_core = (share_before.fillna(1.0) < PARETO_SHARE).reindex(shops.index).fillna(False)
    return (material | (in_core & ~small)) & ~small


def _scope_label(scope: str, city: str | None, distributor: str | None, dsr: str | None) -> str:
    if scope == "dsr":
        bits = [b for b in (city, distributor, dsr) if b]
        return " · ".join(bits) if bits else "DSR"
    if scope == "distributor":
        return f"{city} · {distributor}" if city and distributor else (distributor or city or "Distributor")
    if scope == "city":
        return city or "City"
    return "Country"


def _not_due_flags(shops: pd.DataFrame) -> pd.Series:
    """Closed month: the shop's measured cycle did not fall due inside the month.

    Only a *measured* cycle counts (at least one observed interval). A door
    whose last bill is fresher than its cycle at month-end was simply not due;
    its run-rate Expected is timing, not a miss.
    """
    if shops is None or shops.empty:
        return pd.Series(dtype=bool)
    needed = ("days_since_bill", "api_days", "n_intervals", "billed_mt")
    if any(c not in shops.columns for c in needed):
        return pd.Series(False, index=shops.index)
    dslp = pd.to_numeric(shops["days_since_bill"], errors="coerce")
    api = pd.to_numeric(shops["api_days"], errors="coerce")
    n_int = pd.to_numeric(shops["n_intervals"], errors="coerce").fillna(0)
    billed = pd.to_numeric(shops["billed_mt"], errors="coerce").fillna(0.0)
    measured = n_int >= 1
    return (billed <= ASK_MT) & measured & dslp.notna() & api.notna() & (dslp < api)


def _classify_rows(shops: pd.DataFrame, open_mtd: bool, period: str = "") -> pd.DataFrame:
    out = shops.copy()
    billed = pd.to_numeric(out.get("billed_mt"), errors="coerce").fillna(0.0)
    expected = pd.to_numeric(out.get("expected_mt"), errors="coerce").fillna(0.0)
    remaining = (expected - billed).clip(lower=0)
    ask = pd.to_numeric(out.get("week_target_mt"), errors="coerce").fillna(0.0)
    lapsed = out["is_lapsed"].fillna(False).astype(bool) if "is_lapsed" in out.columns else pd.Series(False, index=out.index)
    # A door that billed inside the month being scored is never a lost door.
    lapsed = lapsed & (billed <= ASK_MT)
    call = out["call_status"].astype(str) if "call_status" in out.columns else pd.Series("", index=out.index)
    unvisited = call.eq("Unvisited")
    visited_no_bill = call.eq("Visited · not billed")
    material = flag_material(out).reindex(out.index).fillna(False).astype(bool)
    not_due = _not_due_flags(out).reindex(out.index).fillna(False).astype(bool) if not open_mtd else pd.Series(False, index=out.index)
    issues = []
    flags = []
    comments = []
    for i, row in out.iterrows():
        rec = row.to_dict()
        rec["billed_mt"] = float(billed.loc[i])
        rec["expected_mt"] = float(expected.loc[i])
        rec["remaining_mt"] = float(remaining.loc[i])
        rec["week_target_mt"] = float(ask.loc[i])
        rec["is_lapsed"] = bool(lapsed.loc[i])
        rec["call_status"] = str(call.loc[i])
        rec["is_material"] = bool(material.loc[i])
        rec["not_due"] = bool(not_due.loc[i])
        issue = _issue(rec, open_mtd, bool(unvisited.loc[i]), bool(visited_no_bill.loc[i]))
        issues.append(issue)
        flags.append(issue in ISSUE_LAG)
        comments.append(_comment(rec, issue, open_mtd, period=period))
    out["remaining_mt"] = remaining
    out["is_lapsed"] = lapsed
    out["is_material"] = material
    out["issue"] = issues
    out["is_issue"] = flags
    out["comment"] = comments
    rank_map = MTD_ISSUE_RANK if open_mtd else CLOSED_ISSUE_RANK
    out["issue_rank"] = out["issue"].map(lambda x: rank_map.get(str(x), 7))
    return out


def miss_tolerance_mt(expected: float) -> float:
    """A shop is off Expected when the difference exceeds 20% and at least 10 kg."""
    return max(MISS_FLOOR_MT, MISS_TOL * max(0.0, float(expected or 0)))


def _issue(row: dict[str, Any], open_mtd: bool, unvisited: bool, visited_no_bill: bool) -> str:
    billed = float(row.get("billed_mt") or 0)
    expected = float(row.get("expected_mt") or 0)
    ask = float(row.get("week_target_mt") or 0)
    lapsed = bool(row.get("is_lapsed"))
    material = bool(row.get("is_material", True))
    has_bill = billed > ASK_MT
    migrated_old = str(row.get("flag") or "") == "Code migrated (old)"
    if migrated_old and not has_bill:
        return ISSUE_MIGRATED
    if open_mtd:
        if not material and not (ask > ASK_MT):
            if billed <= ASK_MT and expected <= ASK_MT:
                return "No run-rate"
            return MIX_OTHER
        if lapsed and not has_bill:
            return ISSUE_LAPSED
        if ask > ASK_MT and unvisited:
            return "Due · unvisited"
        if ask > ASK_MT and visited_no_bill:
            return "Due · no bill"
        if ask > ASK_MT:
            return "Due · light drop"
        if visited_no_bill or (not has_bill and not unvisited):
            return "Visited, no bill"
        if not material and not has_bill:
            return "No run-rate"
        return "On cycle"
    # Closed month. Tail doors are a coverage question, not individual holes.
    if not material:
        if billed <= ASK_MT and expected <= ASK_MT:
            return "No run-rate"
        return MIX_OTHER
    if not has_bill and lapsed:
        return ISSUE_LAPSED
    if not has_bill and bool(row.get("not_due")):
        return ISSUE_NOT_DUE
    tol = miss_tolerance_mt(expected)
    if expected - billed > tol:
        return ISSUE_MISSED
    if billed - expected > tol:
        return "Beat Expected"
    return "On Expected"


def fmt_mt(value: Any, *, zero: str = "0 kg") -> str:
    """Kilograms under 0.1 MT, otherwise MT to two decimals. Never '0.00 MT' for 4 kg."""
    v = _f(value)
    if abs(v) < 0.0005:
        return zero
    if abs(v) < 0.1:
        return f"{v * 1000:,.0f} kg"
    return f"{v:,.2f} MT"


def kg(value: Any) -> int:
    return int(round(_f(value) * 1000))


def _cutoff_text(row: dict[str, Any]) -> str:
    from sndintel.demand import lapse_cutoff_days

    cycle = row.get("cycle_days")
    measured = cycle is not None and not pd.isna(cycle) and float(cycle) > 0
    days = int(round(lapse_cutoff_days(float(cycle) if measured else None, measured=measured)))
    if measured:
        return f"quiet past the {days}-day cut-off (3× its {int(round(float(cycle)))}-day cycle, never under 45)"
    return f"quiet past the {days}-day cut-off for a door with no measured cycle"


def _comment(row: dict[str, Any], issue: str, open_mtd: bool, period: str = "") -> str:
    if open_mtd:
        text = str(row.get("instruction") or "").strip()
        if text:
            if "usual drop" not in text.lower() and "every" not in text.lower():
                extra = _usual_cycle_text(row)
                if extra:
                    text = text.rstrip(".") + "." + extra
            return text
        rec = str(row.get("recommended_action") or "").strip()
        name = str(row.get("store_name") or row.get("store_id") or "Shop")
        base = rec if rec else f"{name} is inside its cycle. Leave it."
        return (base.rstrip(".") + "." + _usual_cycle_text(row)).strip()
    name = str(row.get("store_name") or row.get("store_id") or "Shop")
    billed = float(row.get("billed_mt") or 0)
    expected = float(row.get("expected_mt") or 0)
    remaining = float(row.get("remaining_mt") or 0)
    last = _date_text(row.get("last_bill_date"), period)
    last_drop = float(row.get("last_drop_mt") or 0)
    last_bit = last or "no billed sale"
    drop_bit = fmt_mt(last_drop) if last_drop >= BILL_MT else "—"
    cycle_bit = _usual_cycle_text(row)
    call = str(row.get("call_status") or "")
    unvisited = call == "Unvisited"
    has_bill = billed > ASK_MT
    flag = str(row.get("flag") or "").strip()
    flag_bit = f" Flag: {row.get('flag_detail') or flag}." if flag else ""
    if issue in {ISSUE_UNVISITED, ISSUE_UNBILLED} or (issue == ISSUE_MISSED and not has_bill):
        if unvisited:
            return (
                f"{name} was not visited. Expected {fmt_mt(expected)}, billed 0. "
                f"Last billed {last_bit}.{cycle_bit} "
                f"This is Missed Expected (a live shop short this month), not a lost door.{flag_bit} "
                f"Next month: put on the beat before the first drop is due."
            )
        return (
            f"{name} was visited but billed 0 versus Expected {fmt_mt(expected)}. "
            f"Last billed {last_bit} ({drop_bit}).{cycle_bit} "
            f"This is Missed Expected (a live shop short this month), not a lost door.{flag_bit} "
            f"Next month: do not leave without the usual drop."
        )
    if issue == ISSUE_LAPSED:
        return (
            f"{name} is a lost door: last billed {last_bit}, {_cutoff_text(row)}. "
            f"Billed 0 because they stopped buying — not Unbilled and not Missed Expected.{cycle_bit}{flag_bit} "
            f"Next month: recover or drop from the beat."
        )
    if issue == ISSUE_MIGRATED:
        pair = str(row.get("flag_pair") or "").strip()
        return (
            f"{name} stopped billing under this POP code; the same door continues as {pair or 'a new code'}. "
            f"Not a lost door. Fix the universe (retire this code) so its Expected {fmt_mt(expected)} stops showing as a hole."
        )
    if issue == ISSUE_NOT_DUE:
        return (
            f"{name} was not due this month: last billed {last_bit}, inside its usual cycle at month-end.{cycle_bit} "
            f"Run-rate Expected {fmt_mt(expected)} is timing, not a miss. Next month: it falls due — take the usual drop."
        )
    if issue == ISSUE_MISSED:
        return (
            f"{name} billed {fmt_mt(billed)} versus Expected {fmt_mt(expected)}. "
            f"Last drop {drop_bit} on {last_bit}.{cycle_bit}{flag_bit} "
            f"Next month: recover the {fmt_mt(remaining)} hole."
        )
    if issue == "Beat Expected":
        return (
            f"{name} billed {fmt_mt(billed)} versus Expected {fmt_mt(expected)}.{cycle_bit}{flag_bit} "
            f"Keep the same drop next month."
        )
    if issue == "No run-rate":
        return f"{name} has no material Expected this month. Last billed {last_bit}.{cycle_bit}"
    if issue == MIX_OTHER:
        return (
            f"{name} is a tail door for its DSR (Expected {fmt_mt(expected)}, billed {fmt_mt(billed)}). "
            f"Judged in the tail coverage panel, not as an individual hole. Last billed {last_bit}.{cycle_bit}{flag_bit}"
        )
    return f"{name} landed on Expected ({fmt_mt(billed)}). Last billed {last_bit}.{cycle_bit}{flag_bit}"


def _usual_cycle_text(row: dict[str, Any]) -> str:
    usual = float(row.get("expected_drop_mt") or row.get("typical_drop_mt") or 0)
    cycle_n = None
    cycle = row.get("cycle_days")
    try:
        if cycle is not None and not pd.isna(cycle):
            n = int(round(float(cycle)))
            if n > 0:
                cycle_n = n
    except (TypeError, ValueError):
        cycle_n = None
    if usual >= BILL_MT and cycle_n:
        return f" Usual drop is {fmt_mt(usual)} every {cycle_n} days."
    if usual >= BILL_MT:
        return f" Usual drop is {fmt_mt(usual)}."
    return ""


def _sort_rows(shops: pd.DataFrame, open_mtd: bool) -> pd.DataFrame:
    out = shops.copy()
    if open_mtd:
        out["_ask"] = pd.to_numeric(out.get("week_target_mt"), errors="coerce").fillna(0.0)
        out = out.sort_values(
            ["is_issue", "issue_rank", "_ask", "remaining_mt"],
            ascending=[False, True, False, False],
        )
        return out.drop(columns=["_ask"]).reset_index(drop=True)
    out = out.sort_values(
        ["is_issue", "remaining_mt", "issue_rank"],
        ascending=[False, False, True],
    )
    return out.reset_index(drop=True)


def _kpis(shops: pd.DataFrame, open_mtd: bool) -> dict[str, Any]:
    billed = float(pd.to_numeric(shops.get("billed_mt"), errors="coerce").fillna(0).sum()) if not shops.empty else 0.0
    expected = float(pd.to_numeric(shops.get("expected_mt"), errors="coerce").fillna(0).sum()) if not shops.empty else 0.0
    remaining = float(pd.to_numeric(shops.get("remaining_mt"), errors="coerce").fillna(0).sum()) if not shops.empty else 0.0
    matched_target = float(pd.to_numeric(shops.get("shop_target_mt"), errors="coerce").fillna(0).sum()) if not shops.empty else 0.0
    ask = float(pd.to_numeric(shops.get("week_target_mt"), errors="coerce").fillna(0).sum()) if not shops.empty else 0.0
    issues = shops[shops["is_issue"]] if not shops.empty and "is_issue" in shops.columns else shops.iloc[0:0]
    issue_mt = float(pd.to_numeric(issues.get("remaining_mt"), errors="coerce").fillna(0).sum()) if not issues.empty else 0.0
    quiet_n = int((shops["issue"] == "No run-rate").sum()) if not shops.empty and "issue" in shops.columns else 0
    due_n = int((pd.to_numeric(shops.get("week_target_mt"), errors="coerce").fillna(0) > ASK_MT).sum()) if open_mtd and not shops.empty else 0
    tail_n = int((shops["issue"] == MIX_OTHER).sum()) if not shops.empty and "issue" in shops.columns else 0
    not_due_n = int((shops["issue"] == ISSUE_NOT_DUE).sum()) if not shops.empty and "issue" in shops.columns else 0
    flagged_n = int(shops["flag"].astype(str).str.strip().ne("").sum()) if not shops.empty and "flag" in shops.columns else 0
    return {
        "n_tail": tail_n,
        "n_not_due": not_due_n,
        "n_flagged": flagged_n,
        "open_mtd": open_mtd,
        "billed_mt": billed,
        "expected_mt": expected,
        "gap_mt": max(0.0, expected - billed),
        "target_mt": matched_target,
        "matched_target_mt": matched_target,
        "has_plan": matched_target > HOLE_MT,
        "vs_target_mt": billed - matched_target if matched_target > HOLE_MT else 0.0,
        "ask_mt": ask,
        "n_shops": int(len(shops)),
        "n_issues": int(len(issues)),
        "issue_mt": issue_mt,
        "shop_remaining_mt": remaining,
        "n_quiet": quiet_n,
        "n_due": due_n,
        "n_unvisited": int((shops["issue"] == ISSUE_UNVISITED).sum()) if not shops.empty else 0,
        "n_unbilled": int(
            (
                shops["issue"].isin([ISSUE_UNBILLED, "Visited, no bill", "Due · no bill"])
                | (
                    shops["issue"].eq(ISSUE_MISSED)
                    & (pd.to_numeric(shops.get("billed_mt"), errors="coerce").fillna(0).round(2) <= 0)
                )
            ).sum()
        )
        if not shops.empty
        else 0,
        "cover_from_scorecard": False,
    }


def _num_col(frame: pd.DataFrame, name: str) -> pd.Series:
    """Numeric column, zeros when the column is missing."""
    if frame is None or name not in frame.columns:
        return pd.Series(0.0, index=frame.index if frame is not None else None, dtype=float)
    return pd.to_numeric(frame[name], errors="coerce").fillna(0.0)


def _mt2(value: Any) -> float:
    try:
        return round(float(value or 0), 2)
    except (TypeError, ValueError):
        return 0.0


def _f(value: Any, default: float = 0.0) -> float:
    try:
        if value is None:
            return default
        if pd.isna(value):
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


def _fold_key(value: Any) -> str:
    return str(value or "").strip().casefold()


def _series_eq(series: pd.Series, value: str | None) -> pd.Series:
    if value is None or not str(value).strip():
        return pd.Series(True, index=series.index)
    return series.fillna("").astype(str).map(_fold_key).eq(_fold_key(value))


def _scope_unit(
    units: pd.DataFrame | None,
    scope: str,
    city: str | None,
    distributor: str | None,
    dsr: str | None,
) -> pd.Series | None:
    if units is None or units.empty or "grain" not in units.columns:
        return None
    g = units["grain"].astype(str)
    if scope == "national":
        hit = units[g == "national"]
        return None if hit.empty else hit.iloc[0]
    if scope == "city" and city:
        grain_id = units["grain_id"] if "grain_id" in units.columns else pd.Series("", index=units.index)
        hit = units[(g == "city") & _series_eq(grain_id, city)]
        return None if hit.empty else hit.iloc[0]
    if scope == "distributor" and distributor:
        hit = units[g == "distributor"].copy()
        if hit.empty:
            return None
        if city and "city" in hit.columns:
            hit = hit[_series_eq(hit["city"], city)]
        name = hit["grain_id"] if "grain_id" in hit.columns else pd.Series("", index=hit.index)
        alt = hit["distributor"] if "distributor" in hit.columns else pd.Series("", index=hit.index)
        hit = hit[_series_eq(name, distributor) | _series_eq(alt, distributor)]
        return None if hit.empty else hit.iloc[0]
    if scope == "dsr" and dsr:
        hit = units[g == "dsr"].copy()
        if hit.empty:
            return None
        if city and "city" in hit.columns:
            hit = hit[_series_eq(hit["city"], city)]
        if distributor and "distributor" in hit.columns:
            hit = hit[_series_eq(hit["distributor"], distributor)]
        name = hit["dsr_name"] if "dsr_name" in hit.columns else pd.Series("", index=hit.index)
        grain_id = hit["grain_id"].astype(str) if "grain_id" in hit.columns else pd.Series("", index=hit.index)
        want = str(dsr).strip()
        hit = hit[_series_eq(name, dsr) | grain_id.str.startswith(want + " |") | _series_eq(grain_id, dsr)]
        return None if hit.empty else hit.iloc[0]
    return None


def _book_target(
    shop_targets: pd.DataFrame | None,
    scope: str,
    city: str | None,
    distributor: str | None,
    dsr: str | None,
) -> float:
    """Full plan book for this geography, including unmatched names."""
    if shop_targets is None or shop_targets.empty or "target_mt" not in shop_targets.columns:
        return 0.0
    work = shop_targets.copy()
    if scope in {"city", "distributor", "dsr"} and city and "city" in work.columns:
        work = work[_series_eq(work["city"], city)]
    if scope in {"distributor", "dsr"} and distributor and "distributor" in work.columns:
        work = work[_series_eq(work["distributor"], distributor)]
    if scope == "dsr" and dsr and "dsr_name" in work.columns:
        work = work[_series_eq(work["dsr_name"], dsr)]
    return float(pd.to_numeric(work["target_mt"], errors="coerce").fillna(0).sum())


def _ensure_plan(units: pd.DataFrame | None, shop_targets: pd.DataFrame | None) -> pd.DataFrame | None:
    if units is None or units.empty or shop_targets is None or shop_targets.empty:
        return units
    pace = 1.0
    if "intra_month_frac" in units.columns:
        pace = float(pd.to_numeric(units["intra_month_frac"], errors="coerce").dropna().max() or 1.0) or 1.0
    return attach_plan(units, shop_targets, pace=pace)


def _apply_cover_kpis(
    kpis: dict[str, Any],
    *,
    units: pd.DataFrame | None,
    shop_targets: pd.DataFrame | None,
    scope: str,
    city: str | None,
    distributor: str | None,
    dsr: str | None,
    open_mtd: bool,
    row: pd.Series | None = None,
) -> dict[str, Any]:
    """Cover Billed / Expected / Gap are this pack’s shops. Target is the plan book."""
    billed = float(kpis.get("billed_mt") or 0)
    expected = float(kpis.get("expected_mt") or 0)
    kpis["shop_billed_mt"] = billed
    kpis["shop_expected_mt"] = expected
    if row is None:
        units = _ensure_plan(units, shop_targets)
        row = _scope_unit(units, scope, city, distributor, dsr)
    kpis["official_expected_mt"] = None
    kpis["expected_factor"] = None
    if row is not None:
        kpis["scorecard_billed_mt"] = _f(row.get("volume_mt"))
        official = _official_expected(row, open_mtd) or _f(row.get("expected_mt"))
        kpis["scorecard_expected_mt"] = official
        kpis["official_expected_mt"] = official
        kpis["cover_from_scorecard"] = abs(float(official or 0) - expected) <= HOLE_MT
        if official and official > 1e-9 and expected > 1e-9:
            kpis["expected_factor"] = official / expected
        target = _f(row.get("target_mt"))
        if target > HOLE_MT:
            kpis["target_mt"] = target
            kpis["has_plan"] = True
    kpis["billed_mt"] = billed
    kpis["expected_mt"] = expected
    kpis["expected_full_mt"] = expected
    kpis["gap_mt"] = max(0.0, expected - billed)
    if float(kpis.get("target_mt") or 0) <= HOLE_MT:
        book_t = _book_target(shop_targets, scope, city, distributor, dsr)
        if book_t > HOLE_MT:
            kpis["target_mt"] = book_t
            kpis["has_plan"] = True
    target = float(kpis.get("target_mt") or 0)
    if open_mtd and row is not None:
        pace = _f(row.get("intra_month_frac"), 1.0) or 1.0
        projected = billed / pace if pace > 1e-6 else billed
        kpis["projected_mt"] = projected
        if target > HOLE_MT:
            kpis["vs_target_mt"] = projected - target
    elif target > HOLE_MT:
        kpis["vs_target_mt"] = billed - target
    extra = float(kpis.get("billed_mt") or 0) - float(kpis.get("expected_mt") or 0)
    kpis["beat_extra_mt"] = max(0.0, extra)
    kpis.update(target_credibility(expected, target, float(kpis.get("gap_mt") or 0)))
    return kpis


def target_credibility(expected: float, target: float, execution_gap: float) -> dict[str, Any]:
    """Split a shortfall against Target into execution (vs run-rate) and ambition (Target above run-rate).

    Execution gap is what the beat can recover — Expected minus billed. Ambition
    is the part of the Target that sits above the run-rate; no visit plan closes
    it, only growth does. When Target is under the run-rate the plan undersells.
    """
    expected = float(expected or 0)
    target = float(target or 0)
    if target <= HOLE_MT:
        return {"ambition_mt": 0.0, "execution_gap_mt": float(execution_gap or 0), "target_credibility": ""}
    ambition = target - expected
    if ambition > HOLE_MT:
        text = (
            f"Target {fmt_mt(target)} sits {fmt_mt(ambition)} above the run-rate Expected {fmt_mt(expected)}: "
            f"{fmt_mt(execution_gap)} of the shortfall is execution (recoverable on the beat), "
            f"{fmt_mt(ambition)} is ambition (needs growth, not visits)."
        )
    elif ambition < -HOLE_MT:
        text = (
            f"Target {fmt_mt(target)} is {fmt_mt(-ambition)} under the run-rate Expected {fmt_mt(expected)} — "
            f"the plan undersells this scope; beating Target is not beating the run-rate."
        )
    else:
        text = f"Target {fmt_mt(target)} matches the run-rate Expected; the shortfall is execution."
    return {"ambition_mt": ambition, "execution_gap_mt": float(execution_gap or 0), "target_credibility": text}


def _issue_mix(shops: pd.DataFrame, open_mtd: bool) -> pd.DataFrame:
    if shops is None or shops.empty:
        return pd.DataFrame()
    order = list(MTD_ISSUE_RANK) if open_mtd else list(CLOSED_ISSUE_RANK)
    rows = []
    expected_all = pd.to_numeric(shops.get("expected_mt"), errors="coerce").fillna(0.0)
    billed_all = pd.to_numeric(shops.get("billed_mt"), errors="coerce").fillna(0.0)
    if "is_material" in shops.columns:
        tiny_mask = ~shops["is_material"].fillna(False).astype(bool)
    else:
        tiny_mask = expected_all <= HOLE_MT
    quiet = shops["issue"].astype(str).eq("No run-rate") if "issue" in shops.columns else pd.Series(False, index=shops.index)

    def rec_for(label: str, part: pd.DataFrame, billed: float | None = None, expected: float | None = None) -> dict[str, Any]:
        raw = pd.to_numeric(part["billed_mt"], errors="coerce").fillna(0.0) if billed is None else None
        if billed is None:
            shown = raw.map(lambda x: round(float(x), 2))
            if label in {MIX_OTHER, MIX_TOTAL}:
                billed = float(raw.sum())
            else:
                billed = float(raw.where(shown > 0, 0.0).sum())
        if expected is None:
            expected = float(pd.to_numeric(part["expected_mt"], errors="coerce").fillna(0).sum())
        if label in NO_BILL_ISSUES:
            billed = 0.0
        rec = {
            "Issue": label,
            "Shops": int(len(part)),
            "Billed (MT)": _mt2(billed),
            "Expected (MT)": _mt2(expected),
            "Gap (MT)": _mt2(expected - billed),
            "_raw_billed": billed,
            "_raw_expected": expected,
        }
        if open_mtd:
            rec["Ask (KG)"] = int(round(float(pd.to_numeric(part["week_target_mt"], errors="coerce").fillna(0).sum()) * 1000))
        return rec

    for issue in order:
        if issue in MIX_SKIP:
            continue
        part = shops.loc[(shops["issue"] == issue) & ~tiny_mask]
        if part.empty:
            continue
        rows.append(rec_for(issue, part))
    tiny = shops.loc[tiny_mask & ~quiet]
    if not tiny.empty:
        tiny = tiny.loc[(billed_all.reindex(tiny.index).fillna(0) > ASK_MT) | (expected_all.reindex(tiny.index).fillna(0) > ASK_MT)]
    if tiny is not None and not tiny.empty:
        rows.append(rec_for(MIX_OTHER, tiny))
    named_idx = shops.index.difference(tiny.index if tiny is not None and not tiny.empty else shops.iloc[0:0].index)
    named_idx = named_idx.difference(shops.loc[quiet].index)
    dust = 0.0
    if len(named_idx):
        raw = billed_all.reindex(named_idx).fillna(0.0)
        shown = raw.map(lambda x: round(float(x), 2))
        dust = float(raw.where(shown <= 0, 0.0).sum())
    if _mt2(dust) > 0:
        other_i = next((i for i, rec in enumerate(rows) if rec["Issue"] == MIX_OTHER), None)
        if other_i is None:
            empty = shops.iloc[0:0]
            rows.append(rec_for(MIX_OTHER, empty, billed=0.0, expected=0.0))
            other_i = len(rows) - 1
        billed_other = float(rows[other_i]["_raw_billed"]) + dust
        expected_other = float(rows[other_i]["_raw_expected"])
        rows[other_i]["_raw_billed"] = billed_other
        rows[other_i]["Billed (MT)"] = _mt2(billed_other)
        rows[other_i]["Gap (MT)"] = _mt2(expected_other - billed_other)
    if rows:
        rows.append(rec_for(MIX_TOTAL, shops))
    mix = pd.DataFrame(rows)
    return mix.drop(columns=[c for c in mix.columns if str(c).startswith("_")], errors="ignore")


def _headline(kpis: dict[str, Any], scope_label: str, label: str, open_mtd: bool) -> tuple[str, str]:
    billed = float(kpis.get("billed_mt") or 0)
    expected = float(kpis.get("expected_mt") or 0)
    n_iss = int(kpis.get("n_issues") or 0)
    issue_mt = float(kpis.get("issue_mt") or 0)
    n_shops = int(kpis.get("n_shops") or 0)
    n_quiet = int(kpis.get("n_quiet") or 0)
    n_active = max(0, n_shops - n_quiet)
    if open_mtd:
        n_due = int(kpis.get("n_due") or 0)
        ask_kg = int(round(float(kpis.get("ask_mt") or 0) * 1000))
        headline = (
            f"{scope_label} · {label}: billed {_mt2(billed):.2f} MT so far versus Expected {_mt2(expected):.2f} MT. "
            f"{n_due} shops due this week (Ask {ask_kg:,} KG)."
        )
        weather = (
            "Issues are due this week, visited with no bill, and lost doors. "
            "Still-to-Expected is the full month, not a miss yet. Mix Billed / Expected add to the cover."
        )
        return headline, weather
    hole = float(kpis.get("gap_mt") or 0)
    headline = (
        f"{scope_label} closed {label}: billed {fmt_mt(billed)} versus Expected {fmt_mt(expected)} "
        f"(gap {fmt_mt(hole)})."
    )
    official = kpis.get("official_expected_mt")
    factor = kpis.get("expected_factor")
    recon = ""
    if official is not None and factor is not None and abs(float(factor) - 1.0) > 0.005:
        recon = (
            f" Situation cascade Expected for this scope is {fmt_mt(official)} "
            f"(shop-level sum × {float(factor):.2f}); shops are judged on their own run-rate, not on that factor."
        )
    n_tail = int(kpis.get("n_tail") or 0)
    n_not_due = int(kpis.get("n_not_due") or 0)
    weather = (
        f"Gap vs Expected is Expected − billed = {fmt_mt(hole)}. "
        f"{n_iss} shops missed Expected or are lost doors ({fmt_mt(issue_mt)} of holes); shops that beat Expected net against that in the mix total. "
        f"{n_active} doors had a bill or a run-rate this month"
        + (f"; {n_quiet:,} universe doors with neither are omitted from the mix" if n_quiet else "")
        + (f"; {n_tail:,} tail doors are judged in the coverage panel" if n_tail else "")
        + (f"; {n_not_due:,} long-cycle doors were not due" if n_not_due else "")
        + "."
        + recon
    )
    cred = str(kpis.get("target_credibility") or "").strip()
    if cred:
        weather = weather + " " + cred
    return headline, weather


def _how_to_read(open_mtd: bool) -> list[str]:
    tiny = MIX_OTHER
    pct = int(round(PARETO_SHARE * 100))
    tol = int(round(MISS_TOL * 100))
    if open_mtd:
        return [
            "Cover Billed and Expected are the sum of this pack's shops (each shop on its own run-rate). The Situation cascade figure for the scope is printed beside it with the reconciliation factor.",
            "A shop that has not yet billed its full Expected is not a miss mid-month. Issues are due this week (Ask), visited with no bill, and lost doors.",
            f"{ISSUE_LAPSED} = quiet past max(3× the measured cycle, 45 days); a door with no measured cycle needs 60 quiet days. Billed 0 because they stopped — not a this-week miss.",
            "Ask (KG) is the 90-day expected drop when the depletion ratio is ≥ 0.8. That is the next-order number for this week.",
            f"{tiny}: within each DSR, doors outside the top {pct}% of size (and every door under {kg(MATERIAL_FLOOR_MT)} kg) — judged in the tail coverage panel, not one by one.",
            "Shop rows print kg; cover, roll-up and mix print MT. 'Every N days' appears only when a gap was measured between purchases.",
        ]
    return [
        "Read the roll-up first: it says which DSRs hold the gap and how many doors it sits in. Then the door bridge (who stopped, who started), then the shops.",
        "Cover Billed, Expected and Gap are the sum of this pack's shops, each on its own run-rate (robust last-3 / last-6, shrunk toward its city). Gap = Expected − billed (floored at 0). Mix Total is the same figure.",
        "The Situation cascade Expected for this scope is printed with its factor. Shops are never rescaled to it — a shop that billed its run-rate is On Expected whatever the city did.",
        f"{ISSUE_MISSED} = a live shop billed under Expected by more than {tol}% (and at least {kg(MISS_FLOOR_MT)} kg). Billed 0 is a miss, not a separate Unbilled row. Beat Expected is the mirror.",
        f"{ISSUE_LAPSED} = quiet past max(3× the measured cycle, 45 days); no measured cycle needs 60 quiet days. A door that billed this month is never lapsed.",
        f"{ISSUE_NOT_DUE} = a long-cycle door whose measured cycle had not fallen due by month-end. Its run-rate Expected is timing, not a hole.",
        f"{tiny}: within each DSR, doors outside the top {pct}% of size (max of Expected and billed), and every door under {kg(MATERIAL_FLOOR_MT)} kg. Any door at or above {kg(HOLE_MT)} kg is always core.",
        "Target on the cover is the plan book. The credibility line splits a shortfall into execution (vs run-rate, recoverable on the beat) and ambition (Target above run-rate).",
        "Flags mark POP codes that look like a duplicate or a migrated code (same bills on the same days, or a same-name door that started when this one stopped). Check before calling them lost.",
        "Shop rows print kg; cover, roll-up and mix print MT. 'Every N days' appears only when a gap was measured between purchases. Comment is next month, not this week.",
    ]


def _shop_labels(shops: pd.DataFrame) -> list[str]:
    names = shops.get("store_name", pd.Series("", index=shops.index)).fillna("").astype(str)
    ids = shops.get("store_id", pd.Series("", index=shops.index)).astype(str)
    folded = names.str.strip().str.casefold()
    counts = folded.value_counts()
    labels = []
    for name, sid, key in zip(names, ids, folded):
        label = name.strip() or sid
        if label and key and int(counts.get(key, 0) or 0) > 1:
            tail = sid[-6:] if len(sid) >= 4 else sid
            labels.append(f"{label} · {tail}")
        else:
            labels.append(label)
    return labels


def _present(shops: pd.DataFrame, scope: str, open_mtd: bool, has_plan: bool, period: str = "") -> pd.DataFrame:
    if shops is None or shops.empty:
        return pd.DataFrame()
    labels = _shop_labels(shops)
    rows = []
    for label, (_, r) in zip(labels, shops.iterrows()):
        rec: dict[str, Any] = {"Shop": label}
        if scope == "national":
            rec["City"] = str(r.get("city") or "")
            rec["Distributor"] = str(r.get("distributor") or "")
            rec["DSR"] = str(r.get("dsr_name") or "")
        elif scope == "city":
            rec["Distributor"] = str(r.get("distributor") or "")
            rec["DSR"] = str(r.get("dsr_name") or "")
        elif scope == "distributor":
            rec["DSR"] = str(r.get("dsr_name") or "")
        billed = kg(r.get("billed_mt"))
        expected = kg(r.get("expected_mt"))
        rec["Billed (kg)"] = billed
        rec["Expected (kg)"] = expected
        if not open_mtd:
            rec["Gap (kg)"] = expected - billed
            if has_plan:
                tgt = float(r.get("shop_target_mt") or 0)
                rec["Target (kg)"] = kg(tgt) if tgt > HOLE_MT else None
        rec["Last billed"] = _date_text(r.get("last_bill_date"), period)
        if open_mtd:
            days = r.get("days_since_bill")
            rec["Days since"] = int(round(float(days))) if days is not None and pd.notna(days) else None
            rec["Ask (KG)"] = kg(r.get("week_target_mt"))
        else:
            last_drop = float(r.get("last_drop_mt") or 0)
            rec["Last drop (kg)"] = kg(last_drop) if last_drop >= BILL_MT else 0
            usual = float(r.get("expected_drop_mt") or r.get("typical_drop_mt") or 0)
            rec["Usual drop (kg)"] = kg(usual) if usual >= BILL_MT else 0
        rec["Issue"] = str(r.get("issue") or "")
        rec["Flag"] = str(r.get("flag") or "")
        rec["Comment" if not open_mtd else "Do this"] = str(r.get("comment") or "")
        rows.append(rec)
    return pd.DataFrame(rows)


def _present_flags(shops: pd.DataFrame, scope: str) -> pd.DataFrame:
    if shops is None or shops.empty or "flag" not in shops.columns:
        return pd.DataFrame()
    part = shops.loc[shops["flag"].astype(str).str.strip().ne("")]
    if part.empty:
        return pd.DataFrame()
    rows = []
    for _, r in part.iterrows():
        rec: dict[str, Any] = {"Shop": str(r.get("store_name") or r.get("store_id") or ""), "POP": str(r.get("store_id") or "")}
        if scope in {"national", "city"}:
            rec["Distributor"] = str(r.get("distributor") or "")
        if scope in {"national", "city", "distributor"}:
            rec["DSR"] = str(r.get("dsr_name") or "")
        rec["Flag"] = str(r.get("flag") or "")
        rec["Paired with"] = str(r.get("flag_pair") or "")
        rec["Billed (kg)"] = kg(r.get("billed_mt"))
        rec["Issue"] = str(r.get("issue") or "")
        rec["Detail"] = str(r.get("flag_detail") or "")
        rows.append(rec)
    return pd.DataFrame(rows)


def _child_grain(scope: str) -> tuple[list[str], list[str]] | None:
    """Columns that define the roll-up rows one level below the scope."""
    if scope == "national":
        return ["city", "distributor"], ["City", "Distributor"]
    if scope == "city":
        return ["distributor", "dsr_name"], ["Distributor", "DSR"]
    if scope == "distributor":
        return ["dsr_name"], ["DSR"]
    return None


def _rollup(shops: pd.DataFrame, scope: str, open_mtd: bool) -> pd.DataFrame:
    """One row per unit below the scope: where the gap sits and in how many doors."""
    grain = _child_grain(scope)
    if grain is None or shops is None or shops.empty:
        return pd.DataFrame()
    keys, labels = grain
    work = shops.copy()
    for k in keys:
        if k not in work.columns:
            work[k] = ""
        work[k] = work[k].fillna("").astype(str).replace({"nan": "", "(unmapped)": ""})
    work["_billed"] = _num_col(work, "billed_mt")
    work["_expected"] = _num_col(work, "expected_mt")
    work["_remaining"] = _num_col(work, "remaining_mt")
    work["_ask"] = _num_col(work, "week_target_mt")
    issue = work["issue"].astype(str) if "issue" in work.columns else pd.Series("", index=work.index)
    work["_has_bill"] = work["_billed"] > ASK_MT
    work["_missed"] = issue.eq(ISSUE_MISSED)
    work["_missed_mt"] = work["_remaining"].where(work["_missed"], 0.0)
    work["_lost"] = issue.eq(ISSUE_LAPSED)
    work["_lost_mt"] = work["_expected"].where(work["_lost"], 0.0)
    work["_tail"] = issue.eq(MIX_OTHER)
    work["_active"] = ~issue.eq("No run-rate")
    work["_due"] = work["_ask"] > ASK_MT
    work["_beat"] = issue.eq("Beat Expected")
    work["_beat_mt"] = (work["_billed"] - work["_expected"]).clip(lower=0).where(work["_beat"], 0.0)
    g = work.groupby(keys, dropna=False)
    out = g.agg(
        doors=("_active", "sum"),
        billed_doors=("_has_bill", "sum"),
        billed=("_billed", "sum"),
        expected=("_expected", "sum"),
        missed_n=("_missed", "sum"),
        missed_mt=("_missed_mt", "sum"),
        lost_n=("_lost", "sum"),
        lost_mt=("_lost_mt", "sum"),
        beat_mt=("_beat_mt", "sum"),
        tail_n=("_tail", "sum"),
        due_n=("_due", "sum"),
        ask=("_ask", "sum"),
    ).reset_index()
    out["gap"] = (out["expected"] - out["billed"]).clip(lower=0)
    total_gap = float(out["gap"].sum())
    out["share"] = out["gap"] / total_gap * 100 if total_gap > 1e-9 else 0.0
    out = out.sort_values(["gap", "expected"], ascending=[False, False]).reset_index(drop=True)
    rows = []
    for _, r in out.iterrows():
        rec: dict[str, Any] = {}
        for k, lab in zip(keys, labels):
            rec[lab] = str(r[k])
        rec["Doors"] = int(r["doors"])
        rec["Billed doors"] = int(r["billed_doors"])
        rec["Billed (MT)"] = _mt2(r["billed"])
        rec["Expected (MT)"] = _mt2(r["expected"])
        if open_mtd:
            rec["Due shops"] = int(r["due_n"])
            rec["Ask (KG)"] = kg(r["ask"])
            rec["Lost doors"] = int(r["lost_n"])
        else:
            rec["Gap (MT)"] = _mt2(r["gap"])
            rec["Share of gap (%)"] = int(round(float(r["share"])))
            rec["Missed shops"] = int(r["missed_n"])
            rec["Missed (MT)"] = _mt2(r["missed_mt"])
            rec["Lost doors"] = int(r["lost_n"])
            rec["Lost (MT)"] = _mt2(r["lost_mt"])
            rec["Beat (MT)"] = _mt2(r["beat_mt"])
        rec["Tail doors"] = int(r["tail_n"])
        rows.append(rec)
    if rows:
        tot: dict[str, Any] = {labels[0]: MIX_TOTAL}
        for lab in labels[1:]:
            tot[lab] = ""
        tot["Doors"] = int(out["doors"].sum())
        tot["Billed doors"] = int(out["billed_doors"].sum())
        tot["Billed (MT)"] = _mt2(out["billed"].sum())
        tot["Expected (MT)"] = _mt2(out["expected"].sum())
        if open_mtd:
            tot["Due shops"] = int(out["due_n"].sum())
            tot["Ask (KG)"] = kg(out["ask"].sum())
            tot["Lost doors"] = int(out["lost_n"].sum())
        else:
            tot["Gap (MT)"] = _mt2(max(0.0, float(out["expected"].sum()) - float(out["billed"].sum())))
            tot["Share of gap (%)"] = 100 if total_gap > 1e-9 else 0
            tot["Missed shops"] = int(out["missed_n"].sum())
            tot["Missed (MT)"] = _mt2(out["missed_mt"].sum())
            tot["Lost doors"] = int(out["lost_n"].sum())
            tot["Lost (MT)"] = _mt2(out["lost_mt"].sum())
            tot["Beat (MT)"] = _mt2(out["beat_mt"].sum())
        tot["Tail doors"] = int(out["tail_n"].sum())
        rows.append(tot)
    return pd.DataFrame(rows)


def _month_billed_ids(shop_month: pd.DataFrame | None, period: str | None) -> dict[str, float]:
    if shop_month is None or shop_month.empty or not period:
        return {}
    cur = shop_month.loc[shop_month["period"].astype(str) == str(period)]
    if cur.empty:
        return {}
    vol = pd.to_numeric(cur["volume_mt"], errors="coerce").fillna(0.0)
    billed = cur.assign(_v=vol).loc[vol > ASK_MT].groupby(cur["store_id"].astype(str))["_v"].sum()
    return {str(k): float(v) for k, v in billed.items()}


def _door_bridge(
    shops: pd.DataFrame,
    shop_month: pd.DataFrame | None,
    period: str,
    scope: str,
) -> pd.DataFrame:
    """Billed-door bridge vs last month, per unit below the scope.

    Retained = billed both months. Dropped = billed last month, not this.
    New = billed this month, never before. Reactivated = billed this month
    after at least one quiet month. Volume columns say what the movers were
    worth (dropped doors at last month's bill, new / reactivated at this month's).
    """
    grain = _child_grain(scope)
    if shops is None or shops.empty or shop_month is None or shop_month.empty or not period:
        return pd.DataFrame()
    from sndintel.io_utils import shift_period

    prev = shift_period(period, -1)
    keys, labels = grain if grain is not None else ([], [])
    ids = shops["store_id"].astype(str)
    work = shops[["store_id"] + [k for k in keys if k in shops.columns]].copy()
    work["store_id"] = ids
    for k in keys:
        if k not in work.columns:
            work[k] = ""
        work[k] = work[k].fillna("").astype(str).replace({"nan": "", "(unmapped)": ""})
    now = _month_billed_ids(shop_month, period)
    last = _month_billed_ids(shop_month, prev)
    sm = shop_month.copy()
    sm["store_id"] = sm["store_id"].astype(str)
    sm["period"] = sm["period"].astype(str)
    vol = pd.to_numeric(sm["volume_mt"], errors="coerce").fillna(0.0)
    hist = sm.loc[(vol > ASK_MT) & (sm["period"] < str(period))]
    ever_before = set(hist["store_id"])
    work["_now"] = work["store_id"].map(lambda s: s in now)
    work["_last"] = work["store_id"].map(lambda s: s in last)
    work["_ever"] = work["store_id"].map(lambda s: s in ever_before)
    work["_retained"] = work["_now"] & work["_last"]
    work["_dropped"] = work["_last"] & ~work["_now"]
    work["_new"] = work["_now"] & ~work["_ever"]
    work["_react"] = work["_now"] & ~work["_last"] & work["_ever"]
    work["_dropped_mt"] = work["store_id"].map(lambda s: last.get(s, 0.0)).where(work["_dropped"], 0.0)
    work["_new_mt"] = work["store_id"].map(lambda s: now.get(s, 0.0)).where(work["_new"] | work["_react"], 0.0)
    if keys:
        g = work.groupby(keys, dropna=False)
        agg = g.agg(
            last_n=("_last", "sum"),
            retained=("_retained", "sum"),
            dropped=("_dropped", "sum"),
            new=("_new", "sum"),
            react=("_react", "sum"),
            now_n=("_now", "sum"),
            dropped_mt=("_dropped_mt", "sum"),
            new_mt=("_new_mt", "sum"),
        ).reset_index()
    else:
        agg = pd.DataFrame(
            [
                {
                    "last_n": int(work["_last"].sum()),
                    "retained": int(work["_retained"].sum()),
                    "dropped": int(work["_dropped"].sum()),
                    "new": int(work["_new"].sum()),
                    "react": int(work["_react"].sum()),
                    "now_n": int(work["_now"].sum()),
                    "dropped_mt": float(work["_dropped_mt"].sum()),
                    "new_mt": float(work["_new_mt"].sum()),
                }
            ]
        )
    agg = agg.sort_values(["dropped_mt", "dropped"], ascending=[False, False]).reset_index(drop=True)
    rows = []
    for _, r in agg.iterrows():
        rec: dict[str, Any] = {}
        for k, lab in zip(keys, labels):
            rec[lab] = str(r[k])
        rec["Billed doors LM"] = int(r["last_n"])
        rec["Retained"] = int(r["retained"])
        rec["Dropped"] = int(r["dropped"])
        rec["New"] = int(r["new"])
        rec["Reactivated"] = int(r["react"])
        rec["Billed doors TM"] = int(r["now_n"])
        rec["Dropped doors' LM bill (MT)"] = _mt2(r["dropped_mt"])
        rec["New + reactivated bill (MT)"] = _mt2(r["new_mt"])
        rows.append(rec)
    if rows and keys:
        tot: dict[str, Any] = {labels[0]: MIX_TOTAL}
        for lab in labels[1:]:
            tot[lab] = ""
        tot["Billed doors LM"] = int(agg["last_n"].sum())
        tot["Retained"] = int(agg["retained"].sum())
        tot["Dropped"] = int(agg["dropped"].sum())
        tot["New"] = int(agg["new"].sum())
        tot["Reactivated"] = int(agg["react"].sum())
        tot["Billed doors TM"] = int(agg["now_n"].sum())
        tot["Dropped doors' LM bill (MT)"] = _mt2(agg["dropped_mt"].sum())
        tot["New + reactivated bill (MT)"] = _mt2(agg["new_mt"].sum())
        rows.append(tot)
    return pd.DataFrame(rows)


def _tail_panel(
    shops: pd.DataFrame,
    shop_month: pd.DataFrame | None,
    period: str,
    scope: str,
) -> pd.DataFrame:
    """Tail doors as a coverage question per DSR: how many billed vs how many usually do.

    "Usual billed" is the mean count of tail doors that billed over the three
    prior months. A DSR whose tail billed 30 doors against a usual 55 has a
    coverage hole even if no single door is material.
    """
    if shops is None or shops.empty or "issue" not in shops.columns:
        return pd.DataFrame()
    tail = shops.loc[shops["issue"].astype(str).eq(MIX_OTHER)].copy()
    if tail.empty:
        return pd.DataFrame()
    from sndintel.io_utils import prior_periods

    keys = ["city", "distributor", "dsr_name"] if scope == "national" else (["distributor", "dsr_name"] if scope == "city" else ["dsr_name"])
    labels = {"city": "City", "distributor": "Distributor", "dsr_name": "DSR"}
    for k in keys:
        if k not in tail.columns:
            tail[k] = ""
        tail[k] = tail[k].fillna("").astype(str).replace({"nan": "", "(unmapped)": ""})
    tail["_billed"] = _num_col(tail, "billed_mt")
    tail["_expected"] = _num_col(tail, "expected_mt")
    tail["_has_bill"] = tail["_billed"] > ASK_MT
    tail_ids = set(tail["store_id"].astype(str))
    usual = pd.Series(dtype=float)
    if shop_month is not None and not shop_month.empty and period:
        prev = prior_periods(period, 3)
        sm = shop_month.copy()
        sm["store_id"] = sm["store_id"].astype(str)
        sm["period"] = sm["period"].astype(str)
        sm = sm.loc[sm["store_id"].isin(tail_ids) & sm["period"].isin(prev)]
        vol = pd.to_numeric(sm["volume_mt"], errors="coerce").fillna(0.0)
        sm = sm.loc[vol > ASK_MT]
        if not sm.empty:
            first = tail.drop_duplicates("store_id")
            key_of = {str(sid): tuple(str(v) for v in vals) for sid, vals in zip(first["store_id"], first[keys].itertuples(index=False, name=None))}
            sm = sm.assign(_key=sm["store_id"].map(key_of))
            per_period = sm.groupby(["_key", "period"])["store_id"].nunique()
            usual = per_period.groupby(level=0).sum() / max(len(prev), 1)
    g = tail.groupby(keys, dropna=False)
    agg = g.agg(
        doors=("store_id", "count"),
        billed_doors=("_has_bill", "sum"),
        billed=("_billed", "sum"),
        expected=("_expected", "sum"),
    ).reset_index()
    rows = []
    for _, r in agg.iterrows():
        rec: dict[str, Any] = {labels[k]: str(r[k]) for k in keys}
        key = tuple(str(r[k]) for k in keys)
        usual_n = float(usual.get(key, float("nan"))) if len(usual) else float("nan")
        doors = int(r["doors"])
        billed_n = int(r["billed_doors"])
        rec["Tail doors"] = doors
        rec["Billed doors"] = billed_n
        rec["Usual billed doors"] = int(round(usual_n)) if pd.notna(usual_n) else None
        rec["Doors short"] = int(round(usual_n - billed_n)) if pd.notna(usual_n) and usual_n > billed_n else 0
        rec["ECO (%)"] = int(round(billed_n / doors * 100)) if doors else 0
        rec["Billed (kg)"] = kg(r["billed"])
        rec["Expected (kg)"] = kg(r["expected"])
        rec["Gap (kg)"] = max(0, kg(r["expected"]) - kg(r["billed"]))
        rows.append(rec)
    out = pd.DataFrame(rows)
    if not out.empty:
        out = out.sort_values(["Doors short", "Gap (kg)"], ascending=[False, False]).reset_index(drop=True)
    return out



def _date_text(value: Any, period: str = "") -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    text = str(value).strip()
    if not text or text.lower() in {"nan", "nat", "none"}:
        return ""
    ts = pd.to_datetime(text, errors="coerce")
    if pd.isna(ts):
        return text
    if period and str(period)[:7] and ts.strftime("%Y-%m") != str(period)[:7]:
        return ts.strftime("%d %b %Y")
    return ts.strftime("%d %b")


# --- Excel / PDF --------------------------------------------------------------


def excel_bytes(book: ShopBook) -> bytes:
    from openpyxl.styles import Alignment, Font
    from sndintel.briefing import NAVY as HEX_NAVY, SLATE as HEX_SLATE, _fill

    wb = Workbook()
    ws = wb.active
    ws.title = "00 Cover"
    cut = "MTD" if book.open_mtd else "Closed month"
    ws["A1"] = "SND Intelligence · Shop-wise issues"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=HEX_NAVY)
    ws["A2"] = f"{book.scope_label} · {book.label} · {cut}"
    ws["A2"].font = Font(name="Calibri", size=18, bold=True, color=HEX_NAVY)
    ws["A3"] = book.headline or ""
    ws["A3"].font = Font(name="Calibri", size=12, bold=True, color=HEX_NAVY)
    ws.merge_cells("A3:H3")
    kpis = book.kpis or {}
    if book.open_mtd:
        metric_row = [
            ("Billed so far (MT)", kpis.get("billed_mt")),
            ("Expected (MT)", kpis.get("expected_mt")),
            ("Due shops", kpis.get("n_due")),
            ("Ask (KG)", int(round(float(kpis.get("ask_mt") or 0) * 1000))),
            ("Issue shops", kpis.get("n_issues")),
        ]
    else:
        metric_row = [
            ("Billed (MT)", kpis.get("billed_mt")),
            ("Expected (MT)", kpis.get("expected_mt")),
            ("Gap vs Expected (MT)", kpis.get("gap_mt")),
            ("Issue shops", kpis.get("n_issues")),
        ]
    if kpis.get("official_expected_mt") is not None:
        metric_row.append(("Cascade Expected (MT)", kpis.get("official_expected_mt")))
    if kpis.get("has_plan"):
        metric_row.append(("Target (MT)", kpis.get("target_mt")))
        metric_row.append(("vs Target (MT)", kpis.get("vs_target_mt")))
    for i, (name, val) in enumerate(metric_row, start=1):
        cell = ws.cell(5, i, name)
        cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.fill = _fill(HEX_NAVY)
        show_mt = isinstance(val, (int, float)) and str(name).endswith("(MT)")
        v = ws.cell(6, i, round(float(val), 2) if show_mt else val)
        v.font = Font(size=12, bold=True)
    ws.cell(8, 1, book.weather or "")
    ws.merge_cells("A8:H8")
    ws.cell(8, 1).font = Font(size=9, italic=True, color=HEX_SLATE)
    row = 10
    ws.cell(row, 1, "How to read")
    ws.cell(row, 1).font = Font(bold=True, size=12, color=HEX_NAVY)
    row += 1
    for step in book.how_to_read:
        ws.cell(row, 1, step)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        ws.cell(row, 1).font = Font(size=10, italic=True, color=HEX_SLATE)
        row += 1
    ws.column_dimensions["A"].width = 28
    sheet_n = 1
    if book.rollup is not None and not book.rollup.empty:
        _sheet_table(
            wb,
            f"{sheet_n:02d} Roll-up",
            "Roll-up",
            "One row per unit below this scope. Gap = Expected − billed (floored at 0). Share of gap says where to look first; "
            "Missed / Lost split the holes into live shops short of Expected and doors that stopped.",
            book.rollup,
        )
        sheet_n += 1
    if book.bridge is not None and not book.bridge.empty:
        _sheet_table(
            wb,
            f"{sheet_n:02d} Door bridge",
            "Door bridge",
            "Billed doors this month vs last month. Dropped = billed last month, not this; New = never billed before; "
            "Reactivated = back after a quiet month. Dropped doors are valued at their last-month bill.",
            book.bridge,
        )
        sheet_n += 1
    _sheet_table(
        wb,
        f"{sheet_n:02d} Issue mix",
        "Issue mix",
        f"Mix Total billed, Expected, and Gap (Expected − billed) match the cover. {MIX_OTHER} are the doors outside each DSR's top "
        f"{int(round(PARETO_SHARE * 100))}% by size. Lost-door billed is always 0. Universe doors with no bill and no run-rate are omitted from named rows (they are 0 on the Total).",
        book.mix,
    )
    sheet_n += 1
    if book.tail is not None and not book.tail.empty:
        _sheet_table(
            wb,
            f"{sheet_n:02d} Tail coverage",
            "Tail coverage",
            "Tail doors per DSR: how many billed against how many usually bill (mean of the three prior months). "
            "Doors short is the coverage hole; no single tail door is chased on its own.",
            book.tail,
        )
        sheet_n += 1
    if book.flags is not None and not book.flags.empty:
        _sheet_table(
            wb,
            f"{sheet_n:02d} Code flags",
            "Duplicate / migrated POP codes",
            "POP codes that post the same invoices on the same days under one DSR, or a same-name door that started billing when this one stopped. "
            "Check the universe before treating either as a lost door.",
            book.flags,
        )
        sheet_n += 1
    note_iss = "Shops that are an issue on this cut, largest hole / Ask first. Shop rows are in kg."
    _sheet_table(wb, f"{sheet_n:02d} Issues", "Issues", note_iss, book.issues)
    sheet_n += 1
    note_all = (
        "Every shop in this scope. Mid-month still-to-Expected is not a miss. "
        if book.open_mtd
        else "Every shop in this scope, including those that landed on Expected. "
    )
    _sheet_table(wb, f"{sheet_n:02d} All shops", "All shops", note_all + "Shop rows are in kg. Excel is the full list; the PDF is issues only.", book.shops)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pdf_bytes(book: ShopBook) -> bytes:
    buf = BytesIO()
    write_pdf(book, buf)
    return buf.getvalue()


def write_pdf(book: ShopBook, path: Path | str | BytesIO) -> None:
    from sndintel.situation_report import (
        LINE,
        NAVY,
        WASH,
        WHITE,
        _pdf_styles,
        _pdf_table,
    )

    pagesize = landscape(A4)
    doc = SimpleDocTemplate(
        path if not isinstance(path, (str, Path)) else str(path),
        pagesize=pagesize,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=16 * mm,
        bottomMargin=12 * mm,
        title=f"SND Intelligence · {book.scope_label} shops",
        author="SND Intelligence",
    )
    styles = _pdf_styles()
    usable = pagesize[0] - doc.leftMargin - doc.rightMargin
    cut = "MTD SHOP ISSUES" if book.open_mtd else "CLOSED-MONTH SHOP ISSUES"
    story: list[Any] = [
        Paragraph(cut, styles["kicker"]),
        Paragraph(xml_escape(f"{book.scope_label} · {book.label}"), styles["h1"]),
        Paragraph(xml_escape(book.headline or ""), styles["headline"]),
    ]
    kpis = book.kpis or {}
    if book.open_mtd:
        kpi_cells = [
            ("Billed so far", fmt_mt(kpis.get("billed_mt"))),
            ("Expected", fmt_mt(kpis.get("expected_mt"))),
            ("Due shops", str(int(kpis.get("n_due") or 0))),
            ("Ask", f"{kg(kpis.get('ask_mt')):,} KG"),
            ("Issue shops", str(int(kpis.get("n_issues") or 0))),
        ]
    else:
        kpi_cells = [
            ("Billed", fmt_mt(kpis.get("billed_mt"))),
            ("Expected", fmt_mt(kpis.get("expected_mt"))),
            ("Gap vs Expected", fmt_mt(kpis.get("gap_mt"))),
            ("Issue shops", str(int(kpis.get("n_issues") or 0))),
        ]
    if kpis.get("official_expected_mt") is not None:
        kpi_cells.append(("Cascade Expected", fmt_mt(kpis.get("official_expected_mt"))))
    if kpis.get("has_plan"):
        kpi_cells.append(("Target", fmt_mt(kpis.get("target_mt"))))
        vs = float(kpis.get("vs_target_mt") or 0)
        kpi_cells.append(("vs Target", ("+" if vs > 0 else "") + fmt_mt(vs)))
    labels = [Paragraph(xml_escape(n), styles["kpi_l"]) for n, _ in kpi_cells]
    values = [Paragraph(xml_escape(v), styles["kpi_v"]) for _, v in kpi_cells]
    n = max(len(kpi_cells), 1)
    kpi_table = Table([labels, values], colWidths=[usable / n] * n)
    kpi_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), NAVY),
                ("BACKGROUND", (0, 1), (-1, 1), WASH),
                ("BOX", (0, 0), (-1, -1), 0.4, LINE),
                ("INNERGRID", (0, 0), (-1, -1), 0.25, LINE),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, 0), 5),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 4),
                ("TOPPADDING", (0, 1), (-1, 1), 6),
                ("BOTTOMPADDING", (0, 1), (-1, 1), 6),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    story.append(kpi_table)
    story.append(Spacer(1, 6))
    story.append(Paragraph(xml_escape(book.weather or ""), styles["lead"]))
    story.append(Paragraph("How to read", styles["h3"]))
    for line in book.how_to_read:
        story.append(Paragraph(xml_escape("- " + line), styles["body"]))
    rollup = book.rollup
    if rollup is not None and not rollup.empty:
        heading = Paragraph("Where the gap sits", styles["h2"])
        note = Paragraph(
            "One row per unit below this scope, largest gap first. Missed = live shops short of Expected; "
            "Lost = doors that stopped; Tail = small doors judged in the coverage panel.",
            styles["note"],
        )
        table = _pdf_table(rollup.head(40), styles, usable)
        story.append(KeepTogether([heading, note, Spacer(1, 2), table]))
    bridge = book.bridge
    if bridge is not None and not bridge.empty:
        heading = Paragraph("Door bridge vs last month", styles["h2"])
        note = Paragraph(
            "Retained billed both months. Dropped billed last month and not this (valued at last month's bill). "
            "New never billed before; Reactivated came back after a quiet month.",
            styles["note"],
        )
        table = _pdf_table(bridge.head(40), styles, usable)
        story.append(KeepTogether([heading, note, Spacer(1, 2), table]))
    mix = book.mix
    if mix is not None and not mix.empty:
        heading = Paragraph("Issue mix", styles["h2"])
        note = Paragraph(
            "Mix Total billed, Expected, and Gap (Expected − billed) match the cover. "
            f"{MIX_OTHER} are the doors outside each DSR's top {int(round(PARETO_SHARE * 100))}% by size. "
            f"{ISSUE_LAPSED} billed is always 0 — they stopped, they are not a this-month miss.",
            styles["note"],
        )
        table = _pdf_table(mix, styles, usable)
        story.append(KeepTogether([heading, note, Spacer(1, 2), table]))
    tail = book.tail
    if tail is not None and not tail.empty:
        heading = Paragraph("Tail coverage", styles["h2"])
        note = Paragraph(
            "Small doors per DSR: billed doors against the usual count (mean of the three prior months). "
            "Doors short is the coverage hole to close on the beat.",
            styles["note"],
        )
        table = _pdf_table(tail.head(40), styles, usable)
        story.append(KeepTogether([heading, note, Spacer(1, 2), table]))
    flags_df = book.flags
    if flags_df is not None and not flags_df.empty:
        heading = Paragraph("POP codes to check", styles["h2"])
        note = Paragraph(
            "Same invoices on the same days under one DSR, or a same-name door that started when this one stopped. "
            "Fix the universe before chasing these as lost doors.",
            styles["note"],
        )
        table = _pdf_table(flags_df.head(40), styles, usable)
        story.append(KeepTogether([heading, note, Spacer(1, 2), table]))
    issues = book.issues
    n_iss = 0 if issues is None or issues.empty else int(len(issues))
    shown = issues.head(PDF_ISSUE_N) if n_iss else issues
    heading = Paragraph("Shop issues", styles["h2"])
    extra = f" Showing {PDF_ISSUE_N} of {n_iss}. Excel has every shop." if n_iss > PDF_ISSUE_N else " Excel has every shop in this scope."
    note = Paragraph(
        (
            "Working list: live shops that missed Expected (including billed 0) and lost doors, in kg. "
            "The mix Total is the full pack. "
        )
        + extra.strip(),
        styles["note"],
    )
    table = _pdf_table(shown, styles, usable)
    story.append(PageBreak())
    story.append(KeepTogether([heading, note, Spacer(1, 2)]))
    story.append(table)
    doc.build(story)


def write_excel(book: ShopBook, path: Path | str) -> None:
    Path(path).write_bytes(excel_bytes(book))
