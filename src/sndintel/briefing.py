"""One strategy pack: country → top lagging distributors → DSRs → doors, plus full lists.

Built for a sales head to print or filter. Column names are in English; jargon lives
in the glossary, not the headers. Excel is the working file. PDF is the board pack.
The national summary names the top 10 distributors by Gap, the top 10 DSRs from
those distributors, and the top 50 shops. Report packs and the detailed pack are full lists.
"""

from __future__ import annotations

import html
import json
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.worksheet import Worksheet

from sndintel.coverage import allocate_recoverable_drivers, attach_remarks, sibling_z_frame
from sndintel.hierarchy import _shop_gaps
from sndintel.isolate import empirical_bayes, k_from_ly
from sndintel.io_utils import shift_period
from sndintel.season import fit_seasonality, fit_shop_expected, reconcile_expected


SITUATION_LABEL = {
    "lagging": "Lagging",
    "with_market": "On expected",
    "outperforming": "Ahead",
}

DRIVER_LABEL = {
    "drop_size": "Drop size — same doors, smaller drops",
    "coverage": "Coverage — fewer billed doors",
    "whitespace": "Whitespace — universe not billed",
    "mixed": "Mixed — more than one driver",
    "holding": "Holding",
    "mix": "SKU mix",
}

GLOSSARY = [
    ("Billed this period", "Secondary volume in the month being scored (MTD if the month is still open)."),
    ("AMS last 3 months", "Average monthly secondary volume over the last three *closed* months. A typical recent month, not last year."),
    ("vs AMS", "This period minus AMS × fraction of the month elapsed. Negative = behind the recent run-rate. On day 20 that is billed − AMS × (20 ÷ days in month), not billed − AMS."),
    ("Same month last year", "What this unit billed in the same calendar month a year ago (full closed month). Zero means no August last year — it does not mean Expected should be zero."),
    ("Expected this month", "Recent run-rate: mean of the last three closed months (same window as AMS), blended with the last-six-month median, paced if MTD is open. Same method at country, city, distributor, DSR, and shop. Calendar-month seasonality is not applied — a city with no August history still expects its recent monthly run-rate. Children’s Expecteds are then scaled so they add to the parent."),
    ("Gap", "The hole versus this unit’s own Expected, as a positive number — volume that comes back if the unit billed its recent run-rate. Country Gap is the country miss versus Expected. Zero means billed at or above Expected, not that AMS is irrelevant."),
    ("From drop size (MT)", "Share of the gap explained by smaller (or larger) drops on billed doors. Positive = part of the hole. Negative = billed more than Expected. The three From columns add to Gap when the unit is behind."),
    ("From unvisited shops (MT)", "Share of the gap from universe doors that were not called this period (visit count 0 and not billed). Positive = hole; negative = ahead of Expected."),
    ("From unbilled shops (MT)", "Share of the gap from doors that were visited (or, if no visit file, simply not billed) but did not buy. Positive = hole; negative = ahead of Expected."),
    ("Drop size (MT)", "Average billed volume per billed shop this period (billed MT ÷ billed shops). Not the same as From drop size, which is that driver’s share of Gap."),
    ("Remarks", "Four bullets: trend vs AMS, vs Expected, and YoY; visit coverage vs country; productivity (billed ÷ visited) vs country; drop size vs expected drop vs national average. Expected drop is Expected volume ÷ Expected billed shops — same last-3 / last-6 run-rate as Expected sales, not this month’s shop count and not paced."),
    ("Visit %", "Universe shops visited this period ÷ universe. A billed shop counts as visited even if the visit file missed it."),
    ("Strike %", "Billed shops ÷ universe shops on the live universe list."),
    ("Live universe", "The Universe Shop List is the only book that can sell. POP code is the shop. Names/DSR/distributor/city follow the current list. Closed POPs (not on the list) are dropped from history for scoring."),
    ("Shop lists", "Summary pack: top 50 lagging doors by Gap (after dropping gaps ≤ 0.25 MT). The rest of the hole is one remainder line. Detailed pack still lists every door above 0.25 MT."),
    ("AMS = 0 distributors / DSRs", "Hidden everywhere in the report. No recent three-month run-rate, so they are not a call."),
    (
        "Summary pack lists",
        "Hard caps by Gap, not a statistical cut: top 10 lagging distributors nationally; top 10 lagging DSRs whose shops sit under those 10 distributors; top 50 lagging shops. Remainder lines are the tail. The detailed pack is the full list (AMS > 0 distributors/DSRs; shops above 0.25 MT).",
    ),
]

CALCULATION_NOTES = [
    (
        "Expected this month",
        "Mean of the last three closed months (same window as AMS), blended with the last-six-month median. Calendar-month seasonality is not used: an empty August last year does not zero out a city that has been billing 40 MT/month recently. The same recipe runs at country, city, distributor, DSR, and shop. Children’s Expecteds are then scaled so they add to the parent Expected. If the month is still open, Expected is that full-month run-rate × the intra-month fraction (elapsed days, or learned MTD cuts when those exist).",
    ),
    (
        "Gap",
        "max(0, Expected − billed), after a small empirical-Bayes shrink so a noisy 0.02 MT door cannot outrank a large one. Country Gap is the country miss versus Expected.",
    ),
    (
        "From drop / unvisited / unbilled",
        "Shop-level identity: opportunity is that shop’s Expected (fallback AMS × pace, else last-year × pace). Unvisited = not called and not billed. Unbilled = called (or, with no visit file, simply not billed) and did not buy. Drop size = billed volume versus opportunity on billed doors. Those three are weights, then scaled so they add to Gap. Positive = part of the hole; negative = billed more than Expected.",
    ),
    (
        "Drop size (MT) versus From drop size",
        "Drop size is billed MT ÷ billed shops this period (two decimals). Expected drop is Expected volume ÷ Expected billed shops, using the same last-three-month / last-six-month recipe as Expected sales — not paced Expected ÷ this month’s billed shops. From drop size is that driver’s share of Gap, not the average drop.",
    ),
    (
        "vs AMS",
        "Billed minus (AMS of the last three closed months × fraction of the month elapsed). Negative = behind the recent run-rate. This is not billed minus the full-month AMS.",
    ),
    (
        "Visit % and Strike %",
        "Visit % = visited ÷ universe. A billed shop counts as visited even if the visit file missed it. Strike % = billed ÷ universe. High visit and low strike means conversion (unbilled), not coverage (unvisited).",
    ),
    (
        "Open MTD",
        "Billed is month-to-date. Expected and AMS comparisons are paced. Last year is the full closed same month.",
    ),
    (
        "Rounding and lists",
        "MT, shop counts, and percents print as whole numbers. From-columns are adjusted so they still add to Gap after rounding. Drop size stays two decimals. Distributors and DSRs with AMS = 0 are hidden. The summary pack then names the top 10 lagging distributors by Gap, the top 10 lagging DSRs from those distributors, and the top 50 lagging shops (after the 0.25 MT floor). Remainder lines are the tail. The detailed pack lists every AMS > 0 distributor/DSR and every shop above 0.25 MT.",
    ),
    (
        "Top-N lists",
        "National summary only. Rank lagging units by Gap. Keep 10 distributors, then 10 DSRs that appear on doors under those 10 distributors, then 50 shops (Gap > 0.25 MT). City / distributor / DSR Report packs and the detailed pack are still the full lists.",
    ),
]


def how_to_read_steps(pack: StrategyPack, detailed: bool = False) -> list[str]:
    scope = (pack.scope or "national").lower()
    if scope == "city":
        return [
            "City scorecard versus its own Expected. Gap is billed versus that recent run-rate.",
            "Every distributor in this city with AMS greater than 0.",
            "Every DSR in this city with AMS greater than 0.",
            "Shops in this city with gap greater than 0.25 MT.",
        ]
    if scope == "distributor":
        return [
            "Distributor scorecard versus its own Expected.",
            "DSRs on this distributor’s doors.",
            "Shops under this distributor with gap greater than 0.25 MT.",
        ]
    if scope == "dsr":
        return [
            "DSR scorecard versus its own Expected.",
            "Shops on this beat with gap greater than 0.25 MT.",
        ]
    if detailed:
        return [
            "City detail — every city, highest gap first.",
            "Distributor detail — every distributor with AMS greater than 0.",
            "DSR detail — every DSR with AMS greater than 0.",
            "National shops — every door with gap greater than 0.25 MT.",
        ]
    return [
        "Country by city — every city versus its own Expected. Highest gap first.",
        "Top 10 lagging distributors nationally, biggest Gap first. AMS = 0 is hidden.",
        "Top 10 lagging DSRs from those 10 distributors. Ride-with this list.",
        "Top 50 lagging shops by Gap (doors ≤ 0.25 MT and the tail are the remainder line).",
    ]

SHOP_RECOVERABLE_FLOOR = 0.25
SUMMARY_DIST_N = 10
SUMMARY_DSR_N = 10
SUMMARY_SHOP_N = 50

SCORECARD_METRICS = [
    ("volume_mt", "Billed this period (MT)"),
    ("ams_3m", "AMS last 3 months (MT)"),
    ("vs_ams_mt", "vs AMS (MT)"),
    ("ly_mt", "Same month last year (MT)"),
    ("expected_mt", "Expected this month (MT)"),
    ("recoverable_mt", "Gap (MT)"),
    ("from_drop_size_mt", "From drop size (MT)"),
    ("from_unvisited_mt", "From unvisited shops (MT)"),
    ("from_unbilled_mt", "From unbilled shops (MT)"),
    ("billed", "Billed shops"),
    ("visited", "Visited shops"),
    ("universe", "Universe"),
    ("strike_pct", "Strike %"),
    ("visit_pct", "Visit %"),
    ("drop_size_mt", "Drop size (MT)"),
    ("remarks", "Remarks"),
]

CITY_VIEW = [("grain_id", "City"), *SCORECARD_METRICS]

DIST_IN_CITY_VIEW = [("city", "City"), ("grain_id", "Distributor"), *SCORECARD_METRICS]

DIST_ALL_VIEW = [("grain_id", "Distributor"), ("city", "City"), *SCORECARD_METRICS]

DSR_VIEW = [("grain_id", "DSR"), ("city", "City"), *SCORECARD_METRICS]

SHOP_VIEW = [
    ("store_name", "Shop"),
    ("city", "City"),
    ("distributor", "Distributor"),
    ("dsr_name", "DSR"),
    ("section", "Beat"),
    ("volume_mt", "Billed this period (MT)"),
    ("ams_3m", "AMS last 3 months (MT)"),
    ("vs_ams_mt", "vs AMS (MT)"),
    ("ly_mt", "Same month last year (MT)"),
    ("recoverable_mt", "Gap (MT)"),
    ("visits", "Visits MTD"),
    ("call_status", "Call"),
]


@dataclass
class StrategyPack:
    period: str
    label: str
    headline: str = ""
    weather: str = ""
    problem: str = ""
    action: str = ""
    kpis: dict[str, Any] = field(default_factory=dict)
    cities: pd.DataFrame = field(default_factory=pd.DataFrame)
    city_distributors: pd.DataFrame = field(default_factory=pd.DataFrame)
    city_distributor_shops: pd.DataFrame = field(default_factory=pd.DataFrame)
    lagging_distributors: pd.DataFrame = field(default_factory=pd.DataFrame)
    lagging_dsrs: pd.DataFrame = field(default_factory=pd.DataFrame)
    lagging_shops: pd.DataFrame = field(default_factory=pd.DataFrame)
    lagging_city_names: list[str] = field(default_factory=list)
    lagging_distributor_names: list[str] = field(default_factory=list)
    shop_note: str = ""
    all_distributors: pd.DataFrame = field(default_factory=pd.DataFrame)
    all_dsrs: pd.DataFrame = field(default_factory=pd.DataFrame)
    all_shops: pd.DataFrame = field(default_factory=pd.DataFrame)
    scope: str = "national"
    scope_label: str = ""
    exec_situation: list[str] = field(default_factory=list)
    exec_focus: list[dict[str, str]] = field(default_factory=list)
    exec_error: str = ""
    exec_model: str = ""


def is_national_pack(pack: StrategyPack) -> bool:
    return (pack.scope or "national").strip().lower() in {"", "national", "country"}


def apply_exec_summary(pack: StrategyPack, row: dict[str, Any] | None) -> StrategyPack:
    """Attach a stored national executive summary onto the pack."""
    from dataclasses import replace

    if not row:
        return pack
    situation = _parse_exec_list(row.get("situation_json") or row.get("situation"))
    focus = _parse_exec_focus(row.get("focus_json") or row.get("focus"))
    return replace(
        pack,
        exec_situation=situation,
        exec_focus=focus,
        exec_error=str(row.get("error") or ""),
        exec_model=str(row.get("model") or ""),
    )


def _parse_exec_list(raw: Any) -> list[str]:
    if isinstance(raw, list):
        return [str(x).strip() for x in raw if str(x).strip()]
    if isinstance(raw, str) and raw.strip():
        try:
            parsed = json.loads(raw)
        except json.JSONDecodeError:
            return [raw.strip()]
        if isinstance(parsed, list):
            return [str(x).strip() for x in parsed if str(x).strip()]
        if parsed:
            return [str(parsed).strip()]
    return []


def _parse_exec_focus(raw: Any) -> list[dict[str, str]]:
    items: list[Any]
    if isinstance(raw, list):
        items = raw
    elif isinstance(raw, str) and raw.strip():
        try:
            parsed = json.loads(raw)
        except json.JSONDecodeError:
            return []
        items = parsed if isinstance(parsed, list) else []
    else:
        return []
    out: list[dict[str, str]] = []
    for item in items:
        if not isinstance(item, dict):
            continue
        title = str(item.get("title") or "").strip()
        why = str(item.get("why") or "").strip()
        do = str(item.get("do") or item.get("action") or "").strip()
        if title or why or do:
            out.append({"title": title, "why": why, "do": do})
    return out


def build_strategy_pack(
    units: pd.DataFrame,
    shop_month: pd.DataFrame,
    situation: pd.DataFrame | None = None,
    ledger: pd.DataFrame | None = None,
    national: dict[str, Any] | None = None,
    period: str | None = None,
    visits: pd.DataFrame | None = None,
    exec_summary: dict[str, Any] | None = None,
) -> StrategyPack:
    empty = StrategyPack(period=period or "", label=period or "")
    if units is None or units.empty:
        return empty
    period = period or str(units["period"].dropna().astype(str).max() or "")
    nat_u = units[units["grain"] == "national"]
    nat_row = nat_u.iloc[0] if not nat_u.empty else None
    pace = 1.0
    if nat_row is not None and pd.notna(nat_row.get("intra_month_frac")):
        pace = float(nat_row["intra_month_frac"] or 1.0) or 1.0
    elif "intra_month_frac" in units.columns:
        pace = float(pd.to_numeric(units["intra_month_frac"], errors="coerce").dropna().max() or 1.0)

    sit = _situation_row(situation, national, nat_row)
    label = sit.get("label") or period
    kpis = _kpis(nat_row, units, sit)

    cities = _grain(units, "city")
    dists = _grain(units, "distributor")
    dsrs = _grain(units, "dsr")
    if not cities.empty:
        cities["city"] = cities["grain_id"]
        cities = _attach_ams(cities, shop_month, period, ["city"], ledger, pace)
    if not dists.empty:
        dists["city"] = dists["parent_id"]
        dists["distributor"] = dists["grain_id"]
        dists = _attach_ams(dists, shop_month, period, ["city", "distributor"], ledger, pace)
        dists = _drop_zero_ams(dists)
    if not dsrs.empty:
        dsrs["city"] = dsrs["parent_id"]
        dsrs["dsr_name"] = dsrs["grain_id"]
        dsrs = _attach_ams(dsrs, shop_month, period, ["city", "dsr_name"], ledger, pace)
        dsrs = _drop_zero_ams(dsrs)

    cities = _sort_focus(cities)
    dists = _sort_focus(dists)
    dsrs = _sort_focus(dsrs)

    nat_grain = _grain(units, "national")
    if not nat_grain.empty:
        nat_grain = _attach_national_ams(nat_grain, shop_month, period, ledger, pace)
        nat_grain["grain_id"] = "Country"
        nat_grain["city"] = "Country"
        nat_grain["zone"] = ""
        nat_grain["situation_label"] = "Country"
        nat_grain["share_expected_mt"] = nat_grain.get("expected_mt")
        nat_grain = allocate_recoverable_drivers(nat_grain)
        nat_parent = _parent_stats(nat_grain.iloc[0])
        nat_grain = attach_remarks(nat_grain, {}, None)
    else:
        nat_parent = {}

    cities = allocate_recoverable_drivers(cities)
    dists = allocate_recoverable_drivers(dists)
    dsrs = allocate_recoverable_drivers(dsrs)
    cities = attach_remarks(cities, nat_parent, sibling_z_frame(cities))
    dists = attach_remarks(dists, nat_parent, sibling_z_frame(dists))
    dsrs = attach_remarks(dsrs, nat_parent, sibling_z_frame(dsrs))

    lagging_cities = cities[cities["situation"] == "lagging"] if not cities.empty else cities
    lagging_city_names = [str(x) for x in lagging_cities["grain_id"].tolist()] if not lagging_cities.empty else []

    all_lag_dist = dists[dists["situation"] == "lagging"].copy() if not dists.empty else dists
    all_lag_dist, lag_dist_meta = keep_top_holes(all_lag_dist, SUMMARY_DIST_N)
    all_lag_dist = _sort_focus(all_lag_dist)

    lagging_dist_names: list[str] = []
    if not all_lag_dist.empty:
        lagging_dist_names = [str(x) for x in all_lag_dist["grain_id"].tolist()]

    # Same ten names on both distributor sheets — not a per-city statistical cut.
    city_dists = all_lag_dist
    city_dist_meta = lag_dist_meta

    shops = score_shops(shop_month, cities, period, pace)
    shops = _attach_ams(shops, shop_month, period, ["store_id"], ledger, pace)
    shops = _attach_shop_calls(shops, visits, period)
    hole_floor = SHOP_RECOVERABLE_FLOOR

    city_dist_shops = pd.DataFrame()
    city_dist_shop_meta = {"n_hidden": 0, "hidden_mt": 0.0, "top_n": SUMMARY_SHOP_N}
    if not shops.empty and lagging_dist_names:
        city_dist_shops = shops[shops["distributor"].astype(str).isin(set(lagging_dist_names))].copy()
        city_dist_shops, floor_meta = keep_visit_shops(city_dist_shops, hole_floor)
        city_dist_shops, top_meta = keep_top_holes(city_dist_shops, SUMMARY_SHOP_N)
        city_dist_shop_meta = _merge_hole_meta(floor_meta, top_meta)

    all_lag_dsr = dsrs[dsrs["situation"] == "lagging"].copy() if not dsrs.empty else dsrs
    if all_lag_dsr is None or all_lag_dsr.empty or not lagging_dist_names:
        all_lag_dsr = all_lag_dsr.iloc[0:0].copy() if all_lag_dsr is not None else pd.DataFrame()
        lag_dsr_meta = keep_top_holes(all_lag_dsr, SUMMARY_DSR_N)[1]
    else:
        dsr_keys = _dsr_keys_under_distributors(shop_month, set(lagging_dist_names), period, shops)
        all_lag_dsr = _filter_dsrs_by_keys(all_lag_dsr, dsr_keys)
        all_lag_dsr, lag_dsr_meta = keep_top_holes(all_lag_dsr, SUMMARY_DSR_N)
        all_lag_dsr = _sort_focus(all_lag_dsr)

    lag_shops = pd.DataFrame()
    lag_meta = {"n_hidden": 0, "hidden_mt": 0.0, "top_n": SUMMARY_SHOP_N}
    if not shops.empty:
        lag_shops, floor_meta = keep_visit_shops(shops, hole_floor)
        lag_shops, top_meta = keep_top_holes(lag_shops, SUMMARY_SHOP_N)
        lag_meta = _merge_hole_meta(floor_meta, top_meta)

    all_shops, all_shop_meta = keep_visit_shops(shops, hole_floor) if not shops.empty else (pd.DataFrame(), lag_meta)

    kpis["n_lagging_distributors"] = int(len(all_lag_dist)) if all_lag_dist is not None else 0
    kpis["n_lagging_dsrs"] = int(len(all_lag_dsr)) if all_lag_dsr is not None else 0
    kpis["n_lagging_shops"] = int(lag_meta.get("n_kept") or 0)
    kpis["shop_hole_floor_mt"] = hole_floor
    kpis["n_shops_hidden"] = int(lag_meta.get("n_hidden") or 0)
    kpis["hidden_shop_recoverable_mt"] = float(lag_meta.get("hidden_mt") or 0)
    kpis["n_dists_hidden"] = int(lag_dist_meta.get("n_hidden") or 0)
    kpis["n_dsrs_hidden"] = int(lag_dsr_meta.get("n_hidden") or 0)
    shop_note = (
        f"Top {SUMMARY_SHOP_N} lagging shops by Gap. "
        f"{int(lag_meta.get('n_hidden') or 0)} other doors totalling "
        f"{float(lag_meta.get('hidden_mt') or 0):.0f} MT gap are the remainder line "
        f"(includes gaps ≤ {hole_floor:.2f} MT)."
    )

    pack = StrategyPack(
        period=period,
        label=label,
        headline=sit.get("headline") or "",
        weather=sit.get("weather") or "",
        problem=sit.get("problem") or "",
        action=sit.get("action") or sit.get("action_summary") or "",
        kpis=kpis,
        cities=_present_with_national(nat_grain, cities, CITY_VIEW),
        city_distributors=_present_with_remainder(city_dists, DIST_IN_CITY_VIEW, city_dist_meta, "Distributor", "distributors"),
        city_distributor_shops=_present_shops(city_dist_shops, city_dist_shop_meta),
        lagging_distributors=_present_with_remainder(all_lag_dist, DIST_ALL_VIEW, lag_dist_meta, "Distributor", "distributors"),
        lagging_dsrs=_present_with_remainder(all_lag_dsr, DSR_VIEW, lag_dsr_meta, "DSR", "DSRs"),
        lagging_shops=_present_shops(lag_shops, lag_meta),
        lagging_city_names=lagging_city_names,
        lagging_distributor_names=lagging_dist_names,
        shop_note=shop_note,
        all_distributors=_present(dists, DIST_ALL_VIEW),
        all_dsrs=_present(dsrs, DSR_VIEW),
        all_shops=_present_shops(all_shops, all_shop_meta),
        scope="national",
        scope_label="Country",
    )
    return apply_exec_summary(pack, exec_summary)


def list_report_entities(pack: StrategyPack, report_type: str) -> list[str]:
    """Searchable options for the Report page second dropdown."""
    kind = (report_type or "national").strip().lower()
    if kind == "city":
        if pack.cities is None or pack.cities.empty or "City" not in pack.cities.columns:
            return []
        return sorted(
            {str(x) for x in pack.cities["City"].dropna().astype(str) if str(x) not in {"", "Country"}}
        )
    if kind == "distributor":
        src = pack.all_distributors if pack.all_distributors is not None and not pack.all_distributors.empty else pack.lagging_distributors
        if src is None or src.empty:
            return []
        city_col = "City" if "City" in src.columns else None
        name_col = "Distributor" if "Distributor" in src.columns else None
        if not name_col:
            return []
        out = []
        for _, r in src.iterrows():
            name = str(r[name_col])
            city = str(r[city_col]) if city_col else ""
            out.append(f"{city} · {name}" if city and city != "nan" else name)
        return sorted(set(out))
    if kind == "dsr":
        src = pack.all_dsrs if pack.all_dsrs is not None and not pack.all_dsrs.empty else pack.lagging_dsrs
        if src is None or src.empty:
            return []
        city_col = "City" if "City" in src.columns else None
        name_col = "DSR" if "DSR" in src.columns else None
        if not name_col:
            return []
        out = []
        for _, r in src.iterrows():
            name = str(r[name_col])
            city = str(r[city_col]) if city_col else ""
            out.append(f"{city} · {name}" if city and city != "nan" else name)
        return sorted(set(out))
    return []


def _split_entity(entity: str) -> tuple[str | None, str]:
    text = str(entity or "").strip()
    if " · " in text:
        left, right = text.split(" · ", 1)
        return left.strip(), right.strip()
    return None, text


def _filter_table(df: pd.DataFrame, col: str, value: str) -> pd.DataFrame:
    if df is None or df.empty or col not in df.columns:
        return df if df is not None else pd.DataFrame()
    return df[df[col].astype(str) == str(value)].copy()


def focus_pack(pack: StrategyPack, report_type: str, entity: str) -> StrategyPack:
    """Narrow a national pack to one city, distributor, or DSR plus its children."""
    from dataclasses import replace

    kind = (report_type or "national").strip().lower()
    if kind in {"national", "country", ""}:
        return pack
    pack = replace(
        pack,
        exec_situation=[],
        exec_focus=[],
        exec_error="",
        exec_model="",
    )
    city_key, name = _split_entity(entity)
    if kind == "city":
        city = name
        cities = pack.cities
        keep_cities = cities[cities["City"].astype(str).isin(["Country", city])] if cities is not None and not cities.empty else cities
        dists = _filter_table(pack.all_distributors, "City", city)
        dsrs = _filter_table(pack.all_dsrs, "City", city)
        shops = _filter_table(pack.all_shops, "City", city)
        lag_d = _filter_table(pack.lagging_distributors, "City", city)
        lag_s = _filter_table(pack.lagging_dsrs, "City", city)
        lag_shops = _filter_table(pack.lagging_shops, "City", city)
        city_dists = _filter_table(pack.city_distributors, "City", city)
        city_shops = _filter_table(pack.city_distributor_shops, "City", city)
        headline = f"{city} — city pack"
        return replace(
            pack,
            headline=headline or pack.headline,
            cities=keep_cities,
            city_distributors=city_dists if city_dists is not None and not city_dists.empty else dists,
            city_distributor_shops=city_shops if city_shops is not None and not city_shops.empty else shops,
            lagging_distributors=lag_d,
            lagging_dsrs=lag_s,
            lagging_shops=lag_shops if lag_shops is not None and not lag_shops.empty else shops,
            all_distributors=dists,
            all_dsrs=dsrs,
            all_shops=shops,
            scope="city",
            scope_label=city,
        )
    if kind == "distributor":
        dist = name
        dists_src = pack.all_distributors if pack.all_distributors is not None and not pack.all_distributors.empty else pack.lagging_distributors
        row = pd.DataFrame()
        if dists_src is not None and not dists_src.empty:
            mask = dists_src["Distributor"].astype(str) == dist
            if city_key and "City" in dists_src.columns:
                mask = mask & (dists_src["City"].astype(str) == city_key)
            row = dists_src.loc[mask].copy()
        city = city_key or (str(row.iloc[0]["City"]) if not row.empty and "City" in row.columns else "")
        shops_src = pack.all_shops if pack.all_shops is not None and not pack.all_shops.empty else pack.lagging_shops
        shops = _filter_table(shops_src, "Distributor", dist)
        if city and shops is not None and not shops.empty and "City" in shops.columns:
            shops = shops[shops["City"].astype(str) == city]
        dsrs = pd.DataFrame()
        if shops is not None and not shops.empty and "DSR" in shops.columns:
            names = set(shops["DSR"].dropna().astype(str))
            dsrs_src = pack.all_dsrs if pack.all_dsrs is not None and not pack.all_dsrs.empty else pack.lagging_dsrs
            if dsrs_src is not None and not dsrs_src.empty:
                dsrs = dsrs_src[dsrs_src["DSR"].astype(str).isin(names)].copy()
                if city and "City" in dsrs.columns:
                    dsrs = dsrs[dsrs["City"].astype(str) == city]
        headline = f"{dist} — distributor pack"
        return replace(
            pack,
            headline=headline,
            cities=row,
            city_distributors=row,
            city_distributor_shops=shops,
            lagging_distributors=row,
            lagging_dsrs=dsrs,
            lagging_shops=shops,
            all_distributors=row,
            all_dsrs=dsrs,
            all_shops=shops,
            scope="distributor",
            scope_label=f"{city} · {dist}" if city else dist,
        )
    if kind == "dsr":
        dsr = name
        dsrs_src = pack.all_dsrs if pack.all_dsrs is not None and not pack.all_dsrs.empty else pack.lagging_dsrs
        row = pd.DataFrame()
        if dsrs_src is not None and not dsrs_src.empty:
            mask = dsrs_src["DSR"].astype(str) == dsr
            if city_key and "City" in dsrs_src.columns:
                mask = mask & (dsrs_src["City"].astype(str) == city_key)
            row = dsrs_src.loc[mask].copy()
        city = city_key or (str(row.iloc[0]["City"]) if not row.empty and "City" in row.columns else "")
        shops_src = pack.all_shops if pack.all_shops is not None and not pack.all_shops.empty else pack.lagging_shops
        shops = _filter_table(shops_src, "DSR", dsr)
        if city and shops is not None and not shops.empty and "City" in shops.columns:
            shops = shops[shops["City"].astype(str) == city]
        headline = f"{dsr} — DSR pack"
        return replace(
            pack,
            headline=headline,
            cities=row,
            city_distributors=pd.DataFrame(),
            city_distributor_shops=shops,
            lagging_distributors=pd.DataFrame(),
            lagging_dsrs=row,
            lagging_shops=shops,
            all_distributors=pd.DataFrame(),
            all_dsrs=row,
            all_shops=shops,
            scope="dsr",
            scope_label=f"{city} · {dsr}" if city else dsr,
        )
    return pack


def score_shops(shop_month: pd.DataFrame, cities: pd.DataFrame, period: str, pace: float) -> pd.DataFrame:
    """Shop hole versus its own Expected. Recoverable = that miss after EB shrink."""
    if shop_month is None or shop_month.empty or not period:
        return pd.DataFrame()
    yoy = shift_period(period, -12)
    cur = shop_month[shop_month["period"] == period]
    ly = shop_month[shop_month["period"] == yoy]
    season = fit_seasonality(shop_month, period)
    shop_exp = fit_shop_expected(shop_month, period, season, pace)
    if (
        shop_exp is not None
        and not shop_exp.empty
        and "city" in shop_exp.columns
        and cities is not None
        and not cities.empty
        and "expected_mt" in cities.columns
    ):
        parents = cities.copy()
        if "grain_id" not in parents.columns and "city" in parents.columns:
            parents["grain_id"] = parents["city"]
        shop_exp = shop_exp.copy()
        shop_exp["parent_id"] = shop_exp["city"].astype(str)
        shop_exp = reconcile_expected(shop_exp, parents, intra_frac=pace)
    gaps = _shop_gaps(cur, ly, pace, shop_expected=shop_exp)
    if gaps is None or gaps.empty:
        return pd.DataFrame()
    exp = pd.to_numeric(gaps.get("expected_mt"), errors="coerce").fillna(0.0)
    vol = pd.to_numeric(gaps.get("volume_mt"), errors="coerce").fillna(0.0)
    ly_s = pd.to_numeric(gaps.get("ly_mt"), errors="coerce").fillna(0.0)
    gaps["share_expected_mt"] = exp
    gaps["competitive_mt"] = vol - exp
    k_shop = k_from_ly(gaps["ly_mt"], 0.05)
    size = ly_s.where(ly_s >= exp, exp)
    gaps["isolated_mt"] = [empirical_bayes(c, s, k_shop) for c, s in zip(gaps["competitive_mt"], size)]
    gaps["recoverable_mt"] = gaps["isolated_mt"].clip(upper=0).abs()
    gaps["store_name"] = gaps["store_name"].replace("", pd.NA).fillna(gaps["store_id"])
    return gaps.loc[gaps["recoverable_mt"] > 0].copy()


def visit_shop_floors(shop_month: pd.DataFrame, period: str) -> dict[str, float]:
    """Recoverable cut for shop lists. Size floors are not used — any door above 0.25 MT gap is listed."""
    del shop_month, period
    return {"size_floor": 0.0, "hole_floor": SHOP_RECOVERABLE_FLOOR, "rel": 0.0}


def keep_visit_shops(
    shops: pd.DataFrame,
    min_recoverable: float = SHOP_RECOVERABLE_FLOOR,
    floors: dict[str, float] | None = None,
    max_rows: int | None = None,
) -> tuple[pd.DataFrame, dict[str, float]]:
    """Keep every shop whose recoverable hole is above the cut; roll the rest into a remainder."""
    del max_rows
    if floors:
        min_recoverable = float(floors.get("hole_floor") or min_recoverable)
    empty_meta = {"n_kept": 0, "n_hidden": 0, "hidden_mt": 0.0, "n_pool": 0, "hole_floor": min_recoverable}
    if shops is None or shops.empty:
        return shops if shops is not None else pd.DataFrame(), empty_meta
    out = shops.copy()
    rec = pd.to_numeric(out.get("recoverable_mt"), errors="coerce").fillna(0.0)
    keep = rec > float(min_recoverable)
    kept = _sort_focus(out.loc[keep].copy())
    hidden = out.loc[~keep]
    meta = {
        "n_kept": int(len(kept)),
        "n_hidden": int(len(hidden)),
        "hidden_mt": float(pd.to_numeric(hidden.get("recoverable_mt"), errors="coerce").fillna(0).sum()) if not hidden.empty else 0.0,
        "n_pool": int(len(out)),
        "hole_floor": float(min_recoverable),
    }
    return kept, meta


def keep_top_holes(
    df: pd.DataFrame,
    n: int,
    *,
    value_col: str = "recoverable_mt",
    min_value: float = 0.0,
) -> tuple[pd.DataFrame, dict[str, float]]:
    """Keep the n largest lagging rows by Gap; roll the rest into a remainder."""
    empty_meta: dict[str, float] = {
        "n_kept": 0,
        "n_hidden": 0,
        "hidden_mt": 0.0,
        "n_pool": 0,
        "top_n": int(n),
    }
    if df is None or df.empty or int(n) <= 0:
        return (df if df is not None else pd.DataFrame()), empty_meta
    work = df.copy()
    val = pd.to_numeric(work.get(value_col), errors="coerce").fillna(0.0)
    holes = work.loc[val > float(min_value)].copy()
    if holes.empty:
        return work.iloc[0:0].copy(), {**empty_meta, "n_pool": int(len(work))}
    ranked = holes.assign(_gap=pd.to_numeric(holes[value_col], errors="coerce").fillna(0.0))
    ranked = ranked.sort_values("_gap", ascending=False, kind="mergesort")
    named = ranked.head(int(n)).drop(columns=["_gap"], errors="ignore")
    rest = ranked.iloc[int(n) :].drop(columns=["_gap"], errors="ignore")
    meta = {
        **empty_meta,
        "n_kept": int(len(named)),
        "n_hidden": int(len(rest)),
        "hidden_mt": float(pd.to_numeric(rest.get(value_col), errors="coerce").fillna(0).sum()) if not rest.empty else 0.0,
        "n_pool": int(len(holes)),
    }
    return named, meta


def _dsr_keys_under_distributors(
    shop_month: pd.DataFrame,
    dist_names: set[str],
    period: str,
    shops: pd.DataFrame | None = None,
) -> pd.DataFrame:
    """City + DSR pairs that appear on doors under the named distributors."""
    frames: list[pd.DataFrame] = []
    cols = ["city", "dsr_name"]
    if shop_month is not None and not shop_month.empty and dist_names:
        cur = shop_month
        if "period" in cur.columns:
            this = cur[cur["period"].astype(str) == str(period)]
            if not this.empty:
                cur = this
        if "distributor" in cur.columns:
            mask = cur["distributor"].astype(str).isin(dist_names)
            have = [c for c in cols if c in cur.columns]
            if have:
                frames.append(cur.loc[mask, have].copy())
    if shops is not None and not shops.empty and dist_names and "distributor" in shops.columns:
        mask = shops["distributor"].astype(str).isin(dist_names)
        have = [c for c in cols if c in shops.columns]
        if have:
            frames.append(shops.loc[mask, have].copy())
    if not frames:
        return pd.DataFrame(columns=cols)
    out = pd.concat(frames, ignore_index=True)
    for c in cols:
        if c not in out.columns:
            out[c] = ""
    return out[cols].drop_duplicates()


def _filter_dsrs_by_keys(dsrs: pd.DataFrame, keys: pd.DataFrame) -> pd.DataFrame:
    if dsrs is None or dsrs.empty or keys is None or keys.empty:
        return dsrs.iloc[0:0].copy() if dsrs is not None else pd.DataFrame()
    city_col = "city" if "city" in dsrs.columns else "parent_id"
    dsr_col = "dsr_name" if "dsr_name" in dsrs.columns else "grain_id"
    key_city = "city" if "city" in keys.columns else None
    key_dsr = "dsr_name" if "dsr_name" in keys.columns else None
    if not key_city or not key_dsr:
        return dsrs.iloc[0:0].copy()
    left = dsrs[city_col].astype(str) + "\0" + dsrs[dsr_col].astype(str)
    right = set(keys[key_city].astype(str) + "\0" + keys[key_dsr].astype(str))
    return dsrs.loc[left.isin(right)].copy()


def _merge_hole_meta(floor_meta: dict[str, float], vital_meta: dict[str, float]) -> dict[str, float]:
    out = dict(floor_meta or {})
    out["n_kept"] = int(vital_meta.get("n_kept") or 0)
    out["n_hidden"] = int(floor_meta.get("n_hidden") or 0) + int(vital_meta.get("n_hidden") or 0)
    out["hidden_mt"] = float(floor_meta.get("hidden_mt") or 0) + float(vital_meta.get("hidden_mt") or 0)
    if vital_meta.get("top_n") is not None:
        out["top_n"] = vital_meta["top_n"]
    return out


def _present_shops(df: pd.DataFrame, meta: dict[str, float]) -> pd.DataFrame:
    table = _present(df, SHOP_VIEW)
    return _append_remainder(table, SHOP_VIEW, meta, "Shop", "doors")


def _present_with_remainder(
    df: pd.DataFrame,
    view: list[tuple[str, str]],
    meta: dict[str, float] | None,
    name_field: str,
    noun: str,
) -> pd.DataFrame:
    table = _present(df, view)
    return _append_remainder(table, view, meta or {}, name_field, noun)


def _append_remainder(
    table: pd.DataFrame,
    view: list[tuple[str, str]],
    meta: dict[str, float],
    name_field: str,
    noun: str,
) -> pd.DataFrame:
    n_hidden = int(meta.get("n_hidden") or 0)
    hidden_mt = float(meta.get("hidden_mt") or 0)
    if n_hidden <= 0:
        return table
    rest = {label: None for _, label in view}
    top_n = meta.get("top_n")
    if top_n:
        rest[name_field] = (
            f"Not listed — {n_hidden} {noun} totalling {hidden_mt:.0f} MT gap. "
            f"Ranked by Gap; only the top {int(top_n)} are named on this sheet."
        )
    elif name_field == "Shop":
        rest[name_field] = (
            f"Not listed — {n_hidden} {noun} totalling {hidden_mt:.0f} MT gap. "
            "Gaps of 0.25 MT or less, not a first call."
        )
    else:
        rest[name_field] = (
            f"Not listed — {n_hidden} {noun} totalling {hidden_mt:.0f} MT gap. "
            "The rest of the hole after the named rows."
        )
    rest["Gap (MT)"] = _round_num(hidden_mt)
    extra = pd.DataFrame([rest])
    if table is None or table.empty:
        return extra
    return pd.concat([table, extra], ignore_index=True)


def ams_last_n(
    shop_month: pd.DataFrame,
    period: str,
    keys: list[str],
    n: int = 3,
    ledger: pd.DataFrame | None = None,
) -> pd.DataFrame:
    """Mean volume of the last n closed months, excluding the period being scored."""
    cols = keys + ["ams_3m"]
    if shop_month is None or shop_month.empty or not period:
        return pd.DataFrame(columns=cols)
    hist = shop_month[shop_month["period"].astype(str) != str(period)].copy()
    if hist.empty:
        return pd.DataFrame(columns=cols)
    if ledger is not None and not ledger.empty and "status" in ledger.columns:
        closed = set(ledger.loc[ledger["status"].astype(str) == "closed", "period"].astype(str))
        if closed:
            hist = hist[hist["period"].astype(str).isin(closed)]
    if hist.empty:
        return pd.DataFrame(columns=cols)
    for k in keys:
        if k not in hist.columns:
            hist[k] = "(unmapped)"
        hist[k] = hist[k].fillna("(unmapped)").replace("", "(unmapped)")
    periods = sorted(hist["period"].astype(str).unique())[-n:]
    hist = hist[hist["period"].astype(str).isin(periods)]
    g = hist.groupby(keys + ["period"], as_index=False)["volume_mt"].sum()
    out = g.groupby(keys, as_index=False)["volume_mt"].mean().rename(columns={"volume_mt": "ams_3m"})
    return out


def excel_bytes(pack: StrategyPack) -> bytes:
    buf = BytesIO()
    write_excel(pack, buf)
    return buf.getvalue()


def excel_bytes_detailed(pack: StrategyPack) -> bytes:
    buf = BytesIO()
    write_excel_detailed(pack, buf)
    return buf.getvalue()


def pdf_bytes(pack: StrategyPack, detailed: bool = False) -> bytes:
    from sndintel.pdf import render_pdf

    return render_pdf(pack, detailed=detailed)


def pdf_bytes_detailed(pack: StrategyPack) -> bytes:
    return pdf_bytes(pack, detailed=True)


def html_bytes(pack: StrategyPack) -> bytes:
    return render_html(pack).encode("utf-8")


def html_bytes_detailed(pack: StrategyPack) -> bytes:
    return render_html(pack, detailed=True).encode("utf-8")


def write_excel(pack: StrategyPack, path: Path | str | BytesIO) -> None:
    wb = Workbook()
    _sheet_cover(wb, pack)
    for i, (sheet, heading, note, df) in enumerate(iter_report_sheets(pack, detailed=False)):
        kwargs: dict[str, Any] = {}
        if i == 0 and df is not None and not df.empty and "Gap (MT)" in df.columns:
            cat = "City" if "City" in df.columns else list(df.columns)[0]
            kwargs = dict(bar_col="Gap (MT)", cat_col=cat, freeze="A2")
        _sheet_table(wb, sheet, heading, note, df, **kwargs)
    if path is not None:
        wb.save(path)


def write_excel_detailed(pack: StrategyPack, path: Path | str | BytesIO) -> None:
    wb = Workbook()
    _sheet_cover(wb, pack, detailed=True)
    for sheet, heading, note, df in iter_report_sheets(pack, detailed=True):
        kwargs: dict[str, Any] = {}
        if sheet.startswith("01") and df is not None and not df.empty and "Gap (MT)" in df.columns:
            cat = "City" if "City" in df.columns else list(df.columns)[0]
            kwargs = dict(bar_col="Gap (MT)", cat_col=cat, freeze="A2")
        _sheet_table(wb, sheet, heading, note, df, **kwargs)
    if path is not None:
        wb.save(path)


def iter_report_sheets(pack: StrategyPack, detailed: bool = False) -> list[tuple[str, str, str, pd.DataFrame]]:
    """Ordered (sheet, heading, note, table) for Excel and PDF."""
    scope = (pack.scope or "national").lower()
    label = pack.scope_label or scope
    if scope == "city":
        return [
            (
                "01 City",
                f"{label} versus its Expected",
                "Country row is first when included. Gap is billed versus this unit’s own Expected. From drop / unvisited / unbilled add to Gap.",
                pack.cities,
            ),
            (
                "02 Distributors",
                f"Distributors in {label}",
                "AMS = 0 is hidden. Highest gap first. This city pack lists every distributor with a run-rate, not only the national top 10.",
                pack.all_distributors,
            ),
            (
                "03 DSRs",
                f"DSRs in {label}",
                "Salespeople in this city with AMS greater than 0.",
                pack.all_dsrs,
            ),
            (
                "04 Shops",
                f"Shops in {label}",
                "Every shop with gap greater than 0.25 MT. Shallower doors are one remainder line.",
                pack.all_shops,
            ),
        ]
    if scope == "distributor":
        dist_tbl = pack.all_distributors if pack.all_distributors is not None and not pack.all_distributors.empty else pack.city_distributors
        return [
            (
                "01 Distributor",
                f"{label}",
                "Scorecard versus its own Expected. From drop / unvisited / unbilled add to Gap.",
                dist_tbl,
            ),
            (
                "02 DSRs",
                f"Salespeople under {label}",
                "DSRs on shops billed or listed under this distributor.",
                pack.all_dsrs,
            ),
            (
                "03 Shops",
                f"Shops under {label}",
                "Every shop with gap greater than 0.25 MT. Shallower doors are one remainder line.",
                pack.all_shops,
            ),
        ]
    if scope == "dsr":
        return [
            (
                "01 DSR",
                f"{label}",
                "Scorecard versus its own Expected. From drop / unvisited / unbilled add to Gap.",
                pack.all_dsrs,
            ),
            (
                "02 Shops",
                f"Shops on this beat",
                "Every shop with gap greater than 0.25 MT. Shallower doors are one remainder line.",
                pack.all_shops,
            ),
        ]
    if detailed:
        return [
            (
                "01 City detail",
                "Every city",
                "Full city list, highest gap first. From drop / unvisited / unbilled add to Gap. Strike % = billed shops ÷ universe.",
                pack.cities,
            ),
            (
                "02 Distributor detail",
                "Every distributor with AMS greater than 0",
                "Not just lagging distributors. AMS = 0 is hidden. Sorted highest gap first.",
                pack.all_distributors,
            ),
            (
                "03 DSR detail",
                "Every DSR with AMS greater than 0",
                "Not just lagging DSRs. AMS = 0 is hidden. Sorted highest gap first.",
                pack.all_dsrs,
            ),
            (
                "04 National DSRs",
                "National DSR list",
                "Same as DSR detail — every salesperson with a recent run-rate (AMS > 0).",
                pack.all_dsrs,
            ),
            (
                "05 National shops",
                "National shop list",
                "Every shop with gap greater than 0.25 MT. Shallower doors are one remainder line.",
                pack.all_shops,
            ),
        ]
    return [
        (
            "01 Country by city",
            "Every city versus its own Expected",
            "Gap is billed versus this unit’s own Expected. From drop size / unvisited / unbilled add to Gap (positive = hole; negative = billed more than Expected). Country row is first. Distributors and DSRs with AMS = 0 are hidden.",
            pack.cities,
        ),
        (
            "02 Top 10 distributors",
            "Top 10 lagging distributors by Gap",
            "Biggest Gap first. AMS = 0 is hidden. Remainder line is everyone after the tenth.",
            pack.city_distributors,
        ),
        (
            "03 Top 50 shops of those",
            "Top 50 lagging shops under those distributors",
            (pack.shop_note or "Top 50 shops by Gap under the distributors above.")
            + " Gap is volume that comes back if the door billed its own Expected.",
            pack.city_distributor_shops,
        ),
        (
            "04 Top 10 dists all cities",
            "Top 10 lagging distributors (all cities)",
            "Same ten names as sheet 02 — national rank by Gap, including pockets inside cities that are on expected.",
            pack.lagging_distributors,
        ),
        (
            "05 Top 10 DSRs of those",
            "Top 10 lagging DSRs from those distributors",
            "Salespeople whose shops sit under the top 10 distributors, then the ten biggest Gaps. Ride-with this list.",
            pack.lagging_dsrs,
        ),
        (
            "06 Top 50 shops",
            "Top 50 lagging shops",
            pack.shop_note or "Top 50 doors by Gap after the 0.25 MT floor. The rest of the hole is the remainder line.",
            pack.lagging_shops,
        ),
    ]


def render_html(pack: StrategyPack, detailed: bool = False) -> str:
    k = pack.kpis
    sections = [_html_glossary()]
    if is_national_pack(pack):
        sections.append(_html_exec(pack))
    else:
        sections.append(_html_cover(pack, k, detailed=detailed))
    sections.extend(
        [
        _html_section(
            "1. The country — every city",
            "Gap is billed versus this unit’s own Expected, highest first. From drop size / unvisited / unbilled add to Gap (positive = hole; negative = billed more than Expected). Strike % = billed ÷ universe. Visit % = visited ÷ universe. Country row is first. Remarks are the last column.",
            pack.cities,
        ),
        _html_section(
            "2. Top 10 lagging distributors",
            "Biggest Gap first. AMS = 0 is hidden. Remainder line is everyone after the tenth.",
            pack.city_distributors,
        ),
        _html_section(
            "3. Top 50 shops of those distributors",
            "Top 50 doors by Gap under the distributors above. Remainder line is the tail of that hole.",
            pack.city_distributor_shops,
        ),
        _html_section(
            "4. Top 10 lagging distributors (all cities)",
            "Same ten names as section 2 — national rank by Gap.",
            pack.lagging_distributors,
        ),
        _html_section(
            "5. Top 10 lagging DSRs from those distributors",
            "Salespeople whose shops sit under the top 10 distributors, then the ten biggest Gaps. Ride-with this list.",
            pack.lagging_dsrs,
        ),
        _html_section(
            "6. Top 50 lagging shops",
            pack.shop_note or "Top 50 doors by Gap after the 0.25 MT floor. The rest of the hole is the remainder line.",
            pack.lagging_shops,
        ),
        ]
    )
    if detailed:
        sections.extend(
            [
                _html_section(
                    "City detail — every city",
                    "Full city list with billed shops, strike %, coverage and drop-size split.",
                    pack.cities,
                ),
                _html_section(
                    "Distributor detail — every distributor with AMS > 0",
                    "Not just lagging. Sorted highest gap first.",
                    pack.all_distributors,
                ),
                _html_section(
                    "DSR detail — every DSR with AMS > 0",
                    "Not just lagging. Sorted highest gap first.",
                    pack.all_dsrs,
                ),
                _html_section(
                    "National detail — every DSR with AMS > 0",
                    "Full salesperson list for the country.",
                    pack.all_dsrs,
                ),
                _html_section(
                    "National detail — every shop above 0.25 MT gap",
                    "Every shop with gap greater than 0.25 MT. Shallower doors are one remainder line.",
                    pack.all_shops,
                ),
            ]
        )
    return f"""<!DOCTYPE html>
<html lang="en"><head>
<meta charset="utf-8"/>
<title>SND strategy pack · {html.escape(pack.label)}</title>
<style>
  :root {{ --ink:#0f172a; --muted:#475569; --line:#cbd5e1; --red:#b91c1c; --green:#15803d; --wash:#f8fafc; }}
  body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Helvetica, Arial, sans-serif;
         color: var(--ink); margin: 24px; background: white; }}
  h1 {{ font-size: 22px; margin: 0 0 4px; }}
  h2 {{ font-size: 16px; margin: 28px 0 6px; page-break-after: avoid; }}
  .kicker {{ letter-spacing: .12em; text-transform: uppercase; font-size: 11px; font-weight: 700; color: var(--muted); }}
  .headline {{ font-size: 20px; font-weight: 700; margin: 8px 0 10px; }}
  .lead {{ color: var(--muted); line-height: 1.45; margin: 0 0 8px; max-width: 920px; }}
  .kpis {{ display: flex; flex-wrap: wrap; gap: 10px; margin: 14px 0 8px; }}
  .kpi {{ background: var(--wash); border: 1px solid var(--line); border-radius: 10px; padding: 10px 14px; min-width: 120px; }}
  .kpi b {{ display: block; font-size: 18px; }}
  .kpi span {{ font-size: 11px; color: var(--muted); }}
  table {{ border-collapse: collapse; width: 100%; font-size: 12px; margin: 8px 0 4px; }}
  th {{ background: var(--ink); color: white; text-align: left; padding: 6px 8px; font-weight: 600; }}
  td {{ border-bottom: 1px solid var(--line); padding: 5px 8px; vertical-align: top; }}
  td.num {{ text-align: right; font-variant-numeric: tabular-nums; }}
  tr.lagging td {{ background: #fef2f2; }}
  tr.country td {{ background: #e2e8f0; font-weight: 600; }}
  .note {{ font-size: 12px; color: var(--muted); margin: 0 0 8px; }}
  dl.glossary {{ font-size: 12px; line-height: 1.4; }}
  dl.glossary dt {{ font-weight: 700; margin-top: 8px; }}
  dl.glossary dd {{ margin: 2px 0 0 0; color: var(--muted); }}
  @page {{ size: A4 landscape; margin: 12mm; }}
  @media print {{
    body {{ margin: 0; }}
    section {{ page-break-before: always; }}
    section.cover {{ page-break-before: auto; }}
    h2 {{ page-break-after: avoid; }}
    tr {{ page-break-inside: avoid; }}
  }}
</style></head><body>
{''.join(sections)}
</body></html>
"""


# ----- internals -----

def _grain(units: pd.DataFrame, grain: str) -> pd.DataFrame:
    out = units[units["grain"] == grain].copy()
    if out.empty:
        return out
    out["recoverable_mt"] = pd.to_numeric(out.get("isolated_mt"), errors="coerce").fillna(0).clip(upper=0).abs()
    drop = pd.to_numeric(out.get("from_drop_size_mt"), errors="coerce")
    if drop.isna().all():
        drop = pd.to_numeric(out.get("velocity_effect_mt"), errors="coerce")
    out["from_drop_size_mt"] = drop
    out["from_unvisited_mt"] = pd.to_numeric(out.get("from_unvisited_mt"), errors="coerce")
    out["from_unbilled_mt"] = pd.to_numeric(out.get("from_unbilled_mt"), errors="coerce")
    sit = out["situation"] if "situation" in out.columns else pd.Series("", index=out.index)
    out["situation_label"] = sit.map(SITUATION_LABEL).fillna(sit)
    strike = pd.to_numeric(out.get("strike_rate"), errors="coerce")
    out["strike_pct"] = strike * 100
    visit = pd.to_numeric(out.get("visit_rate"), errors="coerce")
    out["visit_pct"] = visit * 100
    vol = pd.to_numeric(out.get("volume_mt"), errors="coerce")
    if "billed" in out.columns:
        billed = pd.to_numeric(out["billed"], errors="coerce")
        out["drop_size_mt"] = vol / billed.mask(billed <= 0)
    else:
        out["drop_size_mt"] = pd.NA
    return out


def _attach_ams(
    df: pd.DataFrame,
    shop_month: pd.DataFrame,
    period: str,
    keys: list[str],
    ledger: pd.DataFrame | None,
    pace: float,
) -> pd.DataFrame:
    out = df.copy()
    ams = ams_last_n(shop_month, period, keys, n=3, ledger=ledger)
    if ams.empty:
        out["ams_3m"] = pd.NA
        out["vs_ams_mt"] = pd.NA
        return out
    left = out.copy()
    for k in keys:
        if k not in left.columns:
            left[k] = "(unmapped)"
        left[k] = left[k].fillna("(unmapped)").astype(str).replace("", "(unmapped)")
    right = ams.copy()
    for k in keys:
        right[k] = right[k].fillna("(unmapped)").astype(str)
    left["_row"] = range(len(left))
    merged = left.merge(right, on=keys, how="left")
    merged = merged.sort_values("_row").drop(columns=["_row"])
    ams_v = pd.to_numeric(merged["ams_3m"], errors="coerce")
    vol = pd.to_numeric(merged.get("volume_mt"), errors="coerce")
    merged["vs_ams_mt"] = vol - ams_v * float(pace or 1.0)
    return merged


def _drop_zero_ams(df: pd.DataFrame) -> pd.DataFrame:
    """Hide distributors / DSRs with no recent run-rate (AMS missing or 0)."""
    if df is None or df.empty:
        return df
    if "ams_3m" not in df.columns:
        return df
    ams = pd.to_numeric(df["ams_3m"], errors="coerce").fillna(0.0)
    return df.loc[ams > 0].copy()


def _sort_focus(df: pd.DataFrame) -> pd.DataFrame:
    if df is None or df.empty:
        return df
    out = df.copy()
    if "recoverable_mt" in out.columns:
        return out.sort_values(["recoverable_mt", "isolated_mt"], ascending=[False, True])
    if "isolated_mt" in out.columns:
        return out.sort_values("isolated_mt")
    return out


def _present(df: pd.DataFrame, view: list[tuple[str, str]]) -> pd.DataFrame:
    if df is None or df.empty:
        return pd.DataFrame(columns=[label for _, label in view])
    cols = []
    data = {}
    for src, label in view:
        cols.append(label)
        data[label] = df[src] if src in df.columns else pd.NA
    out = pd.DataFrame(data)[cols]
    return _round_display(out).reset_index(drop=True)


def _round_num(val: Any) -> Any:
    if val is None:
        return pd.NA
    try:
        if pd.isna(val):
            return pd.NA
    except (TypeError, ValueError):
        return val
    try:
        return int(round(float(val)))
    except (TypeError, ValueError):
        return val


def _round_display(df: pd.DataFrame) -> pd.DataFrame:
    """Whole numbers for MT, counts, and percents. From-columns still sum to Gap."""
    if df is None or df.empty:
        return df
    out = df.copy()
    rec_col = "Gap (MT)"
    from_cols = [
        c
        for c in ["From drop size (MT)", "From unvisited shops (MT)", "From unbilled shops (MT)"]
        if c in out.columns
    ]
    count_cols = {"Billed shops", "Visited shops", "Universe", "Visits MTD"}
    for col in out.columns:
        if col == "Remarks":
            continue
        name = str(col)
        if name == "Drop size (MT)":
            out[col] = [
                (round(float(v), 2) if v is not None and pd.notna(v) else pd.NA) for v in out[col]
            ]
            continue
        if "(MT)" in name or name.endswith("%") or name in count_cols:
            out[col] = [_round_num(v) for v in out[col]]
    if rec_col in out.columns and len(from_cols) == 3:
        for i in out.index:
            rec = out.loc[i, rec_col]
            rec_i = int(rec) if rec is not None and pd.notna(rec) else 0
            parts = []
            for c in from_cols:
                v = out.loc[i, c]
                parts.append(int(v) if v is not None and pd.notna(v) else 0)
            if rec_i != 0:
                target = rec_i
            else:
                target = int(round(sum(parts)))
            diff = target - sum(parts)
            if diff:
                j = max(range(len(parts)), key=lambda k: abs(parts[k]))
                parts[j] += diff
                out.loc[i, from_cols[j]] = parts[j]
    return out


def _present_with_national(nat: pd.DataFrame, cities: pd.DataFrame, view: list[tuple[str, str]]) -> pd.DataFrame:
    city_tbl = _present(cities, view)
    if nat is None or nat.empty:
        return city_tbl
    nat_tbl = _present(nat, view)
    return pd.concat([nat_tbl, city_tbl], ignore_index=True)


def _attach_national_ams(
    nat: pd.DataFrame,
    shop_month: pd.DataFrame,
    period: str,
    ledger: pd.DataFrame | None,
    pace: float,
) -> pd.DataFrame:
    out = nat.copy()
    if shop_month is None or shop_month.empty or not period:
        out["ams_3m"] = pd.NA
        out["vs_ams_mt"] = pd.NA
        return out
    hist = shop_month[shop_month["period"].astype(str) != str(period)].copy()
    if ledger is not None and not ledger.empty and "status" in ledger.columns:
        closed = set(ledger.loc[ledger["status"].astype(str) == "closed", "period"].astype(str))
        if closed:
            hist = hist[hist["period"].astype(str).isin(closed)]
    if hist.empty:
        out["ams_3m"] = pd.NA
        out["vs_ams_mt"] = pd.NA
        return out
    g = hist.groupby("period")["volume_mt"].sum().sort_index().tail(3)
    ams = float(g.mean()) if len(g) else None
    out["ams_3m"] = ams
    vol = pd.to_numeric(out.get("volume_mt"), errors="coerce")
    out["vs_ams_mt"] = vol - (ams * float(pace or 1.0) if ams is not None else 0.0)
    return out


def _parent_stats(row: pd.Series) -> dict[str, Any]:
    vol = pd.to_numeric(row.get("volume_mt"), errors="coerce")
    ly = pd.to_numeric(row.get("ly_mt"), errors="coerce")
    yoy = None
    if pd.notna(ly) and float(ly) > 1e-9 and pd.notna(vol):
        yoy = 100.0 * (float(vol) - float(ly)) / float(ly)
    billed = pd.to_numeric(row.get("billed"), errors="coerce")
    drop = pd.to_numeric(row.get("drop_size_mt"), errors="coerce")
    if (drop is None or pd.isna(drop)) and pd.notna(vol) and pd.notna(billed) and float(billed) > 0:
        drop = float(vol) / float(billed)
    return {
        "visit_rate": pd.to_numeric(row.get("visit_rate"), errors="coerce"),
        "productivity": pd.to_numeric(row.get("productivity"), errors="coerce"),
        "yoy_pct": yoy,
        "drop_size_mt": float(drop) if drop is not None and pd.notna(drop) else None,
    }


def _attach_shop_calls(shops: pd.DataFrame, visits: pd.DataFrame | None, period: str) -> pd.DataFrame:
    if shops is None or shops.empty:
        return shops
    out = shops.copy()
    out["visits"] = 0
    billed = pd.to_numeric(out.get("volume_mt"), errors="coerce").fillna(0) > 0
    if visits is not None and not visits.empty and "store_id" in visits.columns:
        v = visits.copy()
        v["store_id"] = v["store_id"].astype(str)
        if "period" in v.columns:
            v = v[v["period"].astype(str) == str(period)]
        if not v.empty:
            g = v.groupby("store_id", as_index=False)["visits"].sum()
            out["store_id"] = out["store_id"].astype(str)
            out = out.merge(g, on="store_id", how="left", suffixes=("", "_v"))
            if "visits_v" in out.columns:
                out["visits"] = pd.to_numeric(out["visits_v"], errors="coerce").fillna(0)
                out = out.drop(columns=["visits_v"])
            else:
                out["visits"] = pd.to_numeric(out.get("visits"), errors="coerce").fillna(0)
    visited = (pd.to_numeric(out["visits"], errors="coerce").fillna(0) > 0) | billed
    out["call_status"] = pd.Series("Unvisited", index=out.index)
    out.loc[visited & ~billed, "call_status"] = "Visited · not billed"
    out.loc[billed, "call_status"] = "Billed"
    return out


def _situation_row(situation: pd.DataFrame | None, national: dict | None, nat_row) -> dict[str, Any]:
    out: dict[str, Any] = {}
    if national:
        out.update(national)
    if nat_row is not None:
        out.setdefault("label", None)
    if situation is not None and not situation.empty:
        row = situation.iloc[0]
        out["headline"] = row.get("headline") or out.get("headline")
        out["weather"] = row.get("weather") or out.get("weather")
        out["problem"] = row.get("problem") or out.get("problem")
        out["action"] = row.get("action_summary") or out.get("action")
        try:
            import json as _json
            meta = _json.loads(row["metrics_json"]) if row.get("metrics_json") else {}
            out.update({k: v for k, v in meta.items() if v is not None})
        except Exception:
            pass
    return out


def _kpis(nat_row, units: pd.DataFrame, sit: dict[str, Any]) -> dict[str, Any]:
    cities = units[units["grain"] == "city"] if units is not None else pd.DataFrame()
    n_lag = int((cities["situation"] == "lagging").sum()) if not cities.empty and "situation" in cities.columns else 0
    extra = sit.get("extra_hole_mt")
    if extra is None and not cities.empty and "isolated_mt" in cities.columns:
        extra = float(cities.loc[cities.get("situation", "") == "lagging", "isolated_mt"].sum()) if "situation" in cities.columns else 0.0
    billed = float(nat_row["volume_mt"]) if nat_row is not None else None
    expected = float(nat_row["expected_mt"]) if nat_row is not None else None
    gap = float(nat_row["gap_mt"]) if nat_row is not None else None
    ly = float(nat_row["ly_mt"]) if nat_row is not None else None
    return {
        "billed_mt": billed,
        "expected_mt": expected,
        "gap_mt": gap,
        "ly_mt": ly,
        "extra_hole_mt": float(extra or 0),
        "n_lagging_cities": n_lag,
        "n_history_periods": sit.get("n_history_periods"),
        "n_same_month": sit.get("n_same_month"),
        "intra_month_frac": sit.get("intra_month_frac") or (nat_row.get("intra_month_frac") if nat_row is not None else None),
        "open_mtd": sit.get("open_mtd"),
    }


NAVY = "0F172A"
RED = "B91C1C"
GREEN = "15803D"
SLATE = "64748B"
WASH = "F8FAFC"
WHITE = "FFFFFF"
THIN = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _sheet_cover(wb: Workbook, pack: StrategyPack, detailed: bool = False) -> Worksheet:
    ws = wb.active
    ws.title = "00 Cover"
    ws.sheet_properties.tabColor = NAVY
    ws["A1"] = "SND Intelligence"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=NAVY)
    ws["A2"] = f"{'Detailed pack' if detailed else 'Strategy pack'} · {pack.label}"
    ws["A2"].font = Font(name="Calibri", size=18, bold=True, color=NAVY)
    ws["A3"] = "Glossary first, then the national executive summary (national packs only), then the tables."
    ws["A3"].font = Font(name="Calibri", size=10, italic=True, color=SLATE)
    ws.merge_cells("A3:H3")

    row = 5
    ws.cell(row, 1, "Glossary")
    ws.cell(row, 1).font = Font(bold=True, size=12, color=NAVY)
    row += 1
    ws.cell(row, 1, "Term").font = Font(bold=True, color=WHITE)
    ws.cell(row, 2, "Meaning").font = Font(bold=True, color=WHITE)
    ws.cell(row, 1).fill = _fill(NAVY)
    ws.cell(row, 2).fill = _fill(NAVY)
    row += 1
    for term, meaning in GLOSSARY:
        ws.cell(row, 1, term).font = Font(bold=True)
        ws.cell(row, 2, meaning)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        ws.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 32
        row += 1

    row += 1
    ws.cell(row, 1, "How the figures are calculated")
    ws.cell(row, 1).font = Font(bold=True, size=12, color=NAVY)
    row += 1
    for term, meaning in CALCULATION_NOTES:
        ws.cell(row, 1, term).font = Font(bold=True)
        ws.cell(row, 2, meaning)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        ws.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 48
        row += 1

    if is_national_pack(pack):
        row += 1
        ws.cell(row, 1, "Executive summary")
        ws.cell(row, 1).font = Font(bold=True, size=12, color=NAVY)
        row += 1
        if pack.exec_model and pack.exec_situation:
            ws.cell(
                row,
                1,
                f"Written from this period’s scorecards ({pack.exec_model}). Figures match the tables.",
            )
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
            ws.cell(row, 1).font = Font(italic=True, color=SLATE, size=10)
            row += 1
        if pack.exec_situation:
            ws.cell(row, 1, "Summary of current situation")
            ws.cell(row, 1).font = Font(bold=True, color=NAVY)
            row += 1
            for para in pack.exec_situation:
                ws.cell(row, 1, para)
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
                ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
                ws.row_dimensions[row].height = 56
                row += 1
            ws.cell(row, 1, "Key focus areas")
            ws.cell(row, 1).font = Font(bold=True, color=NAVY)
            row += 1
            for i, item in enumerate(pack.exec_focus, start=1):
                title = item.get("title") or f"Focus {i}"
                why = item.get("why") or ""
                do = item.get("do") or ""
                text = f"{i}. {title}"
                if why:
                    text += f" — {why}"
                if do:
                    text += f" Do this week: {do}"
                ws.cell(row, 1, text)
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
                ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
                ws.row_dimensions[row].height = 48
                row += 1
        elif pack.exec_error:
            ws.cell(row, 1, f"Executive summary not generated. {pack.exec_error}")
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
            ws.cell(row, 1).alignment = Alignment(wrap_text=True)
            row += 1
        else:
            ws.cell(
                row,
                1,
                "No national executive summary stored. Paste an OpenAI key on Upload files and rebuild scorecards.",
            )
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
            row += 1

    row += 1
    ws.cell(row, 1, "How to read the tables")
    ws.cell(row, 1).font = Font(bold=True, size=12, color=NAVY)
    row += 1
    for i, line in enumerate(how_to_read_steps(pack, detailed=detailed), start=1):
        ws.cell(row, 1, f"{i:02d}. {line}")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        ws.cell(row, 1).alignment = Alignment(wrap_text=True)
        row += 1

    ws.column_dimensions["A"].width = 36
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 18
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:2"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    return ws


def _sheet_table(
    wb: Workbook,
    title: str,
    heading: str,
    note: str,
    df: pd.DataFrame,
    freeze: str = "A5",
    bar_col: str | None = None,
    cat_col: str | None = None,
) -> Worksheet:
    ws = wb.create_sheet(title[:31])
    ws.sheet_properties.tabColor = RED if "lagging" in title.lower() or "dists" in title.lower() or "shops" in title.lower() or "DSR" in title else NAVY
    ws["A1"] = heading
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=NAVY)
    ws.merge_cells("A1:L1")
    ws["A2"] = note
    ws["A2"].font = Font(size=10, italic=True, color=SLATE)
    ws.merge_cells("A2:L2")
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[2].height = 32
    if df is None or df.empty:
        ws["A4"] = "No rows at this layer for this period."
        ws.column_dimensions["A"].width = 40
        return ws
    start = 4
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=start):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(r_idx, c_idx, _excel_value(value))
            cell.border = THIN
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if r_idx == start:
                cell.fill = _fill(NAVY)
                cell.font = Font(bold=True, color=WHITE, size=10)
            else:
                cell.font = Font(size=10)
                _format_metric_cell(cell, df.columns[c_idx - 1] if c_idx - 1 < len(df.columns) else "")
                sit = row_tone(df.iloc[r_idx - start - 1])
                if sit == "lagging":
                    cell.fill = PatternFill("solid", fgColor="FEF2F2")
                elif sit == "ahead":
                    cell.fill = PatternFill("solid", fgColor="F0FDF4")
                elif sit == "country":
                    cell.fill = PatternFill("solid", fgColor="E2E8F0")
                    cell.font = Font(size=10, bold=True)
    headers = list(df.columns)
    for i, name in enumerate(headers, start=1):
        width = min(max(len(str(name)) + 2, 12), 28)
        if name in {"What to do", "Shop", "Remarks"}:
            width = 56 if name == "Remarks" else 48
        ws.column_dimensions[get_column_letter(i)].width = width
    if "Remarks" in headers:
        for r in range(start + 1, start + 1 + len(df)):
            ws.row_dimensions[r].height = 68
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(headers))}{start + len(df)}"
    ws.freeze_panes = f"A{start + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.print_title_rows = f"{start}:{start}"
    ws.page_setup.horizontalCentered = True
    if bar_col and cat_col and bar_col in df.columns and cat_col in df.columns and len(df) > 0:
        cat_idx = headers.index(cat_col) + 1
        bar_idx = headers.index(bar_col) + 1
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = bar_col
        chart.y_axis.title = None
        data = Reference(ws, min_col=bar_idx, min_row=start, max_row=start + len(df))
        cats = Reference(ws, min_col=cat_idx, min_row=start + 1, max_row=start + len(df))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.shape = 4
        chart.height = 8
        chart.width = 15
        ws.add_chart(chart, "Q4")
    return ws


def _excel_value(value: Any) -> Any:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if hasattr(value, "item") and not isinstance(value, (bytes, str)):
        try:
            return _excel_value(value.item())
        except Exception:
            return str(value)
    return value


def _format_metric_cell(cell, header: str) -> None:
    h = str(header)
    if h == "Drop size (MT)":
        cell.number_format = "0.00"
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="right")
        return
    if "(MT)" in h:
        cell.number_format = (
            "+0;-0;0"
            if h.startswith("Extra") or h.startswith("vs ") or h.startswith("From ") or "Gap" in h
            else "#,##0"
        )
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="right")
    elif h.endswith("%") or "Strike" in h:
        cell.number_format = "0"
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="right")
    elif h in {"Billed shops", "Visited shops", "Universe", "Visits MTD"}:
        cell.number_format = "#,##0"
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="right")


def row_tone(row: pd.Series) -> str:
    """Row wash: country / lagging (gap > 0) / ahead (billed > expected)."""
    city = str(row.get("City") or "")
    dist = str(row.get("Distributor") or "")
    if city == "Country" or dist == "Country":
        return "country"
    gap = row.get("Gap (MT)")
    billed = row.get("Billed this period (MT)")
    exp = row.get("Expected this month (MT)")
    try:
        if gap is not None and pd.notna(gap) and float(gap) > 0:
            return "lagging"
    except (TypeError, ValueError):
        pass
    try:
        if (
            billed is not None
            and exp is not None
            and pd.notna(billed)
            and pd.notna(exp)
            and float(billed) > float(exp) + 0.5
        ):
            return "ahead"
    except (TypeError, ValueError):
        pass
    return ""


def _html_cover(pack: StrategyPack, k: dict[str, Any], detailed: bool = False) -> str:
    del k
    ol = "".join(f"<li>{html.escape(s)}</li>" for s in how_to_read_steps(pack, detailed=detailed))
    return f"""<section class="cover">
  <div class="kicker">SND Intelligence · {html.escape(pack.scope or "report")} pack</div>
  <h1>{html.escape(pack.scope_label or pack.label)}</h1>
  <p class="headline">{html.escape(pack.headline or "Scorecards ready")}</p>
  <p class="lead">{html.escape(pack.weather or "")}</p>
  <p class="note">How to read this pack</p>
  <ol class="note">{ol}</ol>
</section>
"""


def _html_exec(pack: StrategyPack) -> str:
    bits = [
        '<section class="cover">',
        '<div class="kicker">Executive summary · national</div>',
        f"<h1>{html.escape(pack.label)}</h1>",
        "<h2>Executive summary</h2>",
    ]
    if pack.exec_model and pack.exec_situation:
        bits.append(
            f'<p class="note">Written from this period’s scorecards ({html.escape(pack.exec_model)}). Figures match the tables.</p>'
        )
    if pack.exec_situation:
        bits.append("<h2>Summary of current situation</h2>")
        for para in pack.exec_situation:
            bits.append(f'<p class="lead">{html.escape(para)}</p>')
        bits.append("<h2>Key focus areas</h2>")
        items = []
        for i, item in enumerate(pack.exec_focus, start=1):
            title = html.escape(item.get("title") or f"Focus {i}")
            why = html.escape(item.get("why") or "")
            do = html.escape(item.get("do") or "")
            body = f"<b>{i}. {title}</b>"
            if why:
                body += f" — {why}"
            if do:
                body += f" <b>Do this week.</b> {do}"
            items.append(f"<li>{body}</li>")
        bits.append(f'<ol class="note">{"".join(items)}</ol>')
    elif pack.exec_error:
        bits.append(f'<p class="lead">Executive summary not generated. {html.escape(pack.exec_error)}</p>')
    else:
        bits.append(
            '<p class="lead">No national executive summary stored. Paste an OpenAI key on Upload files and rebuild scorecards.</p>'
        )
    bits.append("</section>")
    return "\n".join(bits)


def _html_section(title: str, note: str, df: pd.DataFrame) -> str:
    return f"""<section>
  <h2>{html.escape(title)}</h2>
  <p class="note">{html.escape(note)}</p>
  {_df_html(df)}
</section>
"""


def _html_glossary() -> str:
    items = "".join(
        f"<dt>{html.escape(t)}</dt><dd>{html.escape(d)}</dd>" for t, d in GLOSSARY
    )
    calc = "".join(
        f"<dt>{html.escape(t)}</dt><dd>{html.escape(d)}</dd>" for t, d in CALCULATION_NOTES
    )
    return f"""<section>
  <h2>Glossary</h2>
  <p class="note">Read this first. Every later table uses these words. From drop / unvisited / unbilled add to Gap.</p>
  <dl class="glossary">{items}</dl>
  <h2>How the figures are calculated</h2>
  <dl class="glossary">{calc}</dl>
</section>
"""


def _df_html(df: pd.DataFrame) -> str:
    if df is None or df.empty:
        return "<p class='note'>No rows at this layer for this period.</p>"
    heads = "".join(f"<th>{html.escape(str(c))}</th>" for c in df.columns)
    body = []
    for _, row in df.iterrows():
        klass = row_tone(row) or ""
        tds = "".join(f"<td>{_html_td(row[c], c)}</td>" for c in df.columns)
        body.append(f"<tr class='{klass}'>{tds}</tr>")
    return f"<table><thead><tr>{heads}</tr></thead><tbody>{''.join(body)}</tbody></table>"


def _html_td(val: Any, col: str) -> str:
    text = html.escape(_html_cell(val, col))
    if str(col) == "Remarks":
        text = text.replace("\n", "<br/>")
    return text


def _html_cell(val: Any, col: str) -> str:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return "—"
    if isinstance(val, (int, float)) and str(col) == "Drop size (MT)":
        return f"{val:.2f}"
    if isinstance(val, (int, float)) and "(MT)" in str(col):
        if str(col).startswith("Extra") or str(col).startswith("vs ") or str(col).startswith("From "):
            return f"{val:+.0f}"
        return f"{val:,.0f}"
    if isinstance(val, (int, float)) and ("Strike" in str(col) or str(col).endswith("%")):
        return f"{val:.0f}"
    if isinstance(val, float):
        return f"{val:.0f}"
    return str(val)


def write_html(pack: StrategyPack, path: Path | str, detailed: bool = False) -> None:
    Path(path).write_text(render_html(pack, detailed=detailed), encoding="utf-8")


def write_pdf(pack: StrategyPack, path: Path | str, detailed: bool = False) -> None:
    from sndintel.pdf import write_pdf as _write

    _write(pack, path, detailed=detailed)
