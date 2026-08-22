"""This-week action engine: purchase cycle, leftover cover, due / another visit / lapsing."""

from calendar import monthrange
from datetime import date, timedelta
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from sndintel.action import (
    ACTION_CALL,
    ACTION_CONVERT,
    ACTION_HOLD,
    ACTION_LIFT,
    ACTION_RECOVER,
    SHOP_FLOOR_MT,
    build_action_pack,
)
from sndintel.action_report import excel_bytes, iter_action_sheets
from sndintel.io_utils import period_key, shift_period


def _period_range(end: str, n: int) -> list[str]:
    out = [end]
    cur = end
    for _ in range(n - 1):
        cur = shift_period(cur, -1)
        out.append(cur)
    return list(reversed(out))


def _cycle_dates(last: date, every: int, n: int) -> list[date]:
    out = []
    cur = last
    for _ in range(n):
        out.append(cur)
        cur = cur - timedelta(days=every)
    return list(reversed(out))


def _bill(store_id, name, city, dist, dsr, day: date, vol: float) -> dict:
    return {
        "store_id": store_id,
        "store_name": name,
        "city": city,
        "distributor": dist,
        "dsr_name": dsr,
        "section": city,
        "sale_date": day.isoformat(),
        "year": day.year,
        "month": day.month,
        "day": day.day,
        "period": f"{day.year:04d}-{day.month:02d}",
        "volume_mt": vol,
    }


def _months_from_days(day_rows: list[dict]) -> list[dict]:
    frame = pd.DataFrame(day_rows)
    rows = []
    for (sid, period), g in frame.groupby(["store_id", "period"], sort=False):
        first = g.iloc[0]
        rows.append(
            {
                "store_id": sid,
                "store_name": first["store_name"],
                "city": first["city"],
                "distributor": first["distributor"],
                "dsr_name": first["dsr_name"],
                "section": first["section"],
                "period": period,
                "year": int(first["year"]),
                "month": int(first["month"]),
                "volume_mt": float(g["volume_mt"].sum()),
                "sku_count": 1,
                "billed": 1,
                "zone": None,
            }
        )
    return rows


def _ledger(periods: list[str], as_of: int) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "period": p,
                "status": "closed" if p != periods[-1] else "mtd_open",
                "as_of_day": monthrange(int(p[:4]), int(p[5:7]))[1] if p != periods[-1] else as_of,
                "days_in_month": monthrange(int(p[:4]), int(p[5:7]))[1],
            }
            for p in periods
        ]
    )


def _attrs(sid, name):
    return {
        "store_id": sid,
        "store_name": name,
        "city": "Karachi",
        "distributor": "Eva Foods",
        "dsr_name": "Amir",
        "section": "Clifton",
        "in_universe": 1,
    }


def _panel(as_of: int = 22):
    """Six doors with distinct buying patterns. Score date is 22 Aug 2026."""
    periods = _period_range("2026-08", 20)
    city, dist, dsr = "Karachi", "Eva Foods", "Amir"
    day_rows: list[dict] = []

    for d in _cycle_dates(date(2026, 7, 28), 15, 20):
        day_rows.append(_bill("DUE1", "Due Mart", city, dist, dsr, d, 1.0))
    for d in _cycle_dates(date(2026, 7, 28), 15, 20):
        day_rows.append(_bill("VIS1", "Visited Mart", city, dist, dsr, d, 1.0))
    for d in _cycle_dates(date(2026, 8, 6), 15, 20):
        vol = 0.25 if (d.year, d.month) == (2026, 8) else 1.0
        day_rows.append(_bill("LITE1", "Lite Mart", city, dist, dsr, d, vol))
    for d in _cycle_dates(date(2026, 6, 20), 15, 16):
        day_rows.append(_bill("LAPSE1", "Lapse Mart", city, dist, dsr, d, 1.0))
    for d, vol in (
        (date(2026, 2, 15), 2.0),
        (date(2026, 3, 15), 2.0),
        (date(2026, 4, 15), 2.0),
        (date(2026, 5, 15), 1.0),
        (date(2026, 6, 15), 1.0),
        (date(2026, 7, 15), 1.0),
    ):
        day_rows.append(_bill("DOWN1", "Down Mart", city, dist, dsr, d, vol))
    for d, vol in (
        (date(2026, 2, 20), 2.0),
        (date(2026, 3, 20), 2.0),
        (date(2026, 4, 20), 2.0),
        (date(2026, 5, 20), 1.0),
        (date(2026, 6, 20), 1.0),
        (date(2026, 7, 5), 2.0),
        (date(2026, 7, 20), 2.0),
    ):
        day_rows.append(_bill("HOLD1", "Hold Mart", city, dist, dsr, d, vol))

    month_rows = _months_from_days(day_rows)
    month_rows.append(
        {
            "store_id": "GHOST",
            "store_name": "Ghost",
            "city": "Lahore",
            "distributor": "Ghost Dist",
            "dsr_name": "Ghost DSR",
            "section": "X",
            "period": "2024-12",
            "year": 2024,
            "month": 12,
            "volume_mt": 0.10,
            "sku_count": 1,
            "billed": 1,
            "zone": None,
        }
    )
    stores = pd.DataFrame(
        [
            _attrs("DUE1", "Due Mart"),
            _attrs("VIS1", "Visited Mart"),
            _attrs("LITE1", "Lite Mart"),
            _attrs("LAPSE1", "Lapse Mart"),
            _attrs("DOWN1", "Down Mart"),
            _attrs("HOLD1", "Hold Mart"),
            {
                "store_id": "GHOST",
                "store_name": "Ghost",
                "city": "Lahore",
                "distributor": "Ghost Dist",
                "dsr_name": "Ghost DSR",
                "section": "X",
                "in_universe": 1,
            },
        ]
    )
    visits = pd.DataFrame([{"store_id": "VIS1", "period": "2026-08", "visits": 2}])
    return pd.DataFrame(month_rows), stores, pd.DataFrame(day_rows), _ledger(periods, as_of), visits


def test_cycle_cover_and_lapse_name_the_right_doors():
    shop_month, stores, shop_day, ledger, visits = _panel()
    pack = build_action_pack(shop_month, stores, shop_day, visits=visits, ledger=ledger, period="2026-08")
    raw = pack.raw_shops.set_index("store_id")

    assert "DUE1" in raw.index
    assert "GHOST" not in raw.index
    assert raw.loc["DUE1", "action"] == ACTION_CALL
    assert raw.loc["VIS1", "action"] == ACTION_CONVERT
    assert raw.loc["LITE1", "action"] == ACTION_LIFT
    assert raw.loc["LAPSE1", "action"] == ACTION_RECOVER
    assert raw.loc["DOWN1", "action"] == ACTION_RECOVER
    assert raw.loc["HOLD1", "action"] == ACTION_HOLD

    assert float(raw.loc["DUE1", "cycle_days"]) < 20
    assert float(raw.loc["DUE1", "days_since_bill"]) >= 15
    assert float(raw.loc["HOLD1", "cover_left_days"]) > 7
    assert float(raw.loc["HOLD1", "last_month_mt"]) >= 3.2
    assert float(raw.loc["LITE1", "billed_mt"]) < 0.5
    assert float(raw.loc["LITE1", "week_target_mt"]) >= SHOP_FLOOR_MT
    assert float(raw.loc["DUE1", "week_target_mt"]) >= SHOP_FLOOR_MT

    assert "Due Mart" in str(raw.loc["DUE1", "instruction"])
    assert "not visited" in str(raw.loc["DUE1", "instruction"]).lower()
    assert "KG" in str(raw.loc["DUE1", "instruction"])
    assert "another visit" in str(raw.loc["LITE1", "instruction"]).lower()
    assert "Hold Mart" in str(raw.loc["HOLD1", "instruction"])
    assert pack.source == "cycle"
    assert "Ask (KG)" in pack.calls.columns
    assert "Billed (KG)" in pack.calls.columns
    assert "AMS (KG)" in pack.calls.columns
    assert "Billed (KG)" in pack.dsrs.columns
    assert "AMS (KG)" in pack.dsrs.columns
    assert "Billed (KG)" in pack.distributors.columns
    assert int(pack.calls.loc[pack.calls["Shop"] == "Due Mart", "Ask (KG)"].iloc[0]) == 1000

    assert "Due Mart" in set(pack.calls["Shop"].astype(str))
    assert "Visited Mart" in set(pack.converts["Shop"].astype(str))
    assert "Lite Mart" in set(pack.lifts["Shop"].astype(str))
    assert "Lapse Mart" in set(pack.lapses["Shop"].astype(str))
    assert "Down Mart" in set(pack.lapses["Shop"].astype(str))
    assert "Hold Mart" not in set(pack.calls["Shop"].astype(str))
    assert "Due Mart" not in set(pack.lapses["Shop"].astype(str))


def test_unvisited_due_ranks_ahead_of_visited_due():
    shop_month, stores, shop_day, ledger, visits = _panel()
    pack = build_action_pack(shop_month, stores, shop_day, visits=visits, ledger=ledger, period="2026-08")
    raw = pack.raw_shops.set_index("store_id")
    assert float(raw.loc["DUE1", "value_score"]) > float(raw.loc["VIS1", "value_score"])


def test_dsr_instruction_names_counts_and_tonnes():
    shop_month, stores, shop_day, ledger, visits = _panel()
    pack = build_action_pack(shop_month, stores, shop_day, visits=visits, ledger=ledger, period="2026-08")
    assert not pack.dsrs.empty
    row = pack.dsrs.iloc[0]
    assert row["DSR"] == "Amir"
    assert "Push Amir" in str(row["Do this"])
    assert "this week" in str(row["Do this"]).lower()
    assert "KG" in str(row["Do this"])
    assert "Ask this week (KG)" in pack.dsrs.columns
    assert "AMS (KG)" in pack.country.columns


def test_action_workbook_has_due_another_visit_and_lapsing():
    shop_month, stores, shop_day, ledger, visits = _panel()
    pack = build_action_pack(shop_month, stores, shop_day, visits=visits, ledger=ledger, period="2026-08")
    names = [s[0] for s in iter_action_sheets(pack)]
    assert names == [
        "01 Country",
        "02 Distributors",
        "03 DSRs",
        "04 Due",
        "05 Due visited",
        "06 Another visit",
        "07 Lapsing",
        "08 Backtest",
    ]
    raw = excel_bytes(pack)
    wb = load_workbook(BytesIO(raw))
    assert "00 Cover" in wb.sheetnames
    assert "04 Due" in wb.sheetnames
    assert "06 Another visit" in wb.sheetnames
    assert "07 Lapsing" in wb.sheetnames
    detailed = excel_bytes(pack, detailed=True)
    wb_d = load_workbook(BytesIO(detailed))
    assert "04 Shops" in wb_d.sheetnames
    assert pack.backtest is not None
    if not pack.backtest.empty:
        due_col = "Due precision %" if "Due precision %" in pack.backtest.columns else "Due precision"
        rnd_col = "Random precision %" if "Random precision %" in pack.backtest.columns else "Random precision"
        if due_col in pack.backtest.columns and rnd_col in pack.backtest.columns:
            due_p = pack.backtest.iloc[0][due_col]
            rnd = pack.backtest.iloc[0][rnd_col]
            assert due_p is None or rnd is None or float(due_p) >= float(rnd) - 5

    assert (pack.dsrs["Due"] + pack.dsrs["Due · visited"] + pack.dsrs["Another visit"] + pack.dsrs["Lapsing"]).gt(0).all()


def test_period_key_stable():
    assert period_key(2026, 8) == "2026-08"


def test_billed_this_month_is_not_lapsing():
    """An 8-day cycle that last billed on 5 Aug is another visit on the 22nd, not lapsing."""
    shop_month, stores, shop_day, ledger, visits = _panel()
    city, dist, dsr = "Karachi", "Eva Foods", "Amir"
    extra_days = []
    for d in _cycle_dates(date(2026, 8, 5), 8, 16):
        extra_days.append(_bill("AGAIN1", "Again Mart", city, dist, dsr, d, 1.0))
    shop_day = pd.concat([shop_day, pd.DataFrame(extra_days)], ignore_index=True)
    shop_month = pd.concat([shop_month, pd.DataFrame(_months_from_days(extra_days))], ignore_index=True)
    stores = pd.concat([stores, pd.DataFrame([_attrs("AGAIN1", "Again Mart")])], ignore_index=True)
    pack = build_action_pack(shop_month, stores, shop_day, visits=visits, ledger=ledger, period="2026-08")
    raw = pack.raw_shops.set_index("store_id")
    assert raw.loc["AGAIN1", "action"] == ACTION_LIFT
    assert raw.loc["AGAIN1", "action"] != ACTION_RECOVER


def test_yesterdays_stub_is_not_another_visit():
    shop_month, stores, shop_day, ledger, visits = _panel()
    city, dist, dsr = "Karachi", "Eva Foods", "Amir"
    extra_days = []
    for d in _cycle_dates(date(2026, 7, 25), 15, 16):
        extra_days.append(_bill("STUB1", "Stub Mart", city, dist, dsr, d, 1.0))
    extra_days.append(_bill("STUB1", "Stub Mart", city, dist, dsr, date(2026, 8, 21), 0.05))
    shop_day = pd.concat([shop_day, pd.DataFrame(extra_days)], ignore_index=True)
    shop_month = pd.concat([shop_month, pd.DataFrame(_months_from_days(extra_days))], ignore_index=True)
    stores = pd.concat([stores, pd.DataFrame([_attrs("STUB1", "Stub Mart")])], ignore_index=True)
    pack = build_action_pack(shop_month, stores, shop_day, visits=visits, ledger=ledger, period="2026-08")
    raw = pack.raw_shops.set_index("store_id")
    assert raw.loc["STUB1", "action"] == ACTION_HOLD
