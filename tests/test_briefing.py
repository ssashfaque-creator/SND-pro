"""Strategy pack: AMS, recoverable volume, layered lists, Excel export."""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import Workbook, load_workbook

from sndintel.briefing import (
    SUMMARY_DIST_N,
    SUMMARY_DSR_N,
    SUMMARY_SHOP_N,
    ams_last_n,
    build_strategy_pack,
    excel_bytes,
    excel_bytes_detailed,
    focus_pack,
    keep_top_holes,
    list_report_entities,
    pdf_bytes,
    render_html,
)
from sndintel.coverage import allocate_recoverable_drivers
from sndintel.hierarchy import build_hierarchy_pack


def _with_recent_ams(rows, volume_by_store=None, months=("2026-05", "2026-06", "2026-07")):
    """Give named doors three closed months so AMS is defined (and not zero)."""
    extra = []
    seen = set()
    for r in rows:
        sid = r["store_id"]
        if sid in seen or r["period"] != "2026-08":
            continue
        seen.add(sid)
        if volume_by_store is not None and sid not in volume_by_store:
            continue
        vol = volume_by_store.get(sid, max(float(r["volume_mt"]), 1.0)) if volume_by_store else max(float(r["volume_mt"]), 1.0)
        for per in months:
            extra.append({**r, "period": per, "year": int(per[:4]), "month": int(per[5:7]), "volume_mt": vol})
    return rows + extra


def _row(store_id, period, volume, city, dist, dsr, section="A", name="Shop"):
    return {
        "store_id": store_id,
        "period": period,
        "year": int(period[:4]),
        "month": int(period[5:7]),
        "volume_mt": volume,
        "sku_count": 1,
        "billed": 1 if volume > 0 else 0,
        "distributor": dist,
        "dsr_name": dsr,
        "section": section,
        "store_name": name,
        "zone": "South" if city == "Karachi" else "Central",
        "city": city,
    }


def _stores(rows):
    seen = {}
    for r in rows:
        seen[r["store_id"]] = r
    return pd.DataFrame(
        [
            {
                "store_id": r["store_id"],
                "store_name": r["store_name"],
                "distributor": r["distributor"],
                "dsr_name": r["dsr_name"],
                "zone": r["zone"],
                "city": r["city"],
                "section": r["section"],
            }
            for r in seen.values()
        ]
    )


def test_ams_is_mean_of_last_three_closed_months():
    rows = []
    for per, vol in [("2026-05", 10.0), ("2026-06", 20.0), ("2026-07", 30.0), ("2026-08", 5.0)]:
        rows.append(_row("K1", per, vol, "Karachi", "Eva Foods", "Amir", name="Kifaya"))
    sm = pd.DataFrame(rows)
    ams = ams_last_n(sm, "2026-08", ["city"])
    assert abs(float(ams.iloc[0]["ams_3m"]) - 20.0) < 1e-9


def test_ams_does_not_pull_last_year_into_a_gapped_window():
    """A missing May must not be filled with July last year — that inflated AMS."""
    rows = []
    rows.append(_row("K1", "2025-07", 90.0, "Karachi", "Eva Foods", "Amir", name="Kifaya"))
    rows.append(_row("K1", "2026-06", 10.0, "Karachi", "Eva Foods", "Amir", name="Kifaya"))
    rows.append(_row("K1", "2026-07", 20.0, "Karachi", "Eva Foods", "Amir", name="Kifaya"))
    rows.append(_row("K1", "2026-08", 5.0, "Karachi", "Eva Foods", "Amir", name="Kifaya"))
    sm = pd.DataFrame(rows)
    ams = ams_last_n(sm, "2026-08", ["city"])
    # (0 + 10 + 20) / 3 — May is missing, not 2025-07.
    assert abs(float(ams.iloc[0]["ams_3m"]) - 10.0) < 1e-9


def test_ams_identical_shop_month_copies_are_not_summed():
    rows = []
    for per, vol in [("2026-05", 10.0), ("2026-06", 10.0), ("2026-07", 10.0), ("2026-08", 5.0)]:
        rows.append(_row("K1", per, vol, "Karachi", "Eva Foods", "Amir", name="Kifaya"))
        rows.append(_row("K1", per, vol, "Karachi", "Eva Foods", "Amir", name="Kifaya"))
    sm = pd.DataFrame(rows)
    ams = ams_last_n(sm, "2026-08", ["city"])
    assert abs(float(ams.iloc[0]["ams_3m"]) - 10.0) < 1e-9


def test_pack_ams_stays_full_month_when_mtd_is_open():
    """Day 15 MTD: printed AMS is the full-month last-3 mean; vs AMS applies 15/31."""
    rows = []
    rows.append(_row("K1", "2026-08", 15.0, "Karachi", "Eva Foods", "Amir", name="Kifaya"))
    rows.append(_row("K1", "2025-08", 31.0, "Karachi", "Eva Foods", "Amir", name="Kifaya"))
    rows = _with_recent_ams(rows, volume_by_store={"K1": 31.0})
    sm = pd.DataFrame(rows)
    pack_h = build_hierarchy_pack(
        sm,
        _stores(rows),
        ledger=pd.DataFrame(
            [{"period": "2026-08", "status": "mtd_open", "as_of_day": 15, "days_in_month": 31}]
        ),
    )
    report = build_strategy_pack(
        pack_h.units,
        sm,
        period="2026-08",
        ledger=pd.DataFrame(
            [{"period": "2026-08", "status": "mtd_open", "as_of_day": 15, "days_in_month": 31}]
        ),
    )
    country = report.cities[report.cities["City"] == "Country"].iloc[0]
    ams = float(country["AMS last 3 months (MT)"])
    billed = float(country["Billed this period (MT)"])
    assert abs(ams - 31.0) < 1.0, ams
    assert abs(billed - 15.0) < 1.0
    vs = float(country["vs AMS (MT)"])
    assert abs(vs - (billed - 31.0 * (15.0 / 31.0))) < 1.5


def test_pack_layers_cities_then_those_dists_then_all_dists():
    """Karachi and Local Dist both miss their own Expected. Ghost Dist has AMS = 0 so it is hidden."""
    rows = []
    # Karachi far behind typical August; Lahore also down vs last year (so vs Expected).
    rows.append(_row("K1", "2026-08", 5.0, "Karachi", "Eva Foods", "Amir", name="Kifaya"))
    rows.append(_row("K1", "2025-08", 80.0, "Karachi", "Eva Foods", "Amir", name="Kifaya"))
    rows.append(_row("K2", "2026-08", 15.0, "Karachi", "South Dist", "Amir", name="Hold K"))
    rows.append(_row("K2", "2025-08", 20.0, "Karachi", "South Dist", "Amir", name="Hold K"))
    rows.append(_row("L1", "2026-08", 90.0, "Lahore", "Holding Dist", "Lahore Ace", name="Big L"))
    rows.append(_row("L1", "2025-08", 100.0, "Lahore", "Holding Dist", "Lahore Ace", name="Big L"))
    rows.append(_row("L2", "2026-08", 10.0, "Lahore", "Local Dist", "Lahore Weak", name="Small L"))
    rows.append(_row("L2", "2025-08", 100.0, "Lahore", "Local Dist", "Lahore Weak", name="Small L"))
    rows.append(_row("G1", "2026-08", 1.0, "Karachi", "Ghost Dist", "Ghost DSR", name="Ghost Shop"))
    rows.append(_row("G1", "2025-08", 40.0, "Karachi", "Ghost Dist", "Ghost DSR", name="Ghost Shop"))
    rows = _with_recent_ams(rows, volume_by_store={"K1": 8.0, "K2": 18.0, "L1": 95.0, "L2": 80.0})
    sm = pd.DataFrame(rows)
    pack_h = build_hierarchy_pack(sm, _stores(rows), ledger=pd.DataFrame([{"period": "2026-08", "status": "closed"}]))
    report = build_strategy_pack(pack_h.units, sm, period="2026-08")
    cities = report.cities
    assert "City" in cities.columns
    assert "AMS last 3 months (MT)" in cities.columns
    assert "Gap (MT)" in cities.columns
    assert "Recoverable (MT)" not in cities.columns
    assert "Situation" not in cities.columns
    assert "Drop size (MT)" in cities.columns
    assert "Expected this month (MT)" in cities.columns
    assert "Fair share of country (MT)" not in cities.columns
    assert "Zone" not in cities.columns
    exp_i = list(cities.columns).index("Expected this month (MT)")
    gap_i = list(cities.columns).index("Gap (MT)")
    assert gap_i > exp_i
    assert "Sales-team plan (MT)" in cities.columns
    assert list(cities.columns)[exp_i + 1] == "Sales-team plan (MT)"
    assert list(cities.columns)[-2] == "Drop size (MT)"
    assert "From drop size (MT)" in cities.columns
    assert "From unvisited shops (MT)" in cities.columns
    assert "From unbilled shops (MT)" in cities.columns
    assert "Remarks" in cities.columns
    assert list(cities.columns)[-1] == "Remarks"
    assert "Main driver" not in cities.columns
    assert "From coverage (MT)" not in cities.columns
    assert "Extra vs country (MT)" not in cities.columns
    assert cities.iloc[0]["City"] == "Country"
    body = cities[cities["City"] != "Country"]
    rec = body["Gap (MT)"].tolist()
    assert rec == sorted(rec, reverse=True)
    khi = body[body["City"] == "Karachi"].iloc[0]
    assert float(khi["Gap (MT)"]) > 0
    country = cities[cities["City"] == "Country"].iloc[0]
    # Expected ≈ billed + Gap (Gap is 0 when billed is ahead of the recent run-rate).
    assert abs(
        float(country["Expected this month (MT)"])
        - float(country["Billed this period (MT)"])
        - float(country["Gap (MT)"])
    ) < 8
    remarks = " ".join(cities["Remarks"].dropna().astype(str))
    assert "vs expected" in remarks.lower()
    assert "national average" in remarks.lower()
    # Full lagging-distributor list includes Local Dist (behind its own Expected).
    # Ghost Dist has no volume in the AMS window, so Expected is 0 and it stays hidden.
    names = set(report.lagging_distributors["Distributor"].astype(str))
    assert "Local Dist" in names
    assert "Eva Foods" in names
    assert "Ghost Dist" not in names
    assert "Billed shops" in report.lagging_distributors.columns
    assert "Strike %" in report.lagging_distributors.columns
    assert "Universe" in report.lagging_distributors.columns
    assert "Drop size (MT)" in report.lagging_distributors.columns
    assert "Expected this month (MT)" in report.lagging_distributors.columns
    assert "Fair share of its city (MT)" not in report.lagging_distributors.columns
    assert "Zone" not in report.lagging_distributors.columns
    assert "Remarks" in report.lagging_dsrs.columns
    assert "Strike %" in report.lagging_dsrs.columns
    assert "Ghost DSR" not in set(report.lagging_dsrs["DSR"].astype(str))
    assert "Ghost Dist" not in set(report.all_distributors["Distributor"].astype(str))


def test_excel_and_html_are_readable_packs():
    rows = []
    rows.append(_row("K1", "2026-08", 10.0, "Karachi", "Eva Foods", "Amir", name="Kifaya"))
    rows.append(_row("K1", "2025-08", 30.0, "Karachi", "Eva Foods", "Amir", name="Kifaya"))
    rows.append(_row("K1", "2026-07", 28.0, "Karachi", "Eva Foods", "Amir", name="Kifaya"))
    rows.append(_row("K1", "2026-06", 26.0, "Karachi", "Eva Foods", "Amir", name="Kifaya"))
    rows.append(_row("K1", "2026-05", 24.0, "Karachi", "Eva Foods", "Amir", name="Kifaya"))
    sm = pd.DataFrame(rows)
    pack_h = build_hierarchy_pack(sm, _stores(rows), ledger=pd.DataFrame([{"period": "2026-08", "status": "closed"}]))
    report = build_strategy_pack(pack_h.units, sm, period="2026-08")
    raw = excel_bytes(report)
    wb = load_workbook(BytesIO(raw))
    names = wb.sheetnames
    assert names[0] == "00 Cover"
    assert "01 Country by city" in names
    assert "02 Distributors by volume" in names
    assert "03 Distributors by seriousness" in names
    assert "04 DSRs by volume" in names
    assert "05 DSRs by seriousness" in names
    assert "06 Whales" in names
    assert "07 Top 50 shops" in names
    assert "08 DSR capacity" in names
    assert "02 Top 10 distributors" not in names
    assert "03 Top 10 DSRs" not in names
    assert "04 Top 10 dists all cities" not in names
    assert "05 Top 10 DSRs of those" not in names
    assert "03 Top 50 shops of those" not in names
    assert len([n for n in names if n != "00 Cover"]) == 8
    cover = wb["00 Cover"]
    col_a = [c.value for row in cover.iter_rows(min_col=1, max_col=1, values_only=False) for c in row]
    assert "Glossary" in col_a
    assert "How the figures are calculated" in col_a
    assert "Executive summary" in col_a
    assert col_a.index("Glossary") < col_a.index("Executive summary")
    html = render_html(report)
    assert html.find("Glossary") < html.find("1. Every city versus its own Expected")
    assert "How the figures are calculated" in html
    assert "Executive summary" in html
    detailed = render_html(report, detailed=True)
    assert "Every distributor with AMS greater than 0" in detailed
    assert "Every DSR with AMS greater than 0" in detailed
    assert "National shop list" in detailed
    assert "National DSR list" not in detailed
    raw_d = excel_bytes_detailed(report)
    wb_d = load_workbook(BytesIO(raw_d))
    assert "01 City detail" in wb_d.sheetnames
    assert "02 Distributor detail" in wb_d.sheetnames
    assert "03 DSR detail" in wb_d.sheetnames
    assert "04 National shops" in wb_d.sheetnames
    assert "04 National DSRs" not in wb_d.sheetnames
    assert "05 National shops" not in wb_d.sheetnames
    pdf = pdf_bytes(report)
    assert pdf.startswith(b"%PDF")
    pdf_d = pdf_bytes(report, detailed=True)
    assert pdf_d.startswith(b"%PDF")


def test_excel_export_treats_pandas_na_as_blank():
    """Detailed pack used to crash: openpyxl cannot store pd.NA."""
    from sndintel.briefing import _excel_value, _sheet_table

    assert _excel_value(pd.NA) is None
    assert _excel_value(float("nan")) is None
    wb = Workbook()
    df = pd.DataFrame(
        {
            "Shop": ["A"],
            "Gap (MT)": [pd.NA],
            "AMS last 3 months (MT)": [pd.NA],
            "Remarks": [pd.NA],
        }
    )
    _sheet_table(wb, "NA sheet", "Heading", "Note", df)
    ws = wb["NA sheet"]
    assert ws.cell(5, 1).value == "A"
    assert ws.cell(5, 2).value is None
    assert ws.cell(5, 3).value is None


def test_from_columns_sum_to_recoverable_and_are_positive_when_behind():
    df = pd.DataFrame(
        {
            "recoverable_mt": [100.0, 0.0, 50.0],
            "isolated_mt": [-100.0, 40.0, -50.0],
            "volume_mt": [80.0, 140.0, 90.0],
            "share_expected_mt": [180.0, 100.0, 140.0],
            "from_drop_size_mt": [-40.0, 20.0, 10.0],
            "from_unvisited_mt": [-30.0, 5.0, -20.0],
            "from_unbilled_mt": [-80.0, 15.0, -30.0],
        }
    )
    out = allocate_recoverable_drivers(df)
    behind = out.iloc[0]
    assert abs(behind["from_drop_size_mt"] + behind["from_unvisited_mt"] + behind["from_unbilled_mt"] - 100.0) < 1e-9
    assert behind["from_drop_size_mt"] >= -1e-9
    assert behind["from_unvisited_mt"] >= -1e-9
    assert behind["from_unbilled_mt"] >= -1e-9
    ahead = out.iloc[1]
    assert ahead["recoverable_mt"] == 0.0
    assert ahead["from_drop_size_mt"] <= 1e-9
    assert abs(ahead["from_drop_size_mt"] + ahead["from_unvisited_mt"] + ahead["from_unbilled_mt"] + 40.0) < 1e-9
    mixed = out.iloc[2]
    assert abs(mixed["from_drop_size_mt"] + mixed["from_unvisited_mt"] + mixed["from_unbilled_mt"] - 50.0) < 1e-9
    assert mixed["from_drop_size_mt"] == 0.0  # gain is not a hole


def test_pack_from_columns_sum_to_recoverable_after_rounding():
    rows = []
    rows.append(_row("K1", "2026-08", 5.0, "Karachi", "Eva Foods", "Amir", name="Kifaya"))
    rows.append(_row("K1", "2025-08", 80.0, "Karachi", "Eva Foods", "Amir", name="Kifaya"))
    rows.append(_row("K2", "2026-08", 15.0, "Karachi", "South Dist", "Amir", name="Hold K"))
    rows.append(_row("K2", "2025-08", 20.0, "Karachi", "South Dist", "Amir", name="Hold K"))
    rows.append(_row("L1", "2026-08", 90.0, "Lahore", "Holding Dist", "Lahore Ace", name="Big L"))
    rows.append(_row("L1", "2025-08", 100.0, "Lahore", "Holding Dist", "Lahore Ace", name="Big L"))
    rows.append(_row("L2", "2026-08", 10.0, "Lahore", "Local Dist", "Lahore Weak", name="Small L"))
    rows.append(_row("L2", "2025-08", 100.0, "Lahore", "Local Dist", "Lahore Weak", name="Small L"))
    rows = _with_recent_ams(rows, volume_by_store={"K1": 8.0, "K2": 18.0, "L1": 95.0, "L2": 80.0})
    sm = pd.DataFrame(rows)
    pack_h = build_hierarchy_pack(sm, _stores(rows), ledger=pd.DataFrame([{"period": "2026-08", "status": "closed"}]))
    report = build_strategy_pack(pack_h.units, sm, period="2026-08")
    cities = report.cities
    assert list(cities.columns)[-1] == "Remarks"
    from_cols = ["From drop size (MT)", "From unvisited shops (MT)", "From unbilled shops (MT)"]
    for _, row in cities.iterrows():
        rec = row["Gap (MT)"]
        rec_i = 0 if rec is None or pd.isna(rec) else int(rec)
        parts = [0 if row[c] is None or pd.isna(row[c]) else int(row[c]) for c in from_cols]
        if rec_i > 0:
            assert sum(parts) == rec_i
            assert all(p >= 0 for p in parts)
        # whole numbers in the table
        for col in ["Billed this period (MT)", "Gap (MT)"]:
            val = row[col]
            if val is not None and pd.notna(val):
                assert float(val) == float(int(round(float(val))))
    cities_opt = list_report_entities(report, "City")
    assert "Karachi" in cities_opt
    focused = focus_pack(report, "City", "Karachi")
    assert focused.scope == "city"
    assert set(focused.cities["City"].astype(str)) <= {"Country", "Karachi"}
    dist_opt = list_report_entities(report, "Distributor")
    assert any("Eva Foods" in x for x in dist_opt)


def test_tiny_shops_are_not_on_visit_lists():
    """Kiryana with recoverable ≤ 0.25 MT is rolled off; any shop above that cut stays, even if AMS is small."""
    rows = []
    rows.append(_row("BIG", "2026-08", 4.0, "Karachi", "Eva Foods", "Amir", name="Kifaya Mart"))
    rows.append(_row("BIG", "2025-08", 12.0, "Karachi", "Eva Foods", "Amir", name="Kifaya Mart"))
    rows.append(_row("HOLD", "2026-08", 12.0, "Karachi", "South Dist", "Amir", name="Hold Super"))
    rows.append(_row("HOLD", "2025-08", 12.0, "Karachi", "South Dist", "Amir", name="Hold Super"))
    # Small AMS (~0.4 MT) but a hole above 0.25 MT — must still appear.
    # Put it in a city that moved with the country so the shop hole is the full drop.
    rows.append(_row("MED", "2026-08", 0.0, "Lahore", "Holding Dist", "Ace", name="Medium Mart"))
    rows.append(_row("MED", "2025-08", 0.40, "Lahore", "Holding Dist", "Ace", name="Medium Mart"))
    for i in range(40):
        rows.append(_row(f"T{i}", "2026-08", 0.0, "Karachi", "Eva Foods", "Amir", name=f"Kiryana {i}"))
        rows.append(_row(f"T{i}", "2025-08", 0.04, "Karachi", "Eva Foods", "Amir", name=f"Kiryana {i}"))
    rows.append(_row("L1", "2026-08", 80.0, "Lahore", "Holding Dist", "Ace", name="Lahore Super"))
    rows.append(_row("L1", "2025-08", 80.0, "Lahore", "Holding Dist", "Ace", name="Lahore Super"))
    rows = _with_recent_ams(
        rows,
        volume_by_store={"BIG": 8.0, "HOLD": 12.0, "MED": 0.40, "L1": 80.0, **{f"T{i}": 0.04 for i in range(40)}},
    )
    sm = pd.DataFrame(rows)
    pack_h = build_hierarchy_pack(sm, _stores(rows), ledger=pd.DataFrame([{"period": "2026-08", "status": "closed"}]))
    report = build_strategy_pack(pack_h.units, sm, period="2026-08")
    assert report.kpis["shop_hole_floor_mt"] == 0.25
    names = report.lagging_shops["Shop"].astype(str)
    assert names.str.contains("Kifaya Mart").any()
    assert names.str.contains("Medium Mart").any()
    assert not names.str.contains("Kiryana").any()
    assert names.str.contains("Not listed").any()
    assert report.kpis["n_shops_hidden"] >= 20


def test_keep_top_holes_caps_and_remainder():
    df = pd.DataFrame(
        {
            "grain_id": [f"D{i}" for i in range(15)],
            "recoverable_mt": [float(15 - i) for i in range(15)],
        }
    )
    kept, meta = keep_top_holes(df, 10)
    assert list(kept["grain_id"]) == [f"D{i}" for i in range(10)]
    assert meta["n_kept"] == 10
    assert meta["n_hidden"] == 5
    assert meta["hidden_mt"] == 15.0  # 5+4+3+2+1
    assert meta["top_n"] == 10


def test_summary_pack_hides_tiny_lagging_distributors_as_remainder():
    """More than 10 lagging dists: summary names the top 10; the rest are remainder."""
    rows = []
    rows.append(_row("BIG", "2026-08", 10.0, "Karachi", "Eva Foods", "Amir", name="Kifaya"))
    rows.append(_row("BIG", "2025-08", 80.0, "Karachi", "Eva Foods", "Amir", name="Kifaya"))
    for i in range(15):
        rows.append(_row(f"S{i}", "2026-08", 0.6, "Karachi", f"Tiny Dist {i}", f"Rep {i}", name=f"Shop {i}"))
        rows.append(_row(f"S{i}", "2025-08", 1.2, "Karachi", f"Tiny Dist {i}", f"Rep {i}", name=f"Shop {i}"))
    rows = _with_recent_ams(
        rows,
        volume_by_store={"BIG": 70.0, **{f"S{i}": 1.2 for i in range(15)}},
    )
    sm = pd.DataFrame(rows)
    pack_h = build_hierarchy_pack(sm, _stores(rows), ledger=pd.DataFrame([{"period": "2026-08", "status": "closed"}]))
    report = build_strategy_pack(pack_h.units, sm, period="2026-08")
    named = [x for x in report.lagging_distributors["Distributor"].astype(str) if not str(x).startswith("Not listed")]
    assert "Eva Foods" in named
    assert len(named) == SUMMARY_DIST_N
    assert report.lagging_distributors["Distributor"].astype(str).str.contains("Not listed").any()
    all_names = set(report.all_distributors["Distributor"].astype(str))
    assert "Tiny Dist 0" in all_names
    assert "Tiny Dist 14" in all_names


def test_summary_dsrs_are_national_top_10():
    """DSRs on the summary are the national top 10 by seriousness, at most 10 named."""
    rows = []
    for i in range(12):
        dist = f"Dist {i:02d}"
        dsr = f"Rep {i:02d}"
        sid = f"S{i:02d}"
        billed = float(i)
        rows.append(_row(sid, "2026-08", billed, "Karachi", dist, dsr, name=f"Shop {i:02d}"))
        rows.append(_row(sid, "2025-08", 30.0, "Karachi", dist, dsr, name=f"Shop {i:02d}"))
    rows = _with_recent_ams(rows, volume_by_store={f"S{i:02d}": 30.0 for i in range(12)})
    sm = pd.DataFrame(rows)
    pack_h = build_hierarchy_pack(sm, _stores(rows), ledger=pd.DataFrame([{"period": "2026-08", "status": "closed"}]))
    report = build_strategy_pack(pack_h.units, sm, period="2026-08")
    dist_named = [x for x in report.lagging_distributors["Distributor"].astype(str) if not str(x).startswith("Not listed")]
    assert dist_named == [f"Dist {i:02d}" for i in range(10)]
    dsr_named = [x for x in report.lagging_dsrs["DSR"].astype(str) if not str(x).startswith("Not listed")]
    assert len(dsr_named) == SUMMARY_DSR_N
    assert dsr_named == [f"Rep {i:02d}" for i in range(10)]
    assert "Rep 10" not in dsr_named
    assert "Rep 11" not in dsr_named
    all_dsrs = set(report.all_dsrs["DSR"].astype(str))
    assert "Rep 10" in all_dsrs
    assert "Rep 11" in all_dsrs


def test_summary_includes_serious_dsr_outside_top_10_distributors():
    """A collapsed DSR under an 11th distributor still ranks on the national DSR ten."""
    rows = []
    volume_by_store = {}
    for d in range(10):
        for r in range(5):
            sid = f"D{d:02d}R{r}"
            dist = f"Fat Dist {d:02d}"
            dsr = f"Fat Rep {d:02d}-{r}"
            volume_by_store[sid] = 20.0
            rows.append(_row(sid, "2026-08", 17.0, "Karachi", dist, dsr, name=sid))
            rows.append(_row(sid, "2025-08", 20.0, "Karachi", dist, dsr, name=sid))
    volume_by_store["MILD"] = 30.0
    rows.append(_row("MILD", "2026-08", 24.0, "Lahore", "Thin Dist", "Mild Rep", name="Mild Shop"))
    rows.append(_row("MILD", "2025-08", 30.0, "Lahore", "Thin Dist", "Mild Rep", name="Mild Shop"))
    rows = _with_recent_ams(rows, volume_by_store=volume_by_store)
    sm = pd.DataFrame(rows)
    pack_h = build_hierarchy_pack(sm, _stores(rows), ledger=pd.DataFrame([{"period": "2026-08", "status": "closed"}]))
    report = build_strategy_pack(pack_h.units, sm, period="2026-08")
    dist_named = [x for x in report.lagging_distributors["Distributor"].astype(str) if not str(x).startswith("Not listed")]
    assert all(str(x).startswith("Fat Dist") for x in dist_named)
    assert "Thin Dist" not in dist_named
    dsr_named = [x for x in report.lagging_dsrs["DSR"].astype(str) if not str(x).startswith("Not listed")]
    assert "Mild Rep" in dsr_named
    assert dsr_named[0] == "Mild Rep"
    assert "Mild Rep" in set(report.all_dsrs["DSR"].astype(str))


def test_summary_shops_are_top_50_most_serious():
    """National shop list names at most 50 doors; detailed still has everyone above 0.25 MT."""
    rows = []
    volume_by_store = {}
    for i in range(60):
        sid = f"P{i:02d}"
        ams = 3.0 + i * 0.1
        volume_by_store[sid] = ams
        rows.append(_row(sid, "2026-08", 0.0, "Karachi", "Eva Foods", "Amir", name=f"Door {i:02d}"))
        rows.append(_row(sid, "2025-08", ams, "Karachi", "Eva Foods", "Amir", name=f"Door {i:02d}"))
    rows = _with_recent_ams(rows, volume_by_store=volume_by_store)
    sm = pd.DataFrame(rows)
    pack_h = build_hierarchy_pack(sm, _stores(rows), ledger=pd.DataFrame([{"period": "2026-08", "status": "closed"}]))
    report = build_strategy_pack(pack_h.units, sm, period="2026-08")
    named = [x for x in report.lagging_shops["Shop"].astype(str) if not str(x).startswith("Not listed")]
    assert len(named) == SUMMARY_SHOP_N
    assert named[0] == "Door 59"
    assert "Door 00" not in named
    assert report.lagging_shops["Shop"].astype(str).str.contains("Not listed").any()
    detailed = [x for x in report.all_shops["Shop"].astype(str) if not str(x).startswith("Not listed")]
    assert len(detailed) == 60


def test_keep_top_holes_ranks_by_seriousness_not_gap_tons():
    df = pd.DataFrame(
        {
            "grain_id": ["whale", "crisis", "micro"],
            "recoverable_mt": [20.0, 12.0, 0.4],
            "_seriousness": [20.0 / (200.0 ** 0.5), 12.0 / (15.0 ** 0.5), 0.4 / (0.4 ** 0.5)],
        }
    )
    kept, meta = keep_top_holes(df, 2, rank_col="_seriousness")
    assert list(kept["grain_id"]) == ["crisis", "whale"]
    assert abs(meta["hidden_mt"] - 0.4) < 1e-9


def test_summary_ranks_collapsed_mid_size_ahead_of_large_modest_miss():
    """A 15 MT book at ~80% miss outranks 200 MT books that are only 10% light."""
    rows = []
    volume_by_store = {}
    for i in range(10):
        sid = f"W{i}"
        volume_by_store[sid] = 200.0
        rows.append(_row(sid, "2026-08", 180.0, "Karachi", f"Whale Dist {i}", f"Whale Rep {i}", name=sid))
        rows.append(_row(sid, "2025-08", 200.0, "Karachi", f"Whale Dist {i}", f"Whale Rep {i}", name=sid))
    volume_by_store["CRISIS"] = 15.0
    rows.append(_row("CRISIS", "2026-08", 3.0, "Lahore", "Crisis Dist", "Crisis Rep", name="Crisis Shop"))
    rows.append(_row("CRISIS", "2025-08", 15.0, "Lahore", "Crisis Dist", "Crisis Rep", name="Crisis Shop"))
    rows = _with_recent_ams(rows, volume_by_store=volume_by_store)
    sm = pd.DataFrame(rows)
    pack_h = build_hierarchy_pack(sm, _stores(rows), ledger=pd.DataFrame([{"period": "2026-08", "status": "closed"}]))
    report = build_strategy_pack(pack_h.units, sm, period="2026-08")
    dist_named = [x for x in report.lagging_distributors["Distributor"].astype(str) if not str(x).startswith("Not listed")]
    assert dist_named[0] == "Crisis Dist"
    assert len(dist_named) == SUMMARY_DIST_N
    assert "Crisis Dist" in dist_named
    assert sum(1 for x in dist_named if str(x).startswith("Whale Dist")) == 9
    dsr_named = [x for x in report.lagging_dsrs["DSR"].astype(str) if not str(x).startswith("Not listed")]
    assert dsr_named[0] == "Crisis Rep"
    shop_named = [x for x in report.lagging_shops["Shop"].astype(str) if not str(x).startswith("Not listed")]
    assert shop_named[0] == "Crisis Shop"
    vol_named = [x for x in report.volume_distributors["Distributor"].astype(str) if not str(x).startswith("Not listed")]
    assert vol_named[0].startswith("Whale Dist")
    assert all(str(x).startswith("Whale Dist") for x in vol_named)
    vol_dsr = [x for x in report.volume_dsrs["DSR"].astype(str) if not str(x).startswith("Not listed")]
    assert vol_dsr[0].startswith("Whale Rep")


def test_two_shahids_never_share_a_scorecard_row():
    rows = []
    volume_by_store = {"K1": 20.0, "L1": 12.0}
    rows.append(_row("K1", "2026-08", 4.0, "Karachi", "Eva Foods", "Shahid", name="Karachi Shop"))
    rows.append(_row("K1", "2025-08", 20.0, "Karachi", "Eva Foods", "Shahid", name="Karachi Shop"))
    rows.append(_row("L1", "2026-08", 2.0, "Lahore", "Lahore Dist", "Shahid", name="Lahore Shop"))
    rows.append(_row("L1", "2025-08", 12.0, "Lahore", "Lahore Dist", "Shahid", name="Lahore Shop"))
    rows = _with_recent_ams(rows, volume_by_store=volume_by_store)
    sm = pd.DataFrame(rows)
    pack_h = build_hierarchy_pack(sm, _stores(rows), ledger=pd.DataFrame([{"period": "2026-08", "status": "closed"}]))
    report = build_strategy_pack(pack_h.units, sm, period="2026-08")
    dsrs = pack_h.units[pack_h.units["grain"] == "dsr"]
    assert len(dsrs) == 2
    assert dsrs["grain_id"].nunique() == 2
    named = report.all_dsrs[report.all_dsrs["DSR"].astype(str) == "Shahid"]
    assert len(named) == 2
    assert set(named["City"].astype(str)) == {"Karachi", "Lahore"}

