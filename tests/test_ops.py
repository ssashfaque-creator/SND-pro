"""Monday / DSR / Friday operating packs and closed-loop scoring."""

import pandas as pd

from sndintel.action import ACTION_CALL, ACTION_CONVERT, persist_action_pack
from sndintel.ops import (
    beat_owner_options,
    build_dsr_beat_pack,
    build_friday_pack,
    build_monday_pack,
    filter_beat_by_owner,
    load_outcomes,
    score_closed_loop,
)
from sndintel.storage import connect, init_db


def _listed(sid, billed=0.0, visits=0, action=ACTION_CALL, ask=1.0):
    return {
        "store_id": sid,
        "store_name": sid,
        "action": action,
        "billed_mt": billed,
        "visits": visits,
        "week_target_mt": ask,
        "city": "Karachi",
        "distributor": "Eva",
        "dsr_name": "Amir",
    }


def test_closed_loop_billed_visited_unbilled_and_not_visited():
    previous = pd.DataFrame(
        [
            _listed("BILL", billed=0.0, visits=0),
            _listed("SEEN", billed=0.0, visits=1, action=ACTION_CONVERT),
            _listed("SKIP", billed=0.0, visits=0),
        ]
    )
    shop_month = pd.DataFrame(
        [
            {"store_id": "BILL", "period": "2026-08", "volume_mt": 0.8},
            {"store_id": "SEEN", "period": "2026-08", "volume_mt": 0.0},
            {"store_id": "SKIP", "period": "2026-08", "volume_mt": 0.0},
        ]
    )
    visits = pd.DataFrame(
        [
            {"store_id": "BILL", "period": "2026-08", "visits": 2},
            {"store_id": "SEEN", "period": "2026-08", "visits": 3},
            {"store_id": "SKIP", "period": "2026-08", "visits": 0},
        ]
    )
    out = score_closed_loop(previous, shop_month, visits, "2026-08")
    got = {row.store_id: row.outcome for row in out.itertuples(index=False)}
    assert got["BILL"] == "billed"
    assert got["SEEN"] == "visited_unbilled"
    assert got["SKIP"] == "not_visited"
    friday = build_friday_pack(out)
    assert "billed after the list" in friday.headline.lower()
    assert not friday.sheets[0][3].empty


def test_dsr_beat_pack_caps_per_person():
    from sndintel.action import ActionPack
    from sndintel.config import DSR_DAY_CAP

    rows = []
    for i in range(30):
        rows.append(
            {
                "store_id": f"S{i}",
                "store_name": f"Shop {i}",
                "dsr_name": "Amir",
                "city": "Karachi",
                "distributor": "Eva",
                "action": ACTION_CALL,
                "week_target_mt": 1.0 - i * 0.01,
                "value_score": 30 - i,
                "instruction": "Call",
                "next_drop_mt": 0.5,
                "days_since_bill": 20,
                "cover_left_days": 0,
            }
        )
    pack = ActionPack(period="2026-08", label="August 2026", raw_shops=pd.DataFrame(rows), headline="test")
    beat = build_dsr_beat_pack(pack)
    listed = beat.sheets[0][3]
    assert len(listed) == DSR_DAY_CAP


def test_beat_filter_keeps_two_people_with_the_same_name_apart():
    from sndintel.action import ActionPack

    rows = []
    for city, dist, n in (("Karachi", "Eva", 5), ("Lahore", "Punjab", 5)):
        for i in range(n):
            rows.append(
                {
                    "store_id": f"{city[:1]}{i}",
                    "store_name": f"{city} shop {i}",
                    "dsr_name": "Shahid",
                    "city": city,
                    "distributor": dist,
                    "action": ACTION_CALL,
                    "week_target_mt": 1.0,
                    "value_score": n - i,
                    "instruction": "Call",
                    "next_drop_mt": 0.4,
                    "days_since_bill": 18,
                    "cover_left_days": 0,
                }
            )
    pack = ActionPack(period="2026-08", label="August 2026", raw_shops=pd.DataFrame(rows), headline="test")
    beat = build_dsr_beat_pack(pack, per_dsr=12)
    table = beat.sheets[0][3]
    owners = beat_owner_options(table)
    assert owners == ["Shahid · Karachi · Eva", "Shahid · Lahore · Punjab"]
    karachi = filter_beat_by_owner(table, owners[0])
    lahore = filter_beat_by_owner(table, owners[1])
    assert set(karachi["City"]) == {"Karachi"}
    assert set(lahore["City"]) == {"Lahore"}
    assert len(karachi) == 5
    assert len(lahore) == 5


def test_monday_pack_has_capacity_and_whales():
    from sndintel.action import ActionPack
    from sndintel.monday import monday_summary_sheets

    shops = pd.DataFrame(
        [
            {
                "store_id": "W",
                "store_name": "Whale Mart",
                "city": "Karachi",
                "distributor": "Eva",
                "dsr_name": "Amir",
                "ams_3m": 2.0,
                "expected_mt": 2.0,
                "billed_mt": 0.2,
                "remaining_mt": 1.8,
                "week_target_mt": 1.0,
                "call_status": "Unvisited",
                "action": ACTION_CALL,
                "instruction": "Call Whale Mart",
                "value_score": 10,
                "coming_due": False,
            }
        ]
    )
    units = pd.DataFrame(
        [
            {
                "grain": "city",
                "grain_id": "Karachi",
                "volume_mt": 10,
                "expected_mt": 20,
                "ams_3m": 18,
                "isolated_mt": -10,
                "from_unbilled_mt": 8,
                "from_unvisited_mt": 1,
                "from_drop_size_mt": 1,
                "visit_rate": 0.99,
                "strike_rate": 0.4,
                "universe": 800,
                "has_visit_file": 1,
            }
        ]
    )
    pack = ActionPack(
        period="2026-08",
        label="August 2026",
        headline="Karachi is the hole",
        as_of_day=21,
        days_in_month=31,
        days_left=10,
        country=pd.DataFrame([{"Billed (KG)": 200}]),
        distributors=pd.DataFrame([{"Distributor": "Eva"}]),
        raw_shops=shops,
    )
    monday = build_monday_pack(pack, units, visits=pd.DataFrame([{"store_id": "W", "period": "2026-08", "visits": 1}]))
    by_key = {s[0]: s for s in monday.sheets}
    assert "02 City drivers" in by_key
    assert "03 City actions" in by_key
    assert "04 Who to push" in by_key
    assert "05 Highlighted distributors" in by_key
    assert "06 Highlighted stores" in by_key
    drivers = by_key["02 City drivers"][3]
    assert list(drivers.columns[:6]) == ["City", "AMS (MT)", "Billed (MT)", "Expected (MT)", "Gap (MT)", "Universe"]
    assert int(drivers.iloc[0]["Universe"]) == 800
    assert float(drivers.iloc[0]["AMS (MT)"]) == 18.0
    country = by_key["01 Country"][3]
    assert "Billed (MT)" in country.columns
    assert "Still to Expected (MT)" not in country.columns
    who = by_key["04 Who to push"][3]
    assert "Span ×" not in who.columns
    assert "Day cap" not in who.columns
    assert "Why" in who.columns
    whales = by_key["06 Highlighted stores"][3]
    assert "Whale Mart" in set(whales["Shop"].astype(str))
    assert "Ask rest of month (KG)" in whales.columns
    summary_keys = [s[0] for s in monday_summary_sheets(monday.sheets)]
    assert summary_keys.index("05 Highlighted distributors") < summary_keys.index("06 Highlighted stores")
    assert any("Karachi" in w for w in monday.warnings)


def test_city_driver_ams_fills_from_shops_when_units_lack_it():
    from sndintel.monday import city_driver_table

    units = pd.DataFrame(
        [
            {
                "grain": "city",
                "grain_id": "Karachi",
                "volume_mt": 10,
                "expected_mt": 20,
                "isolated_mt": -10,
                "from_unbilled_mt": 8,
                "from_unvisited_mt": 1,
                "from_drop_size_mt": 1,
                "visit_rate": 0.9,
                "strike_rate": 0.4,
            }
        ]
    )
    shops = pd.DataFrame(
        [
            {"city": "Karachi", "ams_3m": 12.0},
            {"city": "Karachi", "ams_3m": 6.0},
        ]
    )
    out = city_driver_table(units, shops)
    assert float(out.iloc[0]["AMS (MT)"]) == 18.0
    assert list(out.columns).index("Universe") == list(out.columns).index("Gap (MT)") + 1


def test_action_buckets_add_back_to_ask():
    from sndintel.action import ACTION_HOLD, ACTION_LIFT, ACTION_RECOVER, action_buckets
    from sndintel.monday import count_ask, country_action_table, fmt_kg

    shops = pd.DataFrame(
        [
            {"action": ACTION_CALL, "coming_due": False, "week_target_mt": 1.0, "expected_mt": 2, "ams_3m": 2, "billed_mt": 0.2, "remaining_mt": 1.8},
            {"action": ACTION_CONVERT, "coming_due": False, "week_target_mt": 2.0, "expected_mt": 3, "ams_3m": 3, "billed_mt": 0.5, "remaining_mt": 2.5},
            {"action": ACTION_LIFT, "coming_due": False, "week_target_mt": 0.5, "expected_mt": 1, "ams_3m": 1, "billed_mt": 0.4, "remaining_mt": 0.6},
            {"action": ACTION_RECOVER, "coming_due": False, "week_target_mt": 0.3, "expected_mt": 1, "ams_3m": 1, "billed_mt": 0.0, "remaining_mt": 1.0},
            {"action": ACTION_HOLD, "coming_due": True, "week_target_mt": 0.2, "expected_mt": 1, "ams_3m": 1, "billed_mt": 0.8, "remaining_mt": 0.2},
        ]
    )
    b = action_buckets(shops)
    parts = b["ask_call"] + b["ask_convert"] + b["ask_lift"] + b["ask_lapse"] + b["ask_coming"]
    assert abs(parts - b["week_target_mt"]) < 1e-9
    assert abs(b["ask_doors"] + b["ask_coming"] - b["week_target_mt"]) < 1e-9
    country = country_action_table(shops)
    assert country.iloc[0]["Unvisited"] == count_ask(1, 1.0)
    assert country.iloc[0]["Coming due"] == count_ask(1, 0.2)
    assert "Doors to visit" in country.columns
    assert "Billed (MT)" in country.columns
    assert abs(float(country.iloc[0]["Billed (MT)"]) - 1.9) < 1e-9
    assert fmt_kg(67.665) == "67,665"


def test_store_sort_clubs_dsr_by_total_ask():
    from sndintel.monday import store_table

    shops = pd.DataFrame(
        [
            {"store_name": "A2", "dsr_name": "A", "city": "Karachi", "distributor": "Eva", "week_target_mt": 0.1, "ams_3m": 1, "billed_mt": 0, "call_status": "Unvisited", "instruction": "a2"},
            {"store_name": "A1", "dsr_name": "A", "city": "Karachi", "distributor": "Eva", "week_target_mt": 2.0, "ams_3m": 1, "billed_mt": 0, "call_status": "Unvisited", "instruction": "a1"},
            {"store_name": "B1", "dsr_name": "B", "city": "Karachi", "distributor": "Eva", "week_target_mt": 1.0, "ams_3m": 1, "billed_mt": 0, "call_status": "Unvisited", "instruction": "b1"},
            {"store_name": "B2", "dsr_name": "B", "city": "Karachi", "distributor": "Eva", "week_target_mt": 0.4, "ams_3m": 1, "billed_mt": 0, "call_status": "Unvisited", "instruction": "b2"},
        ]
    )
    table = store_table(shops)
    assert list(table["Shop"]) == ["A1", "A2", "B1", "B2"]
    assert table.iloc[0]["Ask rest of month (KG)"] == "2,000"


def test_monday_pdf_has_internal_destinations():
    from sndintel.action import ActionPack
    from sndintel.ops import pdf_bytes

    shops = pd.DataFrame(
        [
            {
                "store_id": "W",
                "store_name": "Whale Mart",
                "city": "Karachi",
                "distributor": "Eva Foods",
                "dsr_name": "Amir",
                "ams_3m": 2.0,
                "expected_mt": 2.0,
                "billed_mt": 0.2,
                "remaining_mt": 1.8,
                "week_target_mt": 1.0,
                "call_status": "Unvisited",
                "action": ACTION_CALL,
                "instruction": "Call Whale Mart",
                "coming_due": False,
            }
        ]
    )
    pack = ActionPack(period="2026-08", label="August 2026", headline="Karachi is the hole", raw_shops=shops, as_of_day=21, days_in_month=31, days_left=10)
    monday = build_monday_pack(
        pack,
        exec_summary={
            "model": "gpt-4.1",
            "situation": ["Country billed 10 MT against 20 expected. Karachi is the hole."],
            "focus": [{"title": "Karachi · 10 MT", "why": "Unbilled shops", "do": "Convert the due list"}],
        },
    )
    raw = pdf_bytes(monday)
    assert raw.startswith(b"%PDF")
    assert len(raw) > 800
    assert b"Top of report" in raw
    assert b"Summary of current situation" in raw
    assert b"Key focus areas" in raw
    xls = __import__("sndintel.ops", fromlist=["excel_bytes"]).excel_bytes(monday)
    from openpyxl import load_workbook
    from io import BytesIO

    wb = load_workbook(BytesIO(xls))
    assert "00 Exec" in wb.sheetnames
    assert "01 Country" in wb.sheetnames
    assert any(n.startswith("C ") for n in wb.sheetnames)
    assert any(n.startswith("S ") for n in wb.sheetnames)


def test_persist_action_pack_scores_previous_list(tmp_path):
    from sndintel.action import ActionPack

    db = tmp_path / "w.db"
    init_db(db)
    first = ActionPack(
        period="2026-08",
        headline="first",
        raw_shops=pd.DataFrame(
            [
                {
                    "store_id": "S1",
                    "store_name": "Shop 1",
                    "city": "Karachi",
                    "distributor": "Eva",
                    "dsr_name": "Amir",
                    "action": ACTION_CALL,
                    "instruction": "Call",
                    "billed_mt": 0.0,
                    "visits": 0,
                    "week_target_mt": 1.0,
                    "expected_mt": 1.0,
                    "ams_3m": 1.0,
                }
            ]
        ),
    )
    with connect(db) as conn:
        conn.execute(
            "INSERT INTO shop_month (store_id, period, volume_mt) VALUES (?, ?, ?)",
            ("S1", "2026-08", 0.0),
        )
        persist_action_pack(conn, first)
        conn.execute("UPDATE shop_month SET volume_mt = 0.6 WHERE store_id = 'S1'")
        persist_action_pack(conn, first)
        outcomes = load_outcomes(conn, "2026-08")
    assert not outcomes.empty
    assert outcomes.iloc[0]["outcome"] == "billed"


def test_loaded_pack_still_builds_beat_lists(tmp_path):
    from sndintel.action import ActionPack, persist_action_pack, load_action_pack
    from sndintel.config import DSR_DAY_CAP

    db = tmp_path / "w.db"
    init_db(db)
    rows = []
    for i in range(20):
        rows.append(
            {
                "store_id": f"S{i}",
                "store_name": f"Shop {i}",
                "dsr_name": "Amir",
                "city": "Karachi",
                "distributor": "Eva",
                "action": ACTION_CALL,
                "instruction": "Call",
                "billed_mt": 0.0,
                "visits": 0,
                "week_target_mt": 1.0 - i * 0.02,
                "value_score": 20 - i,
                "coming_due": False,
                "ams_3m": 0.4,
                "expected_mt": 0.5,
            }
        )
    pack = ActionPack(period="2026-08", label="August 2026", headline="Karachi is the hole", raw_shops=pd.DataFrame(rows))
    with connect(db) as conn:
        persist_action_pack(conn, pack)
        loaded = load_action_pack(conn, "2026-08")
    assert loaded.headline == "Karachi is the hole"
    assert loaded.raw_shops is not None and not loaded.raw_shops.empty
    beat = build_dsr_beat_pack(loaded)
    listed = beat.sheets[0][3]
    assert len(listed) == DSR_DAY_CAP
    assert "Shop 0" in set(listed["Shop"].astype(str))
