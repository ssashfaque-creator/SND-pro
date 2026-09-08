"""National / city / distributor situation cascade."""

from __future__ import annotations

from io import BytesIO
import zipfile

import pandas as pd

from sndintel.briefing import build_strategy_pack
from sndintel.hierarchy import build_hierarchy_pack
from sndintel.plan import attach_plan
from sndintel.situation_report import (
    build_field_packs,
    build_situation_pack,
    excel_bytes,
    iter_situation_sheets,
    list_situation_entities,
    pdf_bytes,
    scorecards_for_period,
    zip_field_packs,
)


def _row(store_id, period, volume, city, dist, dsr, section="A", name="Shop"):
    return {
        "store_id": store_id,
        "period": period,
        "year": int(period[:4]),
        "month": int(period[5:7]),
        "volume_mt": volume,
        "sku_count": 1,
        "billed": 1 if volume > 0 else 0,
        "distributor": dist,
        "dsr_name": dsr,
        "section": section,
        "store_name": name,
        "zone": "South" if city == "Karachi" else "Central",
        "city": city,
    }


def _stores(rows):
    seen = {}
    for r in rows:
        seen[r["store_id"]] = r
    return pd.DataFrame(
        [
            {
                "store_id": r["store_id"],
                "store_name": r["store_name"],
                "distributor": r["distributor"],
                "dsr_name": r["dsr_name"],
                "zone": r["zone"],
                "city": r["city"],
                "section": r["section"],
            }
            for r in seen.values()
        ]
    )


def _with_recent_ams(rows, volume_by_store, months=("2026-05", "2026-06", "2026-07")):
    extra = []
    seen = set()
    for r in rows:
        sid = r["store_id"]
        if sid in seen or r["period"] != "2026-08":
            continue
        seen.add(sid)
        if sid not in volume_by_store:
            continue
        vol = volume_by_store[sid]
        for per in months:
            extra.append({**r, "period": per, "year": int(per[:4]), "month": int(per[5:7]), "volume_mt": vol})
    return rows + extra


def _world():
    """Karachi misses its run-rate; Lahore holds or beats it."""
    rows = []
    rows.append(_row("K1", "2026-08", 4.0, "Karachi", "Eva Foods", "Amir", name="Kifaya", section="Clifton"))
    rows.append(_row("K1", "2025-08", 40.0, "Karachi", "Eva Foods", "Amir", name="Kifaya", section="Clifton"))
    rows.append(_row("K2", "2026-08", 3.0, "Karachi", "South Dist", "Karachi Weak", name="Quiet K", section="Korangi"))
    rows.append(_row("K2", "2025-08", 30.0, "Karachi", "South Dist", "Karachi Weak", name="Quiet K", section="Korangi"))
    rows.append(_row("L1", "2026-08", 50.0, "Lahore", "Holding Dist", "Lahore Ace", name="Big L", section="Gulberg"))
    rows.append(_row("L1", "2025-08", 40.0, "Lahore", "Holding Dist", "Lahore Ace", name="Big L", section="Gulberg"))
    rows.append(_row("L2", "2026-08", 20.0, "Lahore", "North Dist", "Lahore Steady", name="Steady L", section="Model Town"))
    rows.append(_row("L2", "2025-08", 18.0, "Lahore", "North Dist", "Lahore Steady", name="Steady L", section="Model Town"))
    rows = _with_recent_ams(rows, {"K1": 40.0, "K2": 30.0, "L1": 42.0, "L2": 19.0})
    sm = pd.DataFrame(rows)
    hier = build_hierarchy_pack(sm, _stores(rows), ledger=pd.DataFrame([{"period": "2026-08", "status": "closed"}]))
    return sm, hier


def test_national_pack_names_lagging_and_ahead_cities():
    sm, hier = _world()
    pack = build_situation_pack(hier.units, period="2026-08")
    assert pack.scope == "national"
    assert pack.headline
    assert pack.situation
    cities_lag = pack.lagging_cities
    assert not cities_lag.empty
    assert "Karachi" in set(cities_lag["City"].astype(str))
    # Lahore billed its run-rate or better — must not be on the lagging hit-list.
    assert "Lahore" not in set(cities_lag["City"].astype(str))
    ahead = pack.ahead_cities
    if not ahead.empty:
        assert "Lahore" in set(ahead["City"].astype(str))
    people = pack.lagging_people
    assert "DSR" in people.columns
    assert pack.steps is not None and not pack.steps.empty
    assert "Volume at stake (MT)" in pack.steps.columns
    assert any("potential" in p.lower() or "path" in p.lower() for p in pack.situation) or len(pack.steps) >= 1


def test_national_pack_lists_underperforming_people():
    sm, hier = _world()
    pack = build_situation_pack(hier.units, period="2026-08")
    names = set(pack.lagging_people["DSR"].astype(str)) if not pack.lagging_people.empty else set()
    # Karachi DSRs missed their run-rate.
    assert names.intersection({"Amir", "Karachi Weak"})
    # Lahore Ace held or beat Expected — not a coaching target.
    ahead = set(pack.ahead_people["DSR"].astype(str)) if not pack.ahead_people.empty else set()
    assert "Lahore Ace" not in names or "Lahore Ace" in ahead


def test_city_pack_is_scoped_and_has_local_steps():
    sm, hier = _world()
    pack = build_situation_pack(hier.units, period="2026-08", scope="city", city="Karachi")
    assert pack.scope == "city"
    assert pack.scope_label == "Karachi"
    assert "Karachi" in pack.headline or "Karachi" in pack.weather
    dists = pack.lagging_distributors
    if not dists.empty:
        assert set(dists["City"].astype(str)) <= {"Karachi"}
        assert "Holding Dist" not in set(dists["Distributor"].astype(str))
    people = pd.concat([pack.lagging_people, pack.ahead_people], ignore_index=True)
    if not people.empty:
        assert set(people["City"].astype(str)) <= {"Karachi"}
    assert not pack.steps.empty


def test_distributor_pack_lists_its_dsrs():
    sm, hier = _world()
    pack = build_situation_pack(
        hier.units, period="2026-08", scope="distributor", city="Karachi", distributor="Eva Foods"
    )
    assert pack.scope == "distributor"
    people = pd.concat([pack.lagging_people, pack.ahead_people], ignore_index=True)
    if not people.empty:
        assert set(people["Distributor"].astype(str)) <= {"Eva Foods"}
        assert "Amir" in set(people["DSR"].astype(str))


def test_field_entity_list_and_zip():
    sm, hier = _world()
    cities = list_situation_entities(hier.units, "city")
    assert "Karachi" in cities
    assert "Lahore" in cities
    dists = list_situation_entities(hier.units, "distributor")
    assert any("Eva Foods" in x for x in dists)
    packs = build_field_packs(hier.units, period="2026-08", kind="city")
    assert len(packs) >= 2
    names = {name for name, _ in packs}
    assert "Karachi" in names and "Lahore" in names
    payload = zip_field_packs(packs, fmt="pdf")
    with zipfile.ZipFile(BytesIO(payload)) as zf:
        files = zf.namelist()
    assert any("Karachi" in f and f.endswith(".pdf") for f in files)
    assert any("Lahore" in f and f.endswith(".pdf") for f in files)


def test_pdf_and_excel_are_real_files():
    sm, hier = _world()
    pack = build_situation_pack(hier.units, period="2026-08")
    pdf = pdf_bytes(pack)
    assert pdf[:4] == b"%PDF"
    assert b"Situation" in pdf or b"SND" in pdf
    xls = excel_bytes(pack)
    assert xls[:2] == b"PK"
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(xls))
    assert "00 Situation" in wb.sheetnames
    assert any("potential" in s.lower() or "lagging" in s.lower() for s in wb.sheetnames)


def test_situation_pack_does_not_need_strategy_pack():
    """Cascade is built from unit scorecards, not the glossary-first strategy PDF."""
    sm, hier = _world()
    strategy = build_strategy_pack(hier.units, sm, period="2026-08")
    sit = build_situation_pack(hier.units, period="2026-08")
    assert "Situation" not in strategy.cities.columns
    assert "Situation" in sit.lagging_cities.columns or sit.lagging_cities.empty
    assert sit.kpis.get("gap_mt", 0) >= 0


def _mtd_world():
    """September 2026 open through day 8 of a 30-day month."""
    rows = []
    rows.append(_row("K1", "2026-09", 8.0, "Karachi", "Eva Foods", "Amir", name="Kifaya", section="Clifton"))
    rows.append(_row("K1", "2025-09", 40.0, "Karachi", "Eva Foods", "Amir", name="Kifaya", section="Clifton"))
    rows.append(_row("K2", "2026-09", 6.0, "Karachi", "South Dist", "Karachi Weak", name="Quiet K", section="Korangi"))
    rows.append(_row("K2", "2025-09", 30.0, "Karachi", "South Dist", "Karachi Weak", name="Quiet K", section="Korangi"))
    rows.append(_row("L1", "2026-09", 12.0, "Lahore", "Holding Dist", "Lahore Ace", name="Big L", section="Gulberg"))
    rows.append(_row("L1", "2025-09", 40.0, "Lahore", "Holding Dist", "Lahore Ace", name="Big L", section="Gulberg"))
    rows.append(_row("L2", "2026-09", 5.0, "Lahore", "North Dist", "Lahore Steady", name="Steady L", section="Model Town"))
    rows.append(_row("L2", "2025-09", 18.0, "Lahore", "North Dist", "Lahore Steady", name="Steady L", section="Model Town"))
    extra = []
    for r in list(rows):
        if r["period"] != "2026-09":
            continue
        hist = {"K1": 40.0, "K2": 30.0, "L1": 42.0, "L2": 19.0}[r["store_id"]]
        for per in ("2026-06", "2026-07", "2026-08"):
            extra.append({**r, "period": per, "year": int(per[:4]), "month": int(per[5:7]), "volume_mt": hist})
    sm = pd.DataFrame(rows + extra)
    ledger = pd.DataFrame(
        [
            {"period": "2026-08", "status": "closed", "as_of_day": 31, "days_in_month": 31},
            {
                "period": "2026-09",
                "status": "mtd_open",
                "as_of_day": 8,
                "days_in_month": 30,
                "execution_date": "2026-09-08",
            },
        ]
    )
    hier = build_hierarchy_pack(sm, _stores(rows + extra), ledger=ledger, period="2026-09")
    targets = pd.DataFrame(
        [
            {
                "store_id": "K1",
                "store_name": "Kifaya",
                "city": "Karachi",
                "distributor": "Eva Foods",
                "dsr_name": "Amir",
                "section": "Clifton",
                "target_mt": 50.0,
                "match_method": "id",
            },
            {
                "store_id": "K2",
                "store_name": "Quiet K",
                "city": "Karachi",
                "distributor": "South Dist",
                "dsr_name": "Karachi Weak",
                "section": "Korangi",
                "target_mt": 40.0,
                "match_method": "id",
            },
            {
                "store_id": "L1",
                "store_name": "Big L",
                "city": "Lahore",
                "distributor": "Holding Dist",
                "dsr_name": "Lahore Ace",
                "section": "Gulberg",
                "target_mt": 55.0,
                "match_method": "id",
            },
            {
                "store_id": "L2",
                "store_name": "Steady L",
                "city": "Lahore",
                "distributor": "North Dist",
                "dsr_name": "Lahore Steady",
                "section": "Model Town",
                "target_mt": 25.0,
                "match_method": "id",
            },
        ]
    )
    pace = 1.0
    if hier.units is not None and not hier.units.empty and "intra_month_frac" in hier.units.columns:
        pace = float(pd.to_numeric(hier.units["intra_month_frac"], errors="coerce").dropna().max() or 1.0)
    units = attach_plan(hier.units, targets, pace=pace)
    return sm, units, ledger, pace


def test_open_mtd_label_is_not_a_calendar_date():
    sm, units, ledger, _pace = _mtd_world()
    pack = build_situation_pack(units, ledger=ledger, period="2026-09")
    blob = " ".join([pack.label, pack.weather, pack.headline] + list(pack.situation) + list(pack.plan_lines))
    assert "8/30" not in blob
    assert "8/30" not in pack.label
    assert "Sep 2026" in pack.label
    assert "8 Sep" in pack.label
    assert "30-day" in pack.label


def test_mtd_projects_month_end_against_full_monthly_target():
    sm, units, ledger, pace = _mtd_world()
    pack = build_situation_pack(units, ledger=ledger, period="2026-09")
    kpis = pack.kpis
    billed = float(kpis["billed_mt"])
    assert kpis["open_mtd"] is True
    assert kpis["target_mt"] == 170
    assert kpis["target_mt"] != kpis.get("target_paced_mt")
    assert abs(kpis["projected_mt"] - billed / pace) < 0.6
    assert abs(kpis["vs_target_mt"] - (kpis["projected_mt"] - 170)) < 0.6
    assert "Projected month-end (MT)" in pack.lagging_cities.columns or pack.lagging_cities.empty or "Projected month-end (MT)" in pack.ahead_cities.columns
    assert pack.plan_lines
    assert any("Projected month-end" in line for line in pack.plan_lines)
    assert any("monthly target" in line.lower() for line in pack.plan_lines)


def test_do_this_is_short_and_plain():
    sm, units, ledger, _pace = _mtd_world()
    pack = build_situation_pack(units, ledger=ledger, period="2026-09")
    blobs = []
    for df in (pack.lagging_cities, pack.lagging_distributors, pack.lagging_people, pack.steps, pack.ahead_cities):
        if df is None or df.empty:
            continue
        col = "Do this" if "Do this" in df.columns else ("Do this week" if "Do this week" in df.columns else None)
        if not col:
            continue
        blobs.extend(str(x) for x in df[col].tolist())
    joined = " ".join(blobs)
    assert "Drivers:" not in joined
    assert "velocity" not in joined.lower()
    assert "depletion clock" not in joined.lower()
    for text in blobs:
        assert len(text) < 220


def test_weak_areas_are_gone():
    sm, hier = _world()
    pack = build_situation_pack(hier.units, period="2026-08")
    headings = [heading for _sheet, heading, _note, _df in iter_situation_sheets(pack)]
    assert not any("weak" in h.lower() for h in headings)
    xls = excel_bytes(pack)
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(xls))
    assert not any("weak" in s.lower() for s in wb.sheetnames)
    assert any("plan" in s.lower() and "action" in s.lower() for s in wb.sheetnames)


def test_closed_month_cut_has_no_projection():
    sm, units, ledger, _pace = _mtd_world()
    closed_units = scorecards_for_period(sm, _stores(sm.to_dict("records")), ledger, "2026-08")
    pack = build_situation_pack(closed_units, ledger=ledger, period="2026-08")
    assert pack.kpis["open_mtd"] is False
    assert "closed month" in pack.label.lower()
    assert pack.lagging_cities.empty or "Projected month-end (MT)" not in pack.lagging_cities.columns
    assert abs(pack.kpis["projected_mt"] - pack.kpis["billed_mt"]) < 0.05
    assert pack.this_week is None or pack.this_week.empty
