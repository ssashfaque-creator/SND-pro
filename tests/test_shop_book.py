"""Shop-wise issues pack: reuse action Expected / Ask; MTD vs closed cuts."""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from sndintel.action import ActionPack, build_action_pack
from sndintel.shop_book import build_shop_book, excel_bytes, list_shop_book_entities, pdf_bytes

from test_situation_report import _mtd_world, _stores, _world


def test_closed_month_lists_missed_expected_largest_hole_first():
    sm, _hier = _world()
    ledger = pd.DataFrame([{"period": "2026-08", "status": "closed"}])
    action = build_action_pack(sm, _stores(sm.to_dict("records")), ledger=ledger, period="2026-08")
    book = build_shop_book(action=action, ledger=ledger, period="2026-08", scope="national")
    assert book.open_mtd is False
    assert "this week" not in (book.headline + book.weather).lower()
    issues = book.issues
    assert not issues.empty
    assert "Kifaya" in set(issues["Shop"].astype(str))
    assert "Quiet K" in set(issues["Shop"].astype(str))
    # Lahore Ace beat or held Expected — not an issue.
    assert "Big L" not in set(issues["Shop"].astype(str))
    gap = pd.to_numeric(issues["Gap (MT)"], errors="coerce").fillna(0)
    assert gap.is_monotonic_decreasing or len(issues) <= 1
    blob = " ".join(issues["Comment"].astype(str).tolist()).lower()
    assert "this week" not in blob
    assert "next month" in blob
    assert "Ask (KG)" not in issues.columns
    assert "Usual drop (MT)" in issues.columns
    assert "Gap (MT)" in issues.columns


def test_closed_expected_matches_action_pack():
    sm, _hier = _world()
    ledger = pd.DataFrame([{"period": "2026-08", "status": "closed"}])
    action = build_action_pack(sm, _stores(sm.to_dict("records")), ledger=ledger, period="2026-08")
    book = build_shop_book(action=action, ledger=ledger, period="2026-08")
    left = action.raw_shops.set_index("store_id")["expected_mt"].astype(float)
    right = book.raw.set_index("store_id")["expected_mt"].astype(float)
    assert set(left.index) == set(right.index)
    pd.testing.assert_series_equal(left.sort_index(), right.sort_index(), check_names=False)


def test_city_and_dsr_scope():
    sm, _hier = _world()
    ledger = pd.DataFrame([{"period": "2026-08", "status": "closed"}])
    action = build_action_pack(sm, _stores(sm.to_dict("records")), ledger=ledger, period="2026-08")
    karachi = build_shop_book(
        action=action, ledger=ledger, period="2026-08", scope="city", city="Karachi"
    )
    assert set(karachi.raw["city"].astype(str)) <= {"Karachi"}
    assert "Big L" not in set(karachi.shops["Shop"].astype(str))
    assert "Distributor" in karachi.shops.columns
    assert "City" not in karachi.shops.columns
    amir = build_shop_book(
        action=action,
        ledger=ledger,
        period="2026-08",
        scope="dsr",
        city="Karachi",
        distributor="Eva Foods",
        dsr="Amir",
    )
    assert set(amir.shops["Shop"].astype(str)) <= {"Kifaya"}
    cities = list_shop_book_entities(action.raw_shops, "city")
    assert "Karachi" in cities and "Lahore" in cities
    dsrs = list_shop_book_entities(action.raw_shops, "dsr")
    assert any("Amir" in x for x in dsrs)


def test_mtd_does_not_flag_still_to_expected_as_the_issue_list():
    sm, _units, ledger, _pace = _mtd_world()
    stores = _stores(sm.to_dict("records"))
    action = build_action_pack(sm, stores, ledger=ledger, period="2026-09")
    book = build_shop_book(action=action, ledger=ledger, period="2026-09")
    assert book.open_mtd is True
    assert "Ask (KG)" in book.shops.columns
    assert "Do this" in book.shops.columns
    assert "Gap (MT)" not in book.shops.columns
    remaining = pd.to_numeric(book.raw["remaining_mt"], errors="coerce").fillna(0)
    # Mid-month almost every shop still has a full-month hole. That must not be the issue list.
    behind = int((remaining > 0.05).sum())
    n_issues = 0 if book.issues.empty else len(book.issues)
    assert behind >= n_issues
    if not book.issues.empty:
        assert set(book.issues["Issue"].astype(str)) <= {
            "Due · unvisited",
            "Due · no bill",
            "Due · light drop",
            "Visited, no bill",
            "Lapsed",
            "Unvisited",
        }
    ask_left = action.raw_shops.set_index("store_id")["week_target_mt"].astype(float)
    ask_right = book.raw.set_index("store_id")["week_target_mt"].astype(float)
    pd.testing.assert_series_equal(ask_left.sort_index(), ask_right.sort_index(), check_names=False)


def test_mtd_attaches_matched_shop_target_without_replacing_expected():
    sm, _units, ledger, _pace = _mtd_world()
    stores = _stores(sm.to_dict("records"))
    action = build_action_pack(sm, stores, ledger=ledger, period="2026-09")
    targets = pd.DataFrame(
        [
            {"store_id": "K1", "store_name": "Kifaya", "target_mt": 50.0, "match_method": "id"},
            {"store_id": "L1", "store_name": "Big L", "target_mt": 55.0, "match_method": "id"},
        ]
    )
    book = build_shop_book(action=action, shop_targets=targets, ledger=ledger, period="2026-09")
    raw = book.raw.set_index("store_id")
    assert float(raw.loc["K1", "shop_target_mt"]) == 50.0
    assert float(raw.loc["K1", "expected_mt"]) != 50.0
    assert float(raw.loc["K1", "expected_mt"]) == float(action.raw_shops.set_index("store_id").loc["K1", "expected_mt"])


def _closed_pack(rows):
    shops = pd.DataFrame(rows)
    if "remaining_mt" not in shops.columns:
        shops["remaining_mt"] = (
            pd.to_numeric(shops["expected_mt"], errors="coerce").fillna(0)
            - pd.to_numeric(shops["billed_mt"], errors="coerce").fillna(0)
        ).clip(lower=0)
    return ActionPack(period="2026-08", label="Aug 2026", open_mtd=False, raw_shops=shops)


def test_unbilled_is_billed_zero_not_a_tiny_invoice():
    """0.05 MT was used as 'no bill', so Unbilled showed billed volume on the Karachi mix."""
    pack = _closed_pack(
        [
            {
                "store_id": "U1",
                "store_name": "Visited Zero",
                "city": "Karachi",
                "distributor": "Dist",
                "dsr_name": "Ali",
                "billed_mt": 0.0,
                "expected_mt": 2.0,
                "call_status": "Visited · not billed",
                "is_lapsed": False,
                "week_target_mt": 0.0,
            },
            {
                "store_id": "T1",
                "store_name": "Tiny Invoice",
                "city": "Karachi",
                "distributor": "Dist",
                "dsr_name": "Ali",
                "billed_mt": 0.02,
                "expected_mt": 2.92,
                "call_status": "Billed",
                "is_lapsed": False,
                "week_target_mt": 0.0,
            },
        ]
    )
    book = build_shop_book(action=pack, period="2026-08")
    by_name = book.raw.set_index("store_name")["issue"].astype(str)
    assert by_name["Visited Zero"] == "Unbilled"
    assert by_name["Tiny Invoice"] == "Missed Expected"
    mix = book.mix.set_index("Issue")
    assert float(mix.loc["Unbilled", "Billed (MT)"]) == 0.0
    assert float(mix.loc["Unbilled", "Expected (MT)"]) == 2.0
    assert "visited but not billed" in str(
        book.issues.loc[book.issues["Shop"] == "Visited Zero", "Comment"].iloc[0]
    ).lower()
    tiny_comment = str(book.issues.loc[book.issues["Shop"] == "Tiny Invoice", "Comment"].iloc[0]).lower()
    assert "visited but not billed" not in tiny_comment
    assert "2.92" in tiny_comment


def test_no_run_rate_is_not_a_row_with_expected():
    """0.05 MT was used as 'no Expected', so the mix showed Expected on No Expected."""
    pack = _closed_pack(
        [
            {
                "store_id": "Q1",
                "store_name": "Quiet",
                "city": "Karachi",
                "distributor": "Dist",
                "dsr_name": "Ali",
                "billed_mt": 0.02,
                "expected_mt": 0.03,
                "call_status": "Billed",
                "is_lapsed": False,
                "week_target_mt": 0.0,
            },
            {
                "store_id": "Z1",
                "store_name": "Universe",
                "city": "Karachi",
                "distributor": "Dist",
                "dsr_name": "Ali",
                "billed_mt": 0.0,
                "expected_mt": 0.0,
                "call_status": "Unvisited",
                "is_lapsed": False,
                "week_target_mt": 0.0,
            },
            {
                "store_id": "M1",
                "store_name": "Real Miss",
                "city": "Karachi",
                "distributor": "Dist",
                "dsr_name": "Ali",
                "billed_mt": 1.0,
                "expected_mt": 4.0,
                "call_status": "Billed",
                "is_lapsed": False,
                "week_target_mt": 0.0,
            },
        ]
    )
    book = build_shop_book(action=pack, period="2026-08")
    by_name = book.raw.set_index("store_name")["issue"].astype(str)
    assert by_name["Universe"] == "No run-rate"
    assert by_name["Quiet"] == "On Expected"
    assert by_name["Real Miss"] == "Missed Expected"
    assert "No Expected" not in set(book.mix["Issue"].astype(str))
    assert "No run-rate" not in set(book.mix["Issue"].astype(str))
    assert book.kpis.get("n_quiet") == 1


def test_pdf_and_excel_are_real_files():
    sm, _hier = _world()
    ledger = pd.DataFrame([{"period": "2026-08", "status": "closed"}])
    action = build_action_pack(sm, _stores(sm.to_dict("records")), ledger=ledger, period="2026-08")
    book = build_shop_book(action=action, ledger=ledger, period="2026-08")
    pdf = pdf_bytes(book)
    assert pdf[:4] == b"%PDF"
    xls = excel_bytes(book)
    assert xls[:2] == b"PK"
    wb = load_workbook(BytesIO(xls))
    assert "00 Cover" in wb.sheetnames
    assert "02 Issues" in wb.sheetnames
    assert "03 All shops" in wb.sheetnames
